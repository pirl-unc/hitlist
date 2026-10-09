"""Offline, checksum-pinned curation of the original canine/DLA MS worksheets."""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
import tempfile
from pathlib import Path

from .curation_yaml import load_curation_yaml
from .provenance import file_digest

_PROFILE_PATH = Path(__file__).parent / "data" / "canine_ligands.yaml"
_AMINO_ACIDS = frozenset("ACDEFGHIKLMNPQRSTVWY")


def _json(path, value):
    path.write_text(json.dumps(value, sort_keys=True, indent=2, allow_nan=False) + "\n")


def _check_asset(path, asset):
    expected = {key: asset[key] for key in ("sha256", "size_bytes")}
    if file_digest(path) != expected:
        raise ValueError(f"Canine source checksum/size mismatch: {path.name}")
    md5 = hashlib.md5()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024**2), b""):
            md5.update(block)
    if md5.hexdigest() != asset["md5"]:
        raise ValueError(f"Canine source published MD5 mismatch: {path.name}")


def _rows(sheet, profile, asset, arm):
    header = tuple(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    if header != tuple(asset["header"]):
        raise ValueError(f"Unexpected canine worksheet header: {asset['file']} / {sheet.title}")
    n_observations = 0
    for number, values in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
        if all(value is None for value in values):
            continue
        sequence, length, accessions, method = values
        if (
            not isinstance(sequence, str)
            or set(sequence) - _AMINO_ACIDS
            or not profile["sequence_length_min"] <= len(sequence) <= profile["sequence_length_max"]
            or isinstance(length, bool)
            or length != len(sequence)
            or method != profile["identification_method"]
            or not isinstance(accessions, str)
            or not accessions.strip()
        ):
            raise ValueError(
                f"Invalid canine source row: {asset['file']} / {sheet.title} / {number}"
            )
        n_observations += 1
        if n_observations > arm["n_observations"]:
            raise ValueError(f"Canine worksheet exceeds reviewed row count: {sheet.title}")
        yield {
            "peptide": sequence,
            "mhc_class": profile["mhc_class"],
            "mhc_restriction": arm["mhc_restriction"],
            "attributed_sample_label": arm["sample_label"],
            "condition_id": arm["condition_id"],
            "donor_label": arm["donor_label"],
            "histology": arm["histology"],
            "ip_antibody": arm["ip_antibody"],
            "source_species": arm["defaults"]["source_organism"],
            "host_species": arm["defaults"]["host"],
            "mhc_species": profile["mhc_species"],
            "source_url": asset["url"],
            "source_file": asset["file"],
            "source_sha256": asset["sha256"],
            "sheet": sheet.title,
            "source_row": number,
            "source_protein_mappings": accessions,
            "identification_method": method,
            "reported_fdr_threshold": profile["reported_fdr_threshold"],
            "fdr_scope": profile["fdr_scope"],
            "peptide_q_value": "",
            "spectrum_id": "",
            "raw_accession": profile["raw_accession"],
            "il_ambiguity": "unresolved",
            "modifications": "not_resolved_in_summary_table",
            "context": arm["context"],
            "reported_dla_typing": arm["reported_dla_typing"],
            "construct_description": arm["construct_description"],
            "search_database": arm["search_database"],
        }
    if n_observations != arm["n_observations"]:
        raise ValueError(f"Canine worksheet row count mismatch: {sheet.title}")


def curate_canine_ligands(directory, output_dir):
    """Curate the six local PMID 42199926 workbooks without downloading.

    Returns the path to a JSON list accepted by ``scan_supplementary(entries=,
    directory=, allow_download=False)``. Original worksheets are required at
    their pinned hashes. Every original 8-30-residue observation is retained;
    human-TAA comparison worksheets, flow assays and predictions are excluded.
    Individual q-values, spectra and I/L discrimination remain unknown.

    Install ``hitlist[curation]`` for the optional workbook reader. Output must
    not exist. CSVs and receipts are deterministic; source/arm information lives
    in the packaged YAML profile and survives in contributor records.
    """
    try:
        import openpyxl
    except ImportError as error:
        raise ImportError("Canine workbook curation requires hitlist[curation]") from error

    directory, destination = Path(directory), Path(output_dir)
    if destination.exists():
        raise FileExistsError(destination)
    profile = load_curation_yaml(_PROFILE_PATH)
    if profile.get("schema_version") != 1:
        raise ValueError("Unsupported canine curation profile")
    for asset in profile["assets"]:
        _check_asset(directory / asset["file"], asset)
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{destination.name}-", dir=destination.parent))
    try:
        entries = []
        for asset in profile["assets"]:
            book = openpyxl.load_workbook(directory / asset["file"], read_only=True, data_only=True)
            try:
                for arm in asset["sheets"]:
                    if arm["sheet"] not in book.sheetnames:
                        raise ValueError(f"Missing reviewed canine worksheet: {arm['sheet']}")
                    rows = iter(_rows(book[arm["sheet"]], profile, asset, arm))
                    first = next(rows)
                    path = staging / arm["file"]
                    with path.open("w", newline="") as stream:
                        writer = csv.DictWriter(stream, fieldnames=list(first), lineterminator="\n")
                        writer.writeheader()
                        writer.writerow(first)
                        writer.writerows(rows)
                    entries.append(
                        {
                            "pmid": profile["pmid"],
                            "file": arm["file"],
                            **file_digest(path),
                            "study_label": profile["study_label"],
                            "source": asset["url"] + "#" + arm["sheet"],
                            "source_asset": {
                                key: value for key, value in asset.items() if key != "sheets"
                            },
                            "source_version": profile["source_version"],
                            "curated_arm": arm,
                            "defaults": arm["defaults"],
                        }
                    )
            finally:
                book.close()
        _json(staging / "manifest.json", entries)
        _json(
            staging / "curation.json",
            {
                "schema_version": 1,
                "pmid": profile["pmid"],
                "profile": profile,
                "n_observations": sum(
                    arm["n_observations"] for a in profile["assets"] for arm in a["sheets"]
                ),
                "artifacts": {path.name: file_digest(path) for path in sorted(staging.iterdir())},
            },
        )
        for asset in profile["assets"]:
            _check_asset(directory / asset["file"], asset)
        destination.mkdir()
        try:
            for path in staging.iterdir():
                if path.name != "manifest.json":
                    path.replace(destination / path.name)
            (staging / "manifest.json").replace(destination / "manifest.json")
        except BaseException:
            shutil.rmtree(destination)
            raise
    finally:
        shutil.rmtree(staging)
    return destination / "manifest.json"
