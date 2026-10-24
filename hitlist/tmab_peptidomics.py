"""Lossless, source-grain native cellular peptide measurements.

The Dasgupta and Gelman supplements report relative isotope-channel peak
ratios from whole-cell peptide extracts.  This module keeps those ratios
separate from MHC presentation and from inferred processing labels.  Source
workbooks are caller-supplied, hash checked, and copied into each immutable
package so that row-level tables can always be replayed from their bytes.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import shutil
import tempfile
from collections import Counter
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from .evidence_snapshot import snapshot_file
from .provenance import file_digest

PACKAGE_SCHEMA_VERSION = 17
CATALOG_SCHEMA_VERSION = 22
REVISION = "tmab_native_peptidomics_v3"
ROOT = Path("tmab")
SOURCE_ROOT = ROOT / "source_assets"
SOURCE_TABLE = ROOT / "source_assets.parquet"
ROWS_TABLE = ROOT / "source_rows.parquet"
MEASUREMENTS_TABLE = ROOT / "measurements.parquet"
CONTEXT_TABLE = ROOT / "protocol_context.parquet"
LINEAGE_TABLE = ROOT / "lineage.parquet"
COVERAGE = ROOT / "coverage.json"
ASSETS = Path(__file__).parent / "data/evidence_revisions/tmab_assets.json"

_PRIMARY = {
    "dasgupta_s1.xlsx": {
        "sheet": "Table S1",
        "header_row": 7,
        "data_row": 8,
        "cell": 1,
        "experiment": 2,
        "charge": 3,
        "observed_mass": 5,
        "theoretical_mass": 6,
        "protein": 7,
        "gene": 8,
        "location": 9,
        "upstream": 10,
        "sequence": 11,
        "downstream": 12,
        "comments": 13,
        "analysis": 14,
        "treatment": 14,
        "average_annotation": 22,
        "ratio_columns": (
            (15, "treated_vs_control", "Rep1"),
            (16, "treated_vs_control", "Rep2"),
            (17, "treated_vs_control", "Rep3"),
            (19, "control_vs_control", "Rep1"),
            (20, "control_vs_control", "Rep2"),
            (21, "treated_vs_control_average", "average"),
        ),
    },
    "gelman_s1.xlsx": {
        "sheet": "Sheet1",
        "header_row": 9,
        "data_row": 10,
        "cell": 1,
        "experiment": 2,
        "charge": 3,
        "observed_mass": 5,
        "theoretical_mass": 6,
        "protein": 7,
        "gene": None,
        "location": 8,
        "upstream": 9,
        "sequence": 10,
        "downstream": 11,
        "comments": 12,
        "analysis": 13,
        "treatment": 13,
        "average_annotation": 21,
        "ratio_columns": (
            (14, "treated_vs_control", "Rep1"),
            (15, "treated_vs_control", "Rep2"),
            (16, "treated_vs_control", "Rep3"),
            (18, "control_vs_control", "Rep1"),
            (19, "control_vs_control", "Rep2"),
            (20, "treated_vs_control_average", "average"),
        ),
    },
}

_HEATMAP = {
    "dasgupta_s2.xlsx": {"sheet": "Table S2", "data_row": 18, "sequence": 5},
    "gelman_s2.xlsx": {"sheet": "Sheet1", "data_row": 15, "sequence": 6},
}

_STUDIES = {
    "dasgupta_2014": {
        "doi": "10.1371/journal.pone.0103604",
        "url": "https://journals.plos.org/plosone/article?id=10.1371/journal.pone.0103604",
        "cells": "HEK293T and SH-SY5Y",
        "protocol_instrument": "Waters Synapt G2",
        "protocol_section": "s2b2",
        "treatment_section": "s2b1",
    },
    "gelman_2013": {
        "doi": "10.1371/journal.pone.0053263",
        "url": "https://journals.plos.org/plosone/article?id=10.1371/journal.pone.0053263",
        "cells": "HEK293T and SH-SY5Y",
        "protocol_instrument": "Waters Synapt G1",
        "protocol_section": "s2d",
        "treatment_section": "s2c",
    },
}


def _schema(columns, numeric=(), integer=()):
    numeric = set(numeric)
    integer = set(integer)
    return pa.schema(
        [
            (
                name,
                pa.float64() if name in numeric else pa.int64() if name in integer else pa.string(),
            )
            for name in columns
        ]
    )


SOURCE_COLUMNS = (
    "source_asset_id",
    "source_id",
    "source_file",
    "source_sha256",
    "source_asset_url",
    "source_role",
)
ROW_COLUMNS = (
    "source_record_id",
    "source_asset_id",
    "source_id",
    "source_file",
    "sheet_name",
    "source_row_index",
    "source_row_locator",
    "source_row_json",
    "source_row_disposition",
    "source_sequence_raw",
    "source_handling_raw",
    "source_observed_mass_raw",
    "source_experiment_raw",
    "source_cell_raw",
    "source_comments_raw",
    "row_identity_status",
)
CONTEXT_COLUMNS = (
    "context_id",
    "study_id",
    "source_id",
    "source_asset_id",
    "source_file",
    "publication_doi",
    "publication_url",
    "organism_name",
    "organism_taxon_id",
    "cell_line_name",
    "measured_pool",
    "extraction_method",
    "treatment_label_raw",
    "treatment_name",
    "treatment_dose",
    "treatment_dose_unit",
    "exposure_duration_raw",
    "duration_hours",
    "duration_role",
    "average_annotation_raw",
    "incubation_duration_hours",
    "harvest_duration_hours",
    "harvest_compound_status",
    "timeline_status",
    "timeline_evidence_url",
    "labeling_method",
    "instrument",
    "publishing_protocol_instrument",
    "instrument_status",
    "protocol_evidence_url",
    "acquisition_mode",
    "assay_method",
    "assay_readout",
    "normalization_policy",
    "context_scope",
    "context_status",
    "wash_interval_raw",
    "wash_presence_status",
    "exposure_interval_status",
    "acquisition_source_scope",
    "source_study_inheritance_status",
)
MEASUREMENT_COLUMNS = (
    "measurement_id",
    "source_record_id",
    "context_id",
    "lineage_group_id",
    "candidate_reuse_key",
    "source_asset_id",
    "source_id",
    "source_file",
    "source_row_index",
    "source_cell",
    "ratio_role",
    "replicate_label",
    "channel_label_raw",
    "peptide_sequence_raw",
    "peptide_sequence",
    "peptide_modifications_raw",
    "peptide_modifications_json",
    "sequence_identity_status",
    "modification_status",
    "parent_label_raw",
    "gene_name",
    "cell_line_name",
    "subcellular_location_raw",
    "charge_raw",
    "observed_mass_raw",
    "theoretical_mass_raw",
    "measured_pool",
    "quantitative_raw",
    "quantitative_value",
    "measurement_unit",
    "measurement_inequality",
    "censoring_status",
    "value_status",
    "peak_overlap_status",
    "eligibility_status",
    "source_reuse_status",
    "acquisition_provenance_status",
    "shared_control_group_id",
    "shared_control_group_status",
    "source_comments_raw",
)
LINEAGE_COLUMNS = (
    "lineage_group_id",
    "measurement_id",
    "source_record_id",
    "source_id",
    "cross_paper_match_key",
    "lineage_status",
    "reuse_status",
    "relation_basis",
    "independent_specimen_status",
    "candidate_reuse_group_id",
    "experiment_label_agreement",
)

SOURCE_SCHEMA = _schema(SOURCE_COLUMNS)
ROW_SCHEMA = _schema(ROW_COLUMNS, integer=("source_row_index",))
CONTEXT_SCHEMA = _schema(
    CONTEXT_COLUMNS,
    numeric=(
        "treatment_dose",
        "duration_hours",
        "incubation_duration_hours",
        "harvest_duration_hours",
    ),
)
MEASUREMENT_SCHEMA = _schema(
    MEASUREMENT_COLUMNS, numeric=("quantitative_value",), integer=("source_row_index",)
)
LINEAGE_SCHEMA = _schema(LINEAGE_COLUMNS)


def _asset_specs():
    return json.loads(ASSETS.read_text())["assets"]


def _digest(path):
    return file_digest(path)["sha256"]


def _source_path(directory, filename):
    directory = Path(directory)
    for candidate in (
        directory / filename,
        directory / "assets" / filename,
        directory / SOURCE_ROOT / filename,
    ):
        if candidate.is_file():
            return candidate
    raise FileNotFoundError(f"Missing TMAB source workbook: {filename}")


def _id(kind, value):
    return f"tmab:{kind}:{hashlib.sha256(str(value).encode()).hexdigest()[:24]}"


def _json_value(value):
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _row_json(values):
    return json.dumps(
        [_json_value(value) for value in values], ensure_ascii=False, separators=(",", ":")
    )


def _text(value):
    return None if value is None else str(value).strip() or None


def _sequence_identity(raw):
    """Parse only supported source notation; never extract annotation letters."""
    text = _text(raw)
    if not text:
        return {
            "sequence": None,
            "modifications": [],
            "summary": None,
            "raw_modifications": None,
            "identification": "unidentified",
            "modification_status": "unresolved",
        }
    tentative = text.startswith("(") and text.endswith(")")
    if tentative:
        text = text[1:-1].strip()
    modifications = []
    summary = []
    tokens = []
    if text.startswith("Ac-"):
        modifications.append(
            {"name": "N_terminal_acetylation", "position": 0, "source_token": "Ac-"}
        )
        summary.append("N_terminal_acetylation")
        tokens.append("Ac-")
        text = text[3:]
    residues = []
    while text:
        if text.startswith("Mox"):
            residues.append("M")
            position = len(residues)
            modifications.append({"name": "oxidation", "position": position, "source_token": "Mox"})
            summary.append(f"oxidation@{position}")
            tokens.append(f"Mox@{position}")
            text = text[3:]
        elif text[0] in "ACDEFGHIKLMNPQRSTVWY":
            residues.append(text[0])
            text = text[1:]
        else:
            return {
                "sequence": None,
                "modifications": [],
                "summary": None,
                "raw_modifications": _text(raw),
                "identification": "unresolved_annotation",
                "modification_status": "unresolved_source_notation",
            }
    if not residues:
        return {
            "sequence": None,
            "modifications": [],
            "summary": None,
            "raw_modifications": _text(raw),
            "identification": "unresolved_annotation",
            "modification_status": "unresolved_source_notation",
        }
    return {
        "sequence": "".join(residues),
        "modifications": modifications,
        "summary": ";".join(summary) or None,
        "raw_modifications": ";".join(tokens) or None,
        "identification": "tentative" if tentative else "reported_sequence",
        "modification_status": "supported_source_notation" if modifications else "none_reported",
    }


def _parse_sequence(raw):
    identity = _sequence_identity(raw)
    return identity["sequence"], identity["summary"]


def _eligibility(identity, role, value, censoring):
    if role != "treated_vs_control":
        return "deferred_control_or_derived_average"
    if identity["identification"] != "reported_sequence":
        return "deferred_" + identity["identification"]
    if value is None or not math.isfinite(value):
        return "deferred_nonquantifiable"
    if value <= 0:
        return "deferred_nonpositive_ratio"
    if censoring == "exact":
        return "candidate_exact"
    if censoring == "source_declared_bound":
        return "candidate_censored"
    return "deferred_nonquantifiable"


def _parse_ratio(raw):
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None, None, "missing_signal", "missing_source_value", "not_overlapping"
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return float(raw), None, "exact", "reported_exact", "not_overlapping"
    text = str(raw).strip()
    lowered = text.lower()
    if lowered in {"po", "peak overlap", "overlap", "can't quantify"}:
        return None, None, "overlapping_peak", "reported_overlap", "explicit_overlap"
    match = re.fullmatch(r"([<>])\s*([0-9]+(?:\.[0-9]+)?)", text)
    if match:
        return (
            float(match.group(2)),
            {"<": "lt", ">": "gt"}[match.group(1)],
            "source_declared_bound",
            "reported_bound",
            "not_overlapping",
        )
    if lowered in {"nd", "n/d", "na", "n.a."}:
        return None, None, "missing_signal", "reported_not_detected", "not_overlapping"
    try:
        return float(text), None, "exact", "reported_exact", "not_overlapping"
    except ValueError:
        return None, None, "unresolved_text", "reported_nonnumeric", "not_overlapping"


def _treatment(raw):
    text = _text(raw) or ""
    label = re.split(r"[/=]", text, maxsplit=1)[0].strip()
    name_match = re.match(r"([A-Za-z][A-Za-z0-9-]*(?:\s+[A-Za-z][A-Za-z-]*)*)", label)
    name = name_match.group(1).lower() if name_match else None
    dose_match = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*(nM|uM|µM|mM)", text, re.I)
    duration_match = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*(min|minutes?|h|hr|hours?)", text, re.I)
    dose = float(dose_match.group(1)) if dose_match else None
    unit = dose_match.group(2) if dose_match else None
    duration = None
    if duration_match:
        duration = float(duration_match.group(1)) * (
            1 / 60 if duration_match.group(2).lower().startswith("min") else 1
        )
    return text or None, name, dose, unit, duration


def _context_for(
    source_id, source_asset_id, source_file, cell, experiment, average_annotation=None
):
    label, name, dose, unit, duration = _treatment(experiment)
    study = _STUDIES[source_id]
    annotation = _text(average_annotation)
    key = json.dumps([source_id, cell, experiment, annotation])
    reused = (name or "").startswith(("epoxomicin", "epoximicin")) or (
        source_id == "dasgupta_2014" and name == "bortezomib"
    )
    # Original raw assertions are retained even if the full timing is unresolved.
    timeline = re.search(
        r"([0-9]+(?:\.[0-9]+)?)\s*(min|hr|h|')\s*\+\s*"
        r"(?:([0-9]+(?:\.[0-9]+)?)\s*(min|hr|h|')?\s*)?wash",
        annotation or "",
        re.I,
    )

    def hours(value, unit):
        return float(value) / 60 if unit.lower() in {"min", "'"} else float(value)

    incubation = hours(timeline[1], timeline[2]) if timeline else None
    harvest = hours(timeline[3], timeline[4] or "min") if timeline and timeline[3] else None
    harvest_presence = (
        "present_per_study_protocol"
        if source_id == "dasgupta_2014" and not reused
        else "unresolved"
    )
    return {
        "context_id": _id("context", key),
        "study_id": source_id,
        "source_id": source_id,
        "source_asset_id": source_asset_id,
        "source_file": source_file,
        "publication_doi": study["doi"],
        "publication_url": study["url"],
        "organism_name": "Homo sapiens",
        "organism_taxon_id": "9606",
        "cell_line_name": cell,
        "measured_pool": "native_cellular_peptide_extract",
        "extraction_method": "whole_cell_peptide_extract",
        "treatment_label_raw": label,
        "treatment_name": name,
        "treatment_dose": dose,
        "treatment_dose_unit": unit,
        "exposure_duration_raw": label,
        "duration_hours": duration,
        "duration_role": "reported_total_interval" if duration is not None else "unresolved",
        "average_annotation_raw": annotation,
        "incubation_duration_hours": incubation,
        "harvest_duration_hours": harvest,
        "harvest_compound_status": harvest_presence,
        "timeline_status": "source_annotation_components" if timeline else "unresolved_components",
        "timeline_evidence_url": study["url"] + "#" + study["treatment_section"],
        "labeling_method": "TMAB-NHS isotopic labeling",
        "instrument": None if reused else study["protocol_instrument"],
        "publishing_protocol_instrument": study["protocol_instrument"],
        "instrument_status": "unresolved_reused_acquisition"
        if reused
        else "reported_study_protocol",
        "protocol_evidence_url": study["url"] + "#" + study["protocol_section"],
        "acquisition_mode": "data_dependent_LC_MS/MS",
        "assay_method": "native cellular peptidomics LC-MS/MS",
        "assay_readout": "relative isotope-channel peak-intensity ratio",
        "normalization_policy": "treated_or_control_channel_relative_to_source_control_average",
        "context_scope": "study_and_experiment_context",
        "context_status": "reported_source_scope",
        "wash_interval_raw": annotation,
        "wash_presence_status": harvest_presence,
        "exposure_interval_status": "source_experiment_label_plus_methods_scope",
        "acquisition_source_scope": "original_acquisition_unresolved"
        if reused
        else "reported_study_protocol",
        "source_study_inheritance_status": "unresolved_for_reused_rows"
        if reused
        else "reported_study_protocol",
    }


def _asset_rows():
    return [
        {
            "source_asset_id": spec["asset_id"],
            "source_id": spec["source_id"],
            "source_file": spec["filename"],
            "source_sha256": spec["sha256"],
            "source_asset_url": spec["url"],
            "source_role": spec["role"],
        }
        for spec in _asset_specs()
    ]


def _candidate_fingerprint(cell, sequence, charge, mass, ratios):
    """Candidate reuse only: source-independent identity and exact ratio vector."""

    def is_number(value):
        return (
            isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)
        )

    if not all(is_number(v) for v in (charge, mass)):
        return None
    if sum(is_number(v) for v in ratios) < 2:
        return None
    values = [float(v) if is_number(v) else v for v in ratios]
    payload = json.dumps([cell, sequence, float(charge), float(mass), values], ensure_ascii=False)
    return _id("reuse-candidate", payload)


def _compile(directory):
    directory = Path(directory)
    assets = _asset_rows()
    by_file = {row["source_file"]: row for row in assets}
    source_rows = []
    measurements = []
    contexts = {}
    source_counter = Counter()
    row_candidates = {}
    candidate_sources = {}
    candidate_labels = {}
    for source_file, config in _PRIMARY.items():
        asset = by_file[source_file]
        path = directory / SOURCE_ROOT / source_file
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[config["sheet"]]
        for row_index, values in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_index < config["data_row"] or not any(value is not None for value in values):
                continue
            values = list(values)
            source_record_id = _id(
                "row", f"{asset['source_asset_id']}|{config['sheet']}|{row_index}"
            )
            cell = _text(values[config["cell"] - 1])
            experiment = _text(values[config["experiment"] - 1])
            sequence_raw = _text(values[config["sequence"] - 1])
            identity = _sequence_identity(sequence_raw)
            context = _context_for(
                asset["source_id"],
                asset["source_asset_id"],
                source_file,
                cell,
                experiment,
                values[config["average_annotation"] - 1]
                if len(values) >= config["average_annotation"]
                else None,
            )
            contexts[context["context_id"]] = context
            source_rows.append(
                {
                    "source_record_id": source_record_id,
                    "source_asset_id": asset["source_asset_id"],
                    "source_id": asset["source_id"],
                    "source_file": source_file,
                    "sheet_name": config["sheet"],
                    "source_row_index": row_index,
                    "source_row_locator": f"{source_file}:{config['sheet']}!A{row_index}",
                    "source_row_json": _row_json(values),
                    "source_row_disposition": "primary_ratio_panel",
                    "source_sequence_raw": sequence_raw,
                    "source_handling_raw": _text(values[config["average_annotation"] - 1])
                    if len(values) >= config["average_annotation"]
                    else None,
                    "source_observed_mass_raw": _text(values[config["observed_mass"] - 1]),
                    "source_experiment_raw": experiment,
                    "source_cell_raw": cell,
                    "source_comments_raw": _text(values[config["comments"] - 1]),
                    "row_identity_status": identity["identification"],
                }
            )
            source_counter[source_file] += 1
            sequence = identity["sequence"]
            ratios = [
                values[col - 1] if col <= len(values) else None
                for col, _, _ in config["ratio_columns"][:5]
            ]
            candidate = _candidate_fingerprint(cell, sequence_raw, values[2], values[4], ratios)
            row_candidates[source_record_id] = candidate
            if candidate is not None:
                candidate_sources.setdefault(candidate, set()).add(asset["source_id"])
                candidate_labels.setdefault(candidate, set()).add(
                    " ".join((experiment or "").lower().split())
                )
            shared_control = _id("control", f"{asset['source_id']}|{cell}|{experiment}")
            for column, role, replicate in config["ratio_columns"]:
                raw = values[column - 1] if column <= len(values) else None
                value, inequality, censoring, value_status, overlap = _parse_ratio(raw)
                source_cell = f"{openpyxl.utils.get_column_letter(column)}{row_index}"
                measurements.append(
                    {
                        "measurement_id": _id("measurement", f"{source_record_id}|{source_cell}"),
                        "source_record_id": source_record_id,
                        "context_id": context["context_id"],
                        "lineage_group_id": _id("lineage", f"{source_record_id}|{source_cell}"),
                        "candidate_reuse_key": candidate,
                        "source_asset_id": asset["source_asset_id"],
                        "source_id": asset["source_id"],
                        "source_file": source_file,
                        "source_row_index": row_index,
                        "source_cell": source_cell,
                        "ratio_role": role,
                        "replicate_label": replicate,
                        "channel_label_raw": f"{role}:{replicate}",
                        "peptide_sequence_raw": sequence_raw,
                        "peptide_sequence": sequence,
                        "peptide_modifications_raw": identity["raw_modifications"],
                        "peptide_modifications_json": json.dumps(
                            identity["modifications"], sort_keys=True
                        ),
                        "sequence_identity_status": identity["identification"],
                        "modification_status": identity["modification_status"],
                        "parent_label_raw": _text(values[config["protein"] - 1]),
                        "gene_name": _text(values[config["gene"] - 1]) if config["gene"] else None,
                        "cell_line_name": cell,
                        "subcellular_location_raw": _text(values[config["location"] - 1]),
                        "charge_raw": _text(values[config["charge"] - 1]),
                        "observed_mass_raw": _text(values[config["observed_mass"] - 1]),
                        "theoretical_mass_raw": _text(values[config["theoretical_mass"] - 1]),
                        "measured_pool": "native_cellular_peptide_extract",
                        "quantitative_raw": _text(raw),
                        "quantitative_value": value,
                        "measurement_unit": "relative_ratio",
                        "measurement_inequality": inequality,
                        "censoring_status": censoring,
                        "value_status": value_status,
                        "peak_overlap_status": overlap,
                        "eligibility_status": _eligibility(identity, role, value, censoring),
                        "source_reuse_status": "unresolved_cross_paper_reuse",
                        "acquisition_provenance_status": context["instrument_status"],
                        "shared_control_group_id": shared_control,
                        "shared_control_group_status": "source_label_proxy_unresolved_denominator",
                        "source_comments_raw": _text(values[config["comments"] - 1]),
                    }
                )

    for source_file, config in _HEATMAP.items():
        asset = by_file[source_file]
        workbook = openpyxl.load_workbook(
            directory / SOURCE_ROOT / source_file, read_only=True, data_only=True
        )
        sheet = workbook[config["sheet"]]
        for row_index, values in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_index < config["data_row"] or not any(value is not None for value in values):
                continue
            values = list(values)
            source_record_id = _id(
                "row", f"{asset['source_asset_id']}|{config['sheet']}|{row_index}"
            )
            source_rows.append(
                {
                    "source_record_id": source_record_id,
                    "source_asset_id": asset["source_asset_id"],
                    "source_id": asset["source_id"],
                    "source_file": source_file,
                    "sheet_name": config["sheet"],
                    "source_row_index": row_index,
                    "source_row_locator": f"{source_file}:{config['sheet']}!A{row_index}",
                    "source_row_json": _row_json(values),
                    "source_row_disposition": "deferred_heatmap_republication",
                    "source_sequence_raw": _text(values[config["sequence"] - 1])
                    if len(values) >= config["sequence"]
                    else None,
                    "source_handling_raw": None,
                    "source_observed_mass_raw": None,
                    "source_experiment_raw": None,
                    "source_cell_raw": None,
                    "source_comments_raw": None,
                    "row_identity_status": "heatmap_republication",
                }
            )
            source_counter[source_file] += 1

    lineage = []
    cross_paper_groups = {group for group, sources in candidate_sources.items() if len(sources) > 1}
    for row in measurements:
        candidate = row_candidates[row["source_record_id"]]
        cross_paper = candidate in cross_paper_groups
        if cross_paper:
            row["lineage_group_id"] = _id(
                "lineage", f"{candidate}|{row['ratio_role']}|{row['replicate_label']}"
            )
        group = row["lineage_group_id"]
        row["source_reuse_status"] = (
            "possible_cross_paper_reuse" if cross_paper else "unresolved_cross_paper_reuse"
        )
        lineage.append(
            {
                "lineage_group_id": group,
                "measurement_id": row["measurement_id"],
                "source_record_id": row["source_record_id"],
                "source_id": row["source_id"],
                "cross_paper_match_key": candidate,
                "lineage_status": "possible_cross_paper_reuse"
                if cross_paper
                else "unresolved_source_lineage",
                "reuse_status": "unresolved_cross_paper_reuse",
                "relation_basis": "same_cell_sequence_charge_mass_ratio_vector"
                if cross_paper
                else "source_cell_identity_only",
                "independent_specimen_status": "unknown",
                "candidate_reuse_group_id": candidate if cross_paper else None,
                "experiment_label_agreement": (
                    "equal_after_case_whitespace_normalization"
                    if len(candidate_labels[candidate]) == 1
                    else "different_source_labels"
                )
                if cross_paper
                else "not_applicable",
            }
        )

    coverage = {
        "n_source_assets": len(assets),
        "n_source_rows": len(source_rows),
        "n_primary_ratio_rows": source_counter["dasgupta_s1.xlsx"]
        + source_counter["gelman_s1.xlsx"],
        "n_deferred_heatmap_rows": source_counter["dasgupta_s2.xlsx"]
        + source_counter["gelman_s2.xlsx"],
        "n_measurements": len(measurements),
        "n_contexts": len(contexts),
        "n_lineage_rows": len(lineage),
        "by_value_status": dict(
            sorted(Counter(row["value_status"] for row in measurements).items())
        ),
        "by_censoring_status": dict(
            sorted(Counter(row["censoring_status"] for row in measurements).items())
        ),
        "by_source_file": dict(sorted(source_counter.items())),
        "n_possible_cross_paper_reuse_groups": len(
            {
                row["lineage_group_id"]
                for row in measurements
                if row["source_reuse_status"] == "possible_cross_paper_reuse"
            }
        ),
        "n_candidate_cross_paper_row_groups": len(cross_paper_groups),
        "n_candidate_row_groups_with_equal_experiment_labels": sum(
            len(candidate_labels[group]) == 1 for group in cross_paper_groups
        ),
        "by_identification_status": dict(
            sorted(Counter(row["sequence_identity_status"] for row in measurements).items())
        ),
        "by_eligibility_status": dict(
            sorted(Counter(row["eligibility_status"] for row in measurements).items())
        ),
        "native_rows_preserved": True,
        "source_assets_snapshotted": True,
        "absolute_quantity_claim": False,
        "processing_or_MHC_label_claim": False,
    }
    return (
        {
            "source_assets": pa.Table.from_pylist(assets, schema=SOURCE_SCHEMA),
            "source_rows": pa.Table.from_pylist(source_rows, schema=ROW_SCHEMA),
            "measurements": pa.Table.from_pylist(measurements, schema=MEASUREMENT_SCHEMA),
            "contexts": pa.Table.from_pylist(list(contexts.values()), schema=CONTEXT_SCHEMA),
            "lineage": pa.Table.from_pylist(lineage, schema=LINEAGE_SCHEMA),
        },
        coverage,
    )


def _manifest_files(staging):
    return {
        str(path.relative_to(staging)): file_digest(path)
        for path in sorted(staging.rglob("*"))
        if path.is_file() and path != staging / "manifest.json"
    }


def verify_tmab_peptidomics_package(directory):
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    if (
        manifest.get("schema_version") != PACKAGE_SCHEMA_VERSION
        or manifest.get("catalog_schema_version") != CATALOG_SCHEMA_VERSION
        or manifest.get("revision") != REVISION
    ):
        raise ValueError("Unsupported TMAB native peptidomics package")
    actual = {
        str(path.relative_to(directory))
        for path in directory.rglob("*")
        if path.is_file() and path != directory / "manifest.json"
    }
    if set(manifest.get("files", {})) != actual:
        raise ValueError("TMAB package inventory mismatch")
    for name, digest in manifest["files"].items():
        if file_digest(directory / name)["sha256"] != digest["sha256"]:
            raise ValueError(f"Changed TMAB artifact: {name}")
    for spec in _asset_specs():
        asset_path = directory / SOURCE_ROOT / spec["filename"]
        if _digest(asset_path) != spec["sha256"]:
            raise ValueError(f"Changed pinned TMAB source asset: {spec['filename']}")
    expected, coverage = _compile(directory)
    paths = {
        "source_assets": SOURCE_TABLE,
        "source_rows": ROWS_TABLE,
        "measurements": MEASUREMENTS_TABLE,
        "contexts": CONTEXT_TABLE,
        "lineage": LINEAGE_TABLE,
    }
    for name, table in expected.items():
        actual_table = pq.read_table(directory / paths[name])
        if not actual_table.equals(table, check_metadata=False):
            raise ValueError(f"TMAB relation differs from source replay: {name}")
    actual_coverage = json.loads((directory / COVERAGE).read_text())
    if actual_coverage != coverage or manifest.get("reconciliation") != coverage:
        raise ValueError("TMAB coverage changed")
    return manifest


def write_tmab_peptidomics_package(source_directory, destination):
    """Build a standalone, hash-pinned native peptide measurement package."""
    source_directory = Path(source_directory)
    destination = Path(destination)
    if destination.exists():
        raise FileExistsError(destination)
    staging = Path(tempfile.mkdtemp(prefix=f".{destination.name}-", dir=destination.parent))
    try:
        (staging / SOURCE_ROOT).mkdir(parents=True, exist_ok=True)
        for spec in _asset_specs():
            source = _source_path(source_directory, spec["filename"])
            if _digest(source) != spec["sha256"]:
                raise ValueError(f"TMAB source hash changed: {spec['filename']}")
            snapshot_file(source, staging / SOURCE_ROOT / spec["filename"])
        tables, coverage = _compile(staging)
        for name, table in tables.items():
            path = {
                "source_assets": SOURCE_TABLE,
                "source_rows": ROWS_TABLE,
                "measurements": MEASUREMENTS_TABLE,
                "contexts": CONTEXT_TABLE,
                "lineage": LINEAGE_TABLE,
            }[name]
            pq.write_table(table, staging / path, compression="zstd")
        (staging / COVERAGE).write_text(json.dumps(coverage, indent=2, sort_keys=True) + "\n")
        from .version import __version__

        manifest = {
            "schema_version": PACKAGE_SCHEMA_VERSION,
            "catalog_schema_version": CATALOG_SCHEMA_VERSION,
            "hitlist_version": __version__,
            "revision": REVISION,
            "package_family": "tmab_peptidomics",
            "source_registry": {
                "dasgupta_2014": {
                    **_STUDIES["dasgupta_2014"],
                    "source_role": "native_cellular_peptidomics",
                },
                "gelman_2013": {
                    **_STUDIES["gelman_2013"],
                    "source_role": "native_cellular_peptidomics",
                },
            },
            "reconciliation": coverage,
            "limitations": [
                "native_cellular_extract_not_MHC_elution",
                "relative_ratio_not_absolute_concentration",
                "treatment_is_not_genetic_perturbation",
                "cross_paper_reuse_and_specimen_independence_unresolved",
                "heatmap_republications_deferred_from_training",
            ],
            "files": _manifest_files(staging),
        }
        (staging / "manifest.json").write_text(
            json.dumps(manifest, indent=2, sort_keys=True) + "\n"
        )
        verify_tmab_peptidomics_package(staging)
        staging.rename(destination)
        return destination / "manifest.json"
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
