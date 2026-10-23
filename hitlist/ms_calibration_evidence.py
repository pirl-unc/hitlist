"""Source-faithful known-load MS response evidence; no concentration admission.

Nominal standards, reporter forms, source column slots and reported quantities
have distinct identities. A source blank is never a quantitative zero. Curves
retain their actual plotting grids without inventing a sample-to-load join.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import statistics
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter, defaultdict
from itertools import pairwise
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq

SOURCE_ID = "pxd004886"
DATA = Path(__file__).parent / "data/evidence_revisions"
ASSET_RECEIPT = DATA / "ms_calibration_assets.json"
SUPPORT_FILES = {"contract.json": DATA / "ms_calibration_contract.json"}
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _schema(text=(), integers=(), numbers=(), flags=()):
    return pa.schema(
        [(key, pa.string()) for key in text]
        + [(key, pa.int64()) for key in integers]
        + [(key, pa.float64()) for key in numbers]
        + [(key, pa.bool_()) for key in flags]
    )


SCHEMAS = {
    "protocol_parameters": _schema(
        (
            "parameter_id",
            "site_label",
            "protocol_stage",
            "scope",
            "parameter_name_raw",
            "value_json",
            "value_status",
            "unit_raw",
            "source_cell",
            "source_locator",
            "parameter_source_cell",
            "stage_source_cell",
            "position_source_cell",
            "site_source_cell",
            "position_unit_raw",
        ),
        numbers=("value", "position_value"),
    ),
    "source_sheets": _schema(
        (
            "sheet_id",
            "asset_path",
            "sheet_name",
            "sheet_state",
            "merged_ranges_json",
            "publication_label",
            "workbook_title",
            "title_status",
        ),
        ("n_rows", "n_columns"),
    ),
    "source_cells": _schema(
        (
            "cell_id",
            "sheet_id",
            "asset_path",
            "sheet_name",
            "cell_ref",
            "excel_type",
            "value_json",
            "formula",
            "xml_element",
            "value_status",
        ),
        ("row_1based", "column_1based"),
        ("value",),
    ),
    "nominal_standards": _schema(
        (
            "standard_id",
            "peptide_sequence",
            "form_label_raw",
            "protein_label_raw",
            "group_label",
            "form_source_cell",
            "protein_source_cell",
            "group_source_cell",
            "selection_status",
            "lot_status",
            "purity_status",
        ),
    ),
    "nominal_loads": _schema(
        (
            "nominal_load_id",
            "group_label",
            "sample_label",
            "unit",
            "unit_label_raw",
            "source_cell",
            "sample_source_cell",
            "group_source_cell",
            "unit_source_cell",
            "assignment_status",
        ),
        numbers=("value",),
    ),
    "reporter_forms": _schema(
        (
            "reporter_id",
            "standard_id",
            "peptide_sequence",
            "label_raw",
            "annotations_json",
            "aggregation_label_raw",
            "membership_status",
            "form_amount_status",
            "source_cells_json",
        ),
    ),
    "source_slots": _schema(
        (
            "slot_id",
            "source_column",
            "site_label",
            "sample_label",
            "site_source_cell",
            "sample_source_cell",
            "run_identity_status",
            "day_status",
            "specimen_status",
            "material_link_precision",
        ),
        ("source_slot_ordinal",),
    ),
    "s4_measurements": _schema(
        (
            "measurement_id",
            "reporter_id",
            "slot_id",
            "nominal_load_id",
            "representation",
            "unit",
            "quantity_scale",
            "value_status",
            "source_cell",
            "source_locator",
            "reported_load_source_cell",
            "calibration_applicability",
        ),
        numbers=("value", "reported_load"),
    ),
    "summary_measurements": _schema(
        (
            "measurement_id",
            "reporter_id",
            "representation",
            "statistic_name",
            "unit",
            "value_status",
            "source_cell",
            "derivation_status",
        ),
        numbers=("value", "recomputed_value", "relative_replay_error"),
    ),
    "curve_rows": _schema(
        (
            "curve_id",
            "reporter_id",
            "site_label",
            "representation",
            "group_label",
            "label_source_cell",
            "group_source_cell",
            "rsq_source_cell",
            "rsq_formula",
            "header_range_raw",
            "response_range_raw",
            "sample_join_status",
        ),
        numbers=("reported_r_squared",),
    ),
    "curve_measurements": _schema(
        (
            "measurement_id",
            "curve_id",
            "reporter_id",
            "representation",
            "unit",
            "value_status",
            "source_cell",
            "source_locator",
            "grid_amount_source_cell",
            "grid_unit_status",
            "sample_join_status",
            "censoring_status",
        ),
        ("source_grid_ordinal",),
        ("value", "reported_grid_amount"),
    ),
    "source_assertions": _schema(
        (
            "assertion_id",
            "subject_id",
            "predicate",
            "value_raw",
            "normalized_value",
            "assertion_status",
            "source_locator",
            "scope",
        ),
    ),
    "normalization_diagnostics": _schema(
        ("diagnostic_id", "slot_id", "diagnostic_status", "method", "source_cells_json"),
        ("n_paired_reporters", "n_normalized_only", "n_raw_only"),
        ("median_ratio", "max_relative_deviation"),
    ),
}


def _json(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False, allow_nan=False)


def _id(kind, value):
    return f"{SOURCE_ID}:{kind}:{hashlib.sha256(str(value).encode()).hexdigest()[:24]}"


def _numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _value_status(value, *, present, formula=False):
    if not present:
        return "absent_source_cell"
    if value == "":
        return "explicit_empty_string"
    if value is None:
        return "empty_formula_cache" if formula else "empty_source_value"
    if _numeric(value):
        return (
            "reported_positive"
            if value > 0
            else "reported_zero"
            if value == 0
            else "reported_negative"
        )
    return "reported_nonnumeric"


def canonical_sequence(label, *, panel=False):
    """A sequence-level panel join, not molecular-form equivalence."""
    value = label if panel else label.split(".", 1)[1]
    value = re.sub(r"\.sum[0-9]+$", "", value)
    value = re.sub(r"\[[^]]+\]", "", value)
    if not re.fullmatch("[ACDEFGHIKLMNPQRSTVWY]+", value):
        raise ValueError(f"Invalid calibration panel sequence: {label}")
    return value


def _mode(label):
    if not isinstance(label, str):
        return None
    normalized = re.sub(r"[-\s]", "", label.lower())
    return {"unnormalized": "raw", "normalized": "normalized"}.get(normalized)


def _column(number):
    label = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        label = chr(65 + remainder) + label
    return label


def _coordinate(value):
    match = re.fullmatch(r"\$?([A-Z]+)\$?([0-9]+)", value)
    if not match:
        raise ValueError(f"Unsupported calibration coordinate: {value}")
    column = 0
    for char in match[1]:
        column = column * 26 + ord(char) - 64
    return column, int(match[2])


def curve_range(formula):
    """Read actual RSQ source arguments, including expanded shared formulas."""
    matches = re.findall(
        r"RSQ\((\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+),(\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+)\)", formula or ""
    )
    if not matches or len(set(matches)) != 1:
        raise ValueError("Unresolved calibration RSQ range")
    header, responses = matches[0]
    left, right = [_coordinate(x) for x in header.split(":")]
    low, high = [_coordinate(x) for x in responses.split(":")]
    if (
        left[1] != right[1]
        or low[1] != high[1]
        or (left[0], right[0]) != (low[0], high[0])
        or right[0] - left[0] != 4
    ):
        raise ValueError("Calibration RSQ range is not five corresponding positions")
    return (
        header,
        responses,
        [(f"{_column(c)}{left[1]}", f"{_column(c)}{low[1]}") for c in range(left[0], right[0] + 1)],
    )


def _workbooks(source_dir, tables):
    import openpyxl

    books = {}
    for number in (1, 2, 4, 8, 13, 14, 15):
        asset = f"supplementary-data-{number}.xlsx"
        cached = openpyxl.load_workbook(source_dir / "assets" / asset, data_only=True)
        formulas = openpyxl.load_workbook(source_dir / "assets" / asset, data_only=False)
        try:
            if len(cached.worksheets) != 1:
                raise ValueError("Unexpected calibration workbook sheets")
            ws, wf = cached.active, formulas.active
            sheet_id = f"{SOURCE_ID}:{asset}:{ws.title}"
            with zipfile.ZipFile(source_dir / "assets" / asset) as archive:
                workbook = ET.fromstring(archive.read("xl/workbook.xml"))
                relationships = {
                    r.attrib["Id"]: r.attrib["Target"]
                    for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
                }
                sheet = workbook.find("s:sheets/s:sheet", NS)
                target = relationships[
                    sheet.attrib[
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    ]
                ]
                target = target.lstrip("/") if target.startswith("/") else "xl/" + target
                nodes = ET.fromstring(archive.read(target)).findall(".//s:sheetData/s:row/s:c", NS)
            cells = {}
            for node in nodes:
                if all(node.find(f"s:{tag}", NS) is None for tag in ("v", "f", "is")):
                    continue
                coordinate = node.attrib["r"]
                cached_cell, formula_cell = ws[coordinate], wf[coordinate]
                formula = formula_cell.value if formula_cell.data_type == "f" else None
                value = cached_cell.value
                record = {
                    "cell_id": f"{sheet_id}:{coordinate}",
                    "sheet_id": sheet_id,
                    "asset_path": asset,
                    "sheet_name": ws.title,
                    "cell_ref": coordinate,
                    "excel_type": formula_cell.data_type,
                    "value_json": _json(value),
                    "formula": formula,
                    "xml_element": ET.canonicalize(
                        ET.tostring(node, encoding="unicode"), rewrite_prefixes=True
                    ),
                    "value_status": _value_status(value, present=True, formula=formula is not None),
                    "row_1based": cached_cell.row,
                    "column_1based": cached_cell.column,
                    "value": float(value) if _numeric(value) else None,
                }
                tables["source_cells"].append(record)
                cells[coordinate] = record
            title = ws["A1"].value
            tables["source_sheets"].append(
                {
                    "sheet_id": sheet_id,
                    "asset_path": asset,
                    "sheet_name": ws.title,
                    "sheet_state": ws.sheet_state,
                    "merged_ranges_json": _json(sorted(map(str, ws.merged_cells.ranges))),
                    "publication_label": f"Supplementary Data {number}",
                    "workbook_title": title,
                    "title_status": "source_label_conflict" if number == 8 else "consistent",
                    "n_rows": ws.max_row,
                    "n_columns": ws.max_column,
                }
            )
            books[number] = cells
        finally:
            cached.close()
            formulas.close()
    return books


class _Book:
    def __init__(self, cells):
        self.cells = cells

    def value(self, coordinate):
        cell = self.cells.get(coordinate)
        return json.loads(cell["value_json"]) if cell else None

    def source(self, coordinate):
        return self.cells.get(coordinate, {}).get("cell_id")

    def status(self, coordinate):
        return self.cells.get(coordinate, {}).get("value_status", "absent_source_cell")


def _write_protocol_parameters(books, tables):
    def add(
        number,
        coordinate,
        parameter,
        *,
        site=None,
        stage=None,
        scope,
        parameter_cell=None,
        stage_cell=None,
        position_cell=None,
        site_cell=None,
        unit=None,
        position_unit=None,
    ):
        book = books[number]
        value = book.value(coordinate)
        position = book.value(position_cell) if position_cell else None
        tables["protocol_parameters"].append(
            {
                "parameter_id": _id("protocol", f"{number}:{coordinate}:{scope}"),
                "site_label": str(site) if site is not None else None,
                "protocol_stage": stage,
                "scope": scope,
                "parameter_name_raw": parameter,
                "value_json": _json(value),
                "value_status": book.status(coordinate),
                "unit_raw": unit,
                "source_cell": book.source(coordinate),
                "source_locator": f"supplementary-data-{number}.xlsx#{coordinate}",
                "parameter_source_cell": book.source(parameter_cell) if parameter_cell else None,
                "stage_source_cell": book.source(stage_cell) if stage_cell else None,
                "position_source_cell": book.source(position_cell) if position_cell else None,
                "site_source_cell": book.source(site_cell) if site_cell else None,
                "position_unit_raw": position_unit,
                "value": float(value) if _numeric(value) else None,
                "position_value": float(position) if _numeric(position) else None,
            }
        )

    book = books[13]
    for row in (6, 7, 8):
        for col in range(2, 13):
            column = _column(col)
            site = book.value(f"{column}4")
            if site != col - 1:
                raise ValueError("Site configuration identity changed")
            add(
                13,
                f"{column}{row}",
                book.value(f"A{row}"),
                site=int(site),
                scope="site_LC_MS_configuration",
                parameter_cell=f"A{row}",
                site_cell=f"{column}4",
            )
    for row in (10, 11):
        add(13, f"A{row}", "source_column_footnote", scope="reported_column_footnote")
    book = books[14]
    stage, stage_cell = None, None
    for row in range(6, 24):
        value = book.value(f"A{row}")
        if isinstance(value, str):
            stage, stage_cell = value, f"A{row}"
        elif _numeric(value):
            if not stage:
                raise ValueError("Gradient stage is unresolved")
            for col in range(2, 13):
                column = _column(col)
                site = book.value(f"{column}4")
                if site != col - 1:
                    raise ValueError("Gradient site identity changed")
                add(
                    14,
                    f"{column}{row}",
                    book.value("B5"),
                    site=int(site),
                    stage=stage,
                    scope="source_stage_and_site_gradient",
                    parameter_cell="B5",
                    stage_cell=stage_cell,
                    position_cell=f"A{row}",
                    site_cell=f"{column}4",
                    unit="percent_mobile_phase_B",
                    position_unit=book.value("A5"),
                )
    book = books[15]
    for row in range(4, 68):
        if book.value(f"A{row}") != row - 3:
            raise ValueError("SWATH window identity changed")
        for col in "BCDFG":
            parameter = book.value(f"{col}3")
            unit = re.search(r"\(([^)]+)\)", parameter)
            add(
                15,
                f"{col}{row}",
                parameter,
                scope="common_SWATH_window_protocol",
                parameter_cell=f"{col}3",
                position_cell=f"A{row}",
                unit=unit[1] if unit else None,
                position_unit=book.value("A3"),
            )


def _build_tables(source_dir):
    tables = {name: [] for name in SCHEMAS}
    books = {number: _Book(cells) for number, cells in _workbooks(source_dir, tables).items()}
    _write_protocol_parameters(books, tables)
    panel, loads, reporters = {}, {}, {}
    book = books[2]
    for row in range(5, 35):
        label = book.value(f"A{row}")
        sequence = canonical_sequence(label, panel=True)
        if sequence in panel:
            raise ValueError("Duplicate nominal standard sequence")
        record = {
            "standard_id": _id("standard", sequence),
            "peptide_sequence": sequence,
            "form_label_raw": label,
            "protein_label_raw": book.value(f"B{row}"),
            "group_label": book.value(f"C{row}"),
            "form_source_cell": book.source(f"A{row}"),
            "protein_source_cell": book.source(f"B{row}"),
            "group_source_cell": book.source(f"C{row}"),
            "selection_status": "LC_MS_selected_and_response_informed_group",
            "lot_status": "unreported",
            "purity_status": "unreported",
        }
        panel[sequence] = record
        tables["nominal_standards"].append(record)
    if Counter(x["group_label"] for x in panel.values()) != dict.fromkeys("ABCDE", 6):
        raise ValueError("Nominal standard groups changed")
    book = books[1]
    for row in range(6, 11):
        group = book.value(f"A{row}")
        for column in "BCDEF":
            sample, value = book.value(f"{column}5"), book.value(f"{column}{row}")
            if not _numeric(value) or value <= 0:
                raise ValueError("Invalid nominal loaded amount")
            record = {
                "nominal_load_id": _id("load", f"{group}:{sample}"),
                "group_label": group,
                "sample_label": sample,
                "unit": "fmol_on_column",
                "unit_label_raw": book.value("B4"),
                "value": value,
                "source_cell": book.source(f"{column}{row}"),
                "sample_source_cell": book.source(f"{column}5"),
                "group_source_cell": book.source(f"A{row}"),
                "unit_source_cell": book.source("B4"),
                "assignment_status": "nominal_standard_load_not_certified_form_molarity",
            }
            loads[group, sample] = record
            tables["nominal_loads"].append(record)
        amounts = [loads[group, f"S{i}"]["value"] for i in range(1, 6)]
        if any(not math.isclose(b / a, 3, rel_tol=1e-12) for a, b in pairwise(amounts)):
            raise ValueError("Nominal threefold dilution changed")

    def reporter(label, source_cell):
        if label not in reporters:
            sequence = canonical_sequence(label)
            standard = panel[sequence]
            suffix = re.search(r"\.(sum[0-9]+)$", label)
            reporters[label] = {
                "reporter_id": _id("reporter", label),
                "standard_id": standard["standard_id"],
                "peptide_sequence": sequence,
                "label_raw": label,
                "annotations_json": _json(re.findall(r"\[([^]]+)\]", label)),
                "aggregation_label_raw": suffix[1] if suffix else None,
                "membership_status": "source_panel_unique_sequence_match_not_form_equivalence",
                "form_amount_status": "nominal_standard_only_form_fraction_unreported",
                "source_cells_json": [],
            }
        reporters[label]["source_cells_json"].append(source_cell)
        return reporters[label]

    book = books[4]
    for column in range(6, 105):
        site = (column - 6) // 9 + 1
        header = f"{_column(6 + (site - 1) * 9)}3"
        if book.value(header) != site:
            raise ValueError("Calibration site header changed")
        tables["source_slots"].append(
            {
                "slot_id": _id("slot", _column(column)),
                "source_column": _column(column),
                "site_label": str(site),
                "sample_label": "S4",
                "site_source_cell": book.source(header),
                "sample_source_cell": book.source("A1"),
                "source_slot_ordinal": (column - 6) % 9 + 1,
                "run_identity_status": "unreported",
                "day_status": "unreported",
                "specimen_status": "unreported",
                "material_link_precision": "source_site_and_column_slot_only",
            }
        )
    mode = None
    representations = defaultdict(set)
    for row in range(1, 75):
        label = book.value(f"A{row}")
        mode = _mode(label) or mode
        if not isinstance(label, str) or "." not in label or mode is None:
            continue
        rep = reporter(label, book.source(f"A{row}"))
        representations[mode].add(label)
        nominal = loads[panel[rep["peptide_sequence"]]["group_label"], "S4"]
        values = []
        for slot in tables["source_slots"]:
            coordinate = f"{slot['source_column']}{row}"
            value = book.value(coordinate)
            if _numeric(value):
                values.append(value)
            tables["s4_measurements"].append(
                {
                    "measurement_id": _id("S4", f"{label}:{slot['slot_id']}:{mode}"),
                    "reporter_id": rep["reporter_id"],
                    "slot_id": slot["slot_id"],
                    "nominal_load_id": nominal["nominal_load_id"],
                    "representation": mode,
                    "unit": "arbitrary_peak_area",
                    "quantity_scale": "linear_fragment_peak_area",
                    "value_status": book.status(coordinate),
                    "source_cell": book.source(coordinate),
                    "source_locator": f"supplementary-data-4.xlsx#{coordinate}",
                    "value": float(value) if _numeric(value) else None,
                    "reported_load_source_cell": book.source(f"B{row}"),
                    "reported_load": book.value(f"B{row}"),
                    "calibration_applicability": "nominal_S4_standard_load_only_not_molar_form_concentration",
                }
            )
        derived = (
            [
                statistics.mean(values),
                statistics.stdev(values),
                100 * statistics.stdev(values) / statistics.mean(values),
            ]
            if len(values) > 1
            else [None] * 3
        )
        for column, statistic, recomputed in zip(
            "CDE",
            ("mean", "sample_standard_deviation", "coefficient_of_variation_percent"),
            derived,
        ):
            coordinate = f"{column}{row}"
            value = book.value(coordinate)
            error = (
                abs(recomputed - value) / max(abs(value), 1e-15)
                if _numeric(value) and recomputed is not None
                else None
            )
            tables["summary_measurements"].append(
                {
                    "measurement_id": _id("summary", coordinate),
                    "reporter_id": rep["reporter_id"],
                    "representation": mode,
                    "statistic_name": statistic,
                    "unit": "percent" if column == "E" else "arbitrary_peak_area",
                    "value_status": book.status(coordinate),
                    "source_cell": book.source(coordinate),
                    "value": float(value) if _numeric(value) else None,
                    "recomputed_value": recomputed,
                    "relative_replay_error": error,
                    "derivation_status": "source_matrix_reproduced"
                    if error is not None and error < 1e-10
                    else "not_computable"
                    if error is None
                    else "source_conflict",
                }
            )
    if representations["raw"] != representations["normalized"] or len(representations["raw"]) != 31:
        raise ValueError("S4 reporter universe changed")

    book = books[8]
    site, mode = None, None
    curve_labels = defaultdict(set)
    for row in range(1, 872):
        label, group = book.value(f"A{row}"), book.value(f"B{row}")
        if isinstance(label, str) and re.fullmatch(r"Site [0-9]+", label.strip()):
            site, mode = label.strip().split()[-1], None
        mode = _mode(label) or mode
        if not isinstance(label, str) or "." not in label or group not in set("ABCDE"):
            continue
        if not site or not mode or label not in representations["raw"]:
            raise ValueError("Unresolved curve source context or reporter")
        rep = reporter(label, book.source(f"A{row}"))
        if group != panel[rep["peptide_sequence"]]["group_label"]:
            raise ValueError("Curve nominal panel group disagrees")
        curve_labels[site, mode].add(label)
        formula = book.cells[f"P{row}"]["formula"]
        header, response_range, pairs = curve_range(formula)
        if any(_coordinate(c)[1] != row for _, c in pairs):
            raise ValueError("Curve formula refers to another reporter")
        curve_id = _id("curve", f"{site}:{mode}:{label}")
        r_squared = book.value(f"P{row}")
        tables["curve_rows"].append(
            {
                "curve_id": curve_id,
                "reporter_id": rep["reporter_id"],
                "site_label": site,
                "representation": mode,
                "group_label": group,
                "label_source_cell": book.source(f"A{row}"),
                "group_source_cell": book.source(f"B{row}"),
                "rsq_source_cell": book.source(f"P{row}"),
                "rsq_formula": formula,
                "header_range_raw": header,
                "response_range_raw": response_range,
                "sample_join_status": "unresolved_plot_grid_to_nominal_sample",
                "reported_r_squared": float(r_squared) if _numeric(r_squared) else None,
            }
        )
        for ordinal, (grid, coordinate) in enumerate(pairs, 1):
            value, amount = book.value(coordinate), book.value(grid)
            if not _numeric(amount) or amount <= 0:
                raise ValueError("Invalid source curve grid amount")
            tables["curve_measurements"].append(
                {
                    "measurement_id": _id("curve_point", f"{curve_id}:{ordinal}"),
                    "curve_id": curve_id,
                    "reporter_id": rep["reporter_id"],
                    "representation": mode,
                    "unit": "arbitrary_peak_area",
                    "value_status": book.status(coordinate),
                    "source_cell": book.source(coordinate),
                    "source_locator": f"supplementary-data-8.xlsx#{coordinate}",
                    "grid_amount_source_cell": book.source(grid),
                    "grid_unit_status": "fmol_plot_axis_not_reconciled_to_nominal_design",
                    "sample_join_status": "unresolved_plot_grid_to_nominal_sample",
                    "censoring_status": "unresolved_omission_or_quantification_filter",
                    "source_grid_ordinal": ordinal,
                    "value": float(value) if _numeric(value) else None,
                    "reported_grid_amount": amount,
                }
            )
    if len(curve_labels) != 22 or any(
        labels != representations["raw"] for labels in curve_labels.values()
    ):
        raise ValueError("Curve site/representation universe changed")
    for record in reporters.values():
        record["source_cells_json"] = _json(sorted(set(record["source_cells_json"])))
        tables["reporter_forms"].append(record)

    measurements = {
        (x["slot_id"], x["reporter_id"], x["representation"]): x for x in tables["s4_measurements"]
    }
    for slot in tables["source_slots"]:
        ratios, only_normalized, only_raw, cells = [], 0, 0, []
        for rep in reporters.values():
            raw, norm = [
                measurements[slot["slot_id"], rep["reporter_id"], mode]
                for mode in ("raw", "normalized")
            ]
            positive = [x["value"] is not None and x["value"] > 0 for x in (raw, norm)]
            if all(positive):
                ratios.append(norm["value"] / raw["value"])
            elif positive[1]:
                only_normalized += 1
            elif positive[0]:
                only_raw += 1
            cells.extend(x["source_cell"] for x in (raw, norm) if x["source_cell"])
        median = statistics.median(ratios) if ratios else None
        deviation = max(abs(x / median - 1) for x in ratios) if median else None
        tables["normalization_diagnostics"].append(
            {
                "diagnostic_id": _id("normalization", slot["slot_id"]),
                "slot_id": slot["slot_id"],
                "diagnostic_status": "unresolved_source_discrepancy"
                if only_normalized or only_raw or (deviation is not None and deviation > 1e-6)
                else "no_discrepancy_in_available_pairs",
                "method": "within_column_reporter_ratio_diagnostic_not_source_normalization_coefficient",
                "source_cells_json": _json(sorted(set(cells))),
                "n_paired_reporters": len(ratios),
                "n_normalized_only": only_normalized,
                "n_raw_only": only_raw,
                "median_ratio": median,
                "max_relative_deviation": deviation,
            }
        )
    contract = json.loads((source_dir / "contracts/contract.json").read_text())
    for index, assertion in enumerate(contract["assertions"]):
        tables["source_assertions"].append(
            {"assertion_id": f"{SOURCE_ID}:curated:{index}", **assertion}
        )
    return tables


def _reconciliation(tables):
    counts = {f"n_{name}_rows": len(rows) for name, rows in tables.items()}
    for mode in ("raw", "normalized"):
        statuses = Counter(
            x["value_status"] for x in tables["s4_measurements"] if x["representation"] == mode
        )
        counts[f"S4_{mode}_statuses"] = dict(sorted(statuses.items()))
    counts["n_normalized_only_values"] = sum(
        x["n_normalized_only"] for x in tables["normalization_diagnostics"]
    )
    counts["n_nonuniform_normalization_columns"] = sum(
        x["max_relative_deviation"] is not None and x["max_relative_deviation"] > 1e-6
        for x in tables["normalization_diagnostics"]
    )
    counts["n_reproduced_summary_values"] = sum(
        x["derivation_status"] == "source_matrix_reproduced" for x in tables["summary_measurements"]
    )
    counts["n_positive_curve_values"] = sum(
        x["value_status"] == "reported_positive" for x in tables["curve_measurements"]
    )
    return counts


class MSCalibrationAdapter:
    source_id = SOURCE_ID
    schema_version = 1
    asset_receipt_path = ASSET_RECEIPT
    support_files = SUPPORT_FILES

    def write(self, source_dir):
        source_dir = Path(source_dir)
        tables = _build_tables(source_dir)
        typed = source_dir / "typed"
        typed.mkdir()
        for name, schema in SCHEMAS.items():
            pq.write_table(
                pa.Table.from_pylist(tables[name], schema=schema),
                typed / f"{name}.parquet",
                compression="zstd",
            )
        return _reconciliation(tables)

    def verify(self, source_dir, reconciliation):
        source_dir = Path(source_dir)
        contract = json.loads((source_dir / "contracts/contract.json").read_text())
        expected = _build_tables(source_dir)
        if reconciliation != _reconciliation(expected):
            raise ValueError("Calibration source reconciliation mismatch")
        for name, schema in SCHEMAS.items():
            table = pq.read_table(source_dir / "typed" / f"{name}.parquet")
            if (
                table.schema != schema
                or table.to_pylist()
                != pa.Table.from_pylist(expected[name], schema=schema).to_pylist()
            ):
                raise ValueError(f"Calibration typed relation differs from source: {name}")
        for key, value in contract["expected_reconciliation"].items():
            if reconciliation.get(key) != value:
                raise ValueError(f"Pinned calibration source invariant changed: {key}")

    def registry(self, reconciliation):
        return {
            "source_id": SOURCE_ID,
            "dataset_accession": "PXD004886",
            "measured_pools": ["synthetic_standard_in_background_digest"],
            "organism_scope": "synthetic_standards_human_HEK293_digest_background",
            "measurement_grain": "reporter_form_by_source_slot_or_curve_position_and_representation",
            "training_admission": "deferred_to_explicit_nominal_response_recipe",
            "quantity_artifact": "typed/s4_measurements.parquet",
            "curve_artifact": "typed/curve_measurements.parquet",
            "n_nominal_standards": reconciliation["n_nominal_standards_rows"],
            "n_reporter_forms": reconciliation["n_reporter_forms_rows"],
            "rights": {
                "article_and_supplements": "CC-BY4.0_with_third_party_credit_exceptions",
                "source_assets": "caller_supplied_not_bundled_in_distribution",
            },
            "limitations": [
                "nominal_load_not_certified_form_molarity",
                "response_selected_panel_and_groups",
                "no_APC_processing_or_intracellular_concentration_label",
                "unknown_run_day_and_specimen_links",
                "curve_grid_to_sample_unresolved",
                "normalization_source_discrepancies",
                "censoring_and_saturation_mapping_unresolved",
            ],
        }
