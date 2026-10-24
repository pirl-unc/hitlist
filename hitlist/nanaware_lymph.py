"""Original lymph and DC peptide quantities with unresolved context kept explicit."""

from __future__ import annotations

import gzip
import hashlib
import json
import math
import re
import shutil
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pyarrow.parquet as pq

from .evidence_snapshot import snapshot_file
from .mapp_evidence import _schema, _Sink
from .provenance import file_digest

PACKAGE_SCHEMA_VERSION = 24
CATALOG_SCHEMA_VERSION = 29
PACKAGE_FAMILY = "nanaware_lymph"
REVISION = "nanaware2024_lymph_quantities_v1"
ROOT = Path("nanaware_lymph")
CONTRACT = Path(__file__).parent / "data/evidence_revisions/nanaware_lymph.json"
SCHEMAS = {
    "assets": _schema(
        ("source_asset_id", "source_file", "source_asset_url", "source_asset_role", "asset_sha256"),
        integers=("size_bytes",),
    ),
    "source_records": _schema(
        ("source_record_id", "source_file", "source_locator", "source_row_json", "block_id"),
        integers=("source_row_index",),
    ),
    "contexts": _schema(
        (
            "context_id",
            "organism_name",
            "organism_taxon_id",
            "tissue_name",
            "cell_type",
            "measured_pool",
            "mhc_class",
            "mhc_genotype",
            "mhc_capture",
            "digestion_enzyme",
            "instrument",
            "acquisition_mode",
            "normalization_policy",
            "condition_text",
            "component_state_status",
            "condition_knockout_genes",
            "context_scope",
            "source_reference_iri",
            "source_row_json",
        )
    ),
    "samples": _schema(
        (
            "sample_id",
            "context_id",
            "sample_label_raw",
            "biological_replicate_raw",
            "specimen_id",
            "aliquot_id",
            "material_link_precision",
            "replicate_link_status",
        )
    ),
    "observations": _schema(
        (
            "observation_id",
            "source_record_id",
            "block_id",
            "peptide_sequence",
            "peptide_raw",
            "peptide_modifications_raw",
            "sequence_parse_status",
            "length_status",
            "flank_n_raw",
            "flank_c_raw",
            "parent_label_raw",
            "start_raw",
            "end_raw",
            "charge_raw",
            "scan_raw",
            "confidence_json",
        )
    ),
    "quantities": _schema(
        (
            "measurement_id",
            "observation_id",
            "source_record_id",
            "sample_id",
            "context_id",
            "peptide_sequence",
            "parent_label_raw",
            "source_field",
            "source_cell",
            "quantity_kind",
            "quantitative_raw",
            "value_status",
            "measurement_unit",
            "quantity_scale",
            "eligibility_status",
        ),
        numbers=("quantitative_value",),
    ),
    "references": _schema(
        (
            "reference_record_id",
            "source_file",
            "source_locator",
            "accession_raw",
            "reference_sequence",
            "sequence_sha256",
            "reference_role",
            "search_database_ref",
            "source_row_json",
        )
    ),
    "identifications": _schema(
        (
            "identification_id",
            "source_file",
            "source_locator",
            "reference_record_id",
            "peptide_sequence",
            "peptide_modifications_raw",
            "source_peptide_ref",
            "start_raw",
            "end_raw",
            "is_decoy_raw",
            "source_row_json",
        )
    ),
    "assertions": _schema(
        (
            "source_record_id",
            "source_scope",
            "assertion_predicate",
            "assertion_value",
            "assertion_status",
            "source_reference_iri",
            "source_locator",
        )
    ),
}


def _json(value):
    return json.dumps(value, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _id(kind, *values):
    return "nanaware2024:" + kind + ":" + hashlib.sha256(_json(values).encode()).hexdigest()


def _raw(value):
    return "" if value is None else str(value)


def _number(value):
    if value is None or value == "":
        return None, "missing"
    if isinstance(value, str) and value.startswith("="):
        return None, "formula_not_evaluated"
    if isinstance(value, bool):
        return None, "invalid_literal"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None, "invalid_literal"
    if not math.isfinite(number):
        return None, "nonfinite"
    return number, "positive" if number > 0 else "zero" if number == 0 else "negative"


def _peptide(raw):
    """Recognize only the source's explicit flank and parenthesized mass syntax."""
    if not isinstance(raw, str):
        return None, None, None, "unsupported_literal"
    sequence, flank_n, flank_c = raw, None, None
    flanked = re.fullmatch(r"([A-Z-])\.(.+)\.([A-Z-])", sequence)
    if flanked:
        flank_n, sequence, flank_c = flanked.groups()
    sequence = re.sub(r"\([+-][0-9]+(?:\.[0-9]+)?\)", "", sequence)
    if not re.fullmatch(r"[ACDEFGHIKLMNPQRSTVWY]+", sequence):
        return None, flank_n, flank_c, "unsupported_or_noncanonical_sequence"
    return sequence, flank_n, flank_c, "recognized_source_notation"


def _rows(path, block):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book[block["sheet"]]
        first_col, last_col = block["column_bounds"]
        rows = sheet.iter_rows(min_col=first_col, max_col=last_col)
        description = next(rows)
        headers = next(rows)
        actual = {c.column_letter: c.value for c in headers if c.value is not None}
        if actual != block["headers"]:
            raise ValueError(f"Source header drift: {block['block_id']}")
        description_actual = {c.column_letter: c.value for c in description if c.value is not None}
        if description_actual != block["description"]:
            raise ValueError(f"Source description drift: {block['block_id']}")
        for number, cells in enumerate(rows, 3):
            values = {
                openpyxl.utils.get_column_letter(index): cell.value
                for index, cell in enumerate(cells, first_col)
            }
            peptide = values[block["columns"]["peptide"]]
            if peptide is None or peptide == "":
                continue
            yield number, values
    finally:
        book.close()


def _workbook_block(assets, block, sinks, samples):
    columns = block["columns"]
    for index, row in _rows(assets / block["filename"], block):
        record = _id("source_record", block["block_id"], index)
        observation = _id("observation", record)
        sinks["source_records"].add(
            {
                "source_record_id": record,
                "source_file": block["filename"],
                "source_locator": f"{block['sheet']}!row:{index}:block:{block['block_id']}",
                "source_row_json": _json({"headers": block["headers"], "cells": row}),
                "block_id": block["block_id"],
                "source_row_index": index,
            }
        )
        sequence, n_flank, c_flank, status = _peptide(row[columns["peptide"]])
        length, length_status = _number(row[columns["length"]])
        length_agreement = (
            "unresolved"
            if sequence is None or length_status != "positive"
            else "matches"
            if length == len(sequence)
            else "conflicting"
        )
        parent = _raw(row[columns["parent"]])
        sinks["observations"].add(
            {
                "observation_id": observation,
                "source_record_id": record,
                "block_id": block["block_id"],
                "peptide_sequence": sequence,
                "peptide_raw": _raw(row[columns["peptide"]]),
                "peptide_modifications_raw": _raw(row.get(columns.get("ptm"))),
                "sequence_parse_status": status,
                "length_status": length_agreement,
                "flank_n_raw": n_flank,
                "flank_c_raw": c_flank,
                "parent_label_raw": parent,
                "start_raw": _raw(row[columns["start"]]),
                "end_raw": _raw(row[columns["end"]]),
                "charge_raw": _raw(row[columns["charge"]]),
                "scan_raw": _raw(row[columns["scan"]]),
                "confidence_json": _json(
                    {key: row[column] for key, column in block["confidence"].items()}
                ),
            }
        )
        for quantity in block["quantities"]:
            column = quantity["column"]
            value, value_status = _number(row[column])
            sample = samples[quantity["sample_id"]]
            sinks["quantities"].add(
                {
                    "measurement_id": _id("measurement", record, column),
                    "observation_id": observation,
                    "source_record_id": record,
                    "sample_id": sample["sample_id"],
                    "context_id": sample["context_id"],
                    "peptide_sequence": sequence,
                    "parent_label_raw": parent,
                    "source_field": block["headers"][column],
                    "source_cell": f"{block['sheet']}!{column}{index}",
                    "quantity_kind": quantity["kind"],
                    "quantitative_raw": _raw(row[column]),
                    "quantitative_value": value,
                    "value_status": value_status,
                    "measurement_unit": "arbitrary_PEAKS_" + quantity["kind"],
                    "quantity_scale": "linear_reported_source_normalization_unqualified",
                    "eligibility_status": "source_evidence_not_training_admitted",
                }
            )


def _identifications(path, sinks):
    payload = gzip.decompress(path.read_bytes()) if path.suffix == ".gz" else path.read_bytes()
    root = ET.fromstring(payload)
    namespace = root.tag.split("}")[0] + "}"
    if root.tag != namespace + "MzIdentML":
        raise ValueError("Expected an original mzIdentML asset")
    references = {}
    for node in root.iter(namespace + "DBSequence"):
        key = node.get("id")
        if key in references:
            raise ValueError("Duplicate source DBSequence ID")
        reference = _id("reference", path.name, key)
        references[key] = reference
        sequence = node.findtext(namespace + "Seq")
        sinks["references"].add(
            {
                "reference_record_id": reference,
                "source_file": path.name,
                "source_locator": "DBSequence:" + key,
                "accession_raw": node.get("accession"),
                "reference_sequence": sequence,
                "sequence_sha256": hashlib.sha256(sequence.encode()).hexdigest()
                if sequence
                else None,
                "reference_role": "original_search_entry_not_measured_parent_origin",
                "search_database_ref": node.get("searchDatabase_ref"),
                "source_row_json": _json(dict(node.attrib)),
            }
        )
    peptides = {}
    for node in root.iter(namespace + "Peptide"):
        key = node.get("id")
        if key in peptides:
            raise ValueError("Duplicate source peptide ID")
        peptides[key] = (
            node.findtext(namespace + "PeptideSequence"),
            _json(
                [
                    {"attributes": dict(m.attrib), "parameters": [dict(c.attrib) for c in m]}
                    for m in node.findall(namespace + "Modification")
                ]
            ),
        )
    evidence_ids = set()
    for node in root.iter(namespace + "PeptideEvidence"):
        key = node.get("id")
        if key in evidence_ids:
            raise ValueError("Duplicate source peptide-evidence ID")
        evidence_ids.add(key)
        sequence, modifications = peptides[node.get("peptide_ref")]
        sinks["identifications"].add(
            {
                "identification_id": _id("identification", path.name, key),
                "source_file": path.name,
                "source_locator": "PeptideEvidence:" + key,
                "reference_record_id": references[node.get("dBSequence_ref")],
                "peptide_sequence": sequence,
                "peptide_modifications_raw": modifications,
                "source_peptide_ref": node.get("peptide_ref"),
                "start_raw": node.get("start"),
                "end_raw": node.get("end"),
                "is_decoy_raw": node.get("isDecoy"),
                "source_row_json": _json(dict(node.attrib)),
            }
        )


def _compile(assets, destination, contract):
    destination.mkdir(parents=True)
    sinks = {name: _Sink(destination, name, schema=schema) for name, schema in SCHEMAS.items()}
    try:
        for asset in contract["assets"]:
            sinks["assets"].add(
                {
                    "source_asset_id": _id("asset", asset["filename"]),
                    "source_file": asset["filename"],
                    "source_asset_url": asset["source_url"],
                    "source_asset_role": asset["role"],
                    "asset_sha256": asset["sha256"],
                    "size_bytes": asset["size_bytes"],
                }
            )
        for name in ("contexts", "samples", "assertions"):
            for row in contract[name]:
                sinks[name].add(row)
        samples = {row["sample_id"]: row for row in contract["samples"]}
        if len(samples) != len(contract["samples"]):
            raise ValueError("Duplicate sample identity in source contract")
        for block in contract["blocks"]:
            _workbook_block(assets, block, sinks, samples)
        for asset in contract["assets"]:
            if asset["role"] == "original_identifications_and_reference_sequences":
                _identifications(assets / asset["filename"], sinks)
        return {"n_" + name + "_rows": sink.n_rows for name, sink in sinks.items()}
    finally:
        for sink in sinks.values():
            sink.close()


def _inventory(directory):
    return {
        str(path.relative_to(directory)): file_digest(path)
        for path in sorted(directory.rglob("*"))
        if path.is_file() and path != directory / "manifest.json"
    }


def verify_nanaware_lymph_package(directory):
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    if (
        manifest.get("schema_version"),
        manifest.get("catalog_schema_version"),
        manifest.get("package_family"),
        manifest.get("revision"),
    ) != (PACKAGE_SCHEMA_VERSION, CATALOG_SCHEMA_VERSION, PACKAGE_FAMILY, REVISION):
        raise ValueError("Unsupported Nanaware lymph package")
    if _inventory(directory) != manifest.get("files"):
        raise ValueError("Nanaware inventory or hashes differ")
    contract = json.loads(CONTRACT.read_text())
    if json.loads((directory / ROOT / "contract.json").read_text()) != contract:
        raise ValueError("Nanaware source contract differs")
    if manifest.get("source_registry") != contract["source_registry"]:
        raise ValueError("Nanaware source registry differs")
    for asset in contract["assets"]:
        if file_digest(directory / ROOT / "source_assets" / asset["filename"]) != {
            key: asset[key] for key in ("sha256", "size_bytes")
        }:
            raise ValueError("Nanaware original source mismatch")
    with tempfile.TemporaryDirectory(prefix="hitlist-lymph-verify-") as temporary:
        replay = Path(temporary) / "typed"
        counts = _compile(directory / ROOT / "source_assets", replay, contract)
        if counts != manifest.get("reconciliation"):
            raise ValueError("Nanaware source reconciliation differs")
        for name, schema in SCHEMAS.items():
            expected = pq.ParquetFile(replay / f"{name}.parquet")
            actual = pq.ParquetFile(directory / ROOT / "typed" / f"{name}.parquet")
            if not actual.schema_arrow.equals(schema):
                raise ValueError(f"Nanaware schema differs: {name}")
            for left, right in zip(expected.iter_batches(), actual.iter_batches(), strict=True):
                if not left.equals(right):
                    raise ValueError(f"Nanaware typed rows differ from original: {name}")
    return manifest


def write_nanaware_lymph_package(source_directory, destination):
    """Snapshot the reviewed source revision; do not create a training cohort."""
    source_directory, destination = Path(source_directory), Path(destination)
    if destination.exists():
        raise FileExistsError(destination)
    contract = json.loads(CONTRACT.read_text())
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{destination.name}-", dir=destination.parent))
    try:
        for asset in contract["assets"]:
            source = source_directory / asset["filename"]
            if file_digest(source) != {key: asset[key] for key in ("sha256", "size_bytes")}:
                raise ValueError(f"Nanaware source mismatch: {asset['filename']}")
            snapshot_file(source, staging / ROOT / "source_assets" / asset["filename"])
        (staging / ROOT / "contract.json").write_text(_json(contract) + "\n")
        counts = _compile(staging / ROOT / "source_assets", staging / ROOT / "typed", contract)
        from .version import __version__

        manifest = {
            "schema_version": PACKAGE_SCHEMA_VERSION,
            "catalog_schema_version": CATALOG_SCHEMA_VERSION,
            "package_family": PACKAGE_FAMILY,
            "revision": REVISION,
            "hitlist_version": __version__,
            "source_registry": contract["source_registry"],
            "reconciliation": counts,
            "files": _inventory(staging),
        }
        (staging / "manifest.json").write_text(
            json.dumps(manifest, indent=2, sort_keys=True) + "\n"
        )
        verify_nanaware_lymph_package(staging)
        staging.rename(destination)
        return destination / "manifest.json"
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
