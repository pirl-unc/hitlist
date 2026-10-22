"""Source-faithful quantitative MS adapter for pinned PXD057588 workbooks.

The six supplementary tables are one source generation. Primary peptide rows,
assay-scoped channels, observed intensities, published fold changes and their
derivation inputs have different grains. Original workbooks and every populated
cell remain available in the enclosing external evidence package.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import Counter
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq

SOURCE_ID = "pxd057588"
ASSET_RECEIPT = Path(__file__).parent / "data" / "evidence_revisions" / "pxd057588_assets.json"
SUPPORT_FILES = {
    "contrasts.json": Path(__file__).parent
    / "data"
    / "evidence_revisions"
    / "pxd057588_contract.json",
    "workbook_inventory.json": Path(__file__).parent
    / "data"
    / "evidence_revisions"
    / "pxd057588_workbook_inventory.json",
    "retrieval_receipts.json": Path(__file__).parent
    / "data"
    / "evidence_revisions"
    / "pxd057588_retrieval_receipts.json",
}
EXPECTED = {
    "n_source_assets": 17,
    "n_primary_peptide_rows": 35_772,
    "n_panels": 15,
    "n_panel_records": 62_911,
    "n_channels": 53,
    "n_source_cells": 851_830,
    "n_intensity_results": 309_126,
    "n_quantity_measurements": 509_850,
    "n_fold_change_measurements": 200_724,
    "n_contrasts": 32,
    "n_derivation_inputs": 802_896,
    "n_source_assertions": 37,
    "n_sdrf_rows": 72,
    "n_sdrf_cells": 3_240,
    "n_deposited_files": 65,
    "n_deposited_raw_files": 50,
    "n_sdrf_annotated_raw_files": 6,
}
_TABLE_CONTEXT = {
    7: ("mhc_bound", "macrophage", None, "polarized_M1_vs_untreated"),
    8: ("mhc_bound", "macrophage", None, "polarized_M2_vs_untreated"),
    9: ("mhc_bound", "macrophage", "tumor_cell", "coculture_vs_untreated"),
    10: ("mhc_bound", "macrophage", "mhc_i_loss_tumor_cell", "coculture_vs_untreated"),
    11: ("bulk_digest_peptide", "macrophage", "tumor_cell", "coculture_vs_untreated"),
    12: ("mhc_bound", "tumor_cell", "macrophage", "coculture_vs_untreated"),
}


def _schema(*fields):
    return pa.schema([pa.field(name, type_) for name, type_ in fields])


SCHEMAS = {
    "source_assets": _schema(
        ("asset_id", pa.string()),
        ("relative_path", pa.string()),
        ("sha256", pa.string()),
        ("size_bytes", pa.int64()),
        ("original_url", pa.string()),
        ("license_raw", pa.string()),
        ("license_status", pa.string()),
    ),
    "source_cells": _schema(
        ("cell_id", pa.string()),
        ("asset_path", pa.string()),
        ("sheet_name", pa.string()),
        ("row_1based", pa.int64()),
        ("column_1based", pa.int64()),
        ("cell_ref", pa.string()),
        ("excel_type", pa.string()),
        ("value_text", pa.string()),
        ("value_number", pa.float64()),
        ("source_role", pa.string()),
    ),
    "panels": _schema(
        ("panel_id", pa.string()),
        ("table_number", pa.int64()),
        ("sequence_header_cell", pa.string()),
        ("panel_role", pa.string()),
        ("n_sequence_records", pa.int64()),
    ),
    "panel_records": _schema(
        ("panel_record_id", pa.string()),
        ("panel_id", pa.string()),
        ("table_number", pa.int64()),
        ("row_1based", pa.int64()),
        ("sequence_cell", pa.string()),
        ("sequence_raw", pa.string()),
        ("parent_label_raw", pa.string()),
    ),
    "peptide_rows": _schema(
        ("peptide_row_id", pa.string()),
        ("table_number", pa.int64()),
        ("row_1based", pa.int64()),
        ("asset_path", pa.string()),
        ("sequence_cell", pa.string()),
        ("sequence_raw", pa.string()),
        ("parent_label_cell", pa.string()),
        ("parent_label_raw", pa.string()),
        ("parent_mapping_status", pa.string()),
        ("measured_pool", pa.string()),
        ("measured_population", pa.string()),
        ("partner_population", pa.string()),
        ("material_link_precision", pa.string()),
        ("selection_scope", pa.string()),
    ),
    "channels": _schema(
        ("channel_id", pa.string()),
        ("table_number", pa.int64()),
        ("channel_column", pa.string()),
        ("channel_label_raw", pa.string()),
        ("sample_label_raw", pa.string()),
        ("channel_label_cell", pa.string()),
        ("sample_label_cell", pa.string()),
        ("is_control_for_table", pa.bool_()),
        ("measured_population", pa.string()),
        ("partner_population", pa.string()),
        ("partner_loss_status", pa.string()),
        ("material_link_precision", pa.string()),
        ("intensity_total_over_retained_rows", pa.float64()),
        ("normalization_policy", pa.string()),
    ),
    "quantity_results": _schema(
        ("result_id", pa.string()),
        ("peptide_row_id", pa.string()),
        ("channel_id", pa.string()),
        ("table_number", pa.int64()),
        ("source_intensity_cell", pa.string()),
        ("measured_pool", pa.string()),
        ("endpoint", pa.string()),
        ("material_link_precision", pa.string()),
    ),
    "quantity_measurements": _schema(
        ("measurement_id", pa.string()),
        ("result_id", pa.string()),
        ("contrast_id", pa.string()),
        ("representation", pa.string()),
        ("reported_value", pa.float64()),
        ("recomputed_value", pa.float64()),
        ("absolute_replay_error", pa.float64()),
        ("unit", pa.string()),
        ("quantity_scale", pa.string()),
        ("value_status", pa.string()),
        ("source_cell", pa.string()),
        ("normalization_policy", pa.string()),
    ),
    "contrasts": _schema(
        ("contrast_id", pa.string()),
        ("table_number", pa.int64()),
        ("treated_channel_id", pa.string()),
        ("control_channel_ids_json", pa.string()),
        ("reported_fc_column", pa.string()),
        ("normalization_policy", pa.string()),
        ("normalization_universe", pa.string()),
        ("operation_order", pa.string()),
        ("method_source_locator", pa.string()),
    ),
    "derivation_inputs": _schema(
        ("derived_measurement_id", pa.string()),
        ("input_measurement_id", pa.string()),
        ("input_role", pa.string()),
        ("control_ordinal", pa.int64()),
    ),
    "source_assertions": _schema(
        ("assertion_id", pa.string()),
        ("subject_id", pa.string()),
        ("predicate", pa.string()),
        ("value_raw", pa.string()),
        ("normalized_value", pa.string()),
        ("assertion_status", pa.string()),
        ("source_locator", pa.string()),
        ("scope", pa.string()),
    ),
    "sdrf_cells": _schema(
        ("sdrf_cell_id", pa.string()),
        ("row_1based", pa.int64()),
        ("column_1based", pa.int64()),
        ("header_raw", pa.string()),
        ("value_raw", pa.string()),
        ("source_name", pa.string()),
        ("data_file", pa.string()),
    ),
    "deposited_files": _schema(
        ("source_file_id", pa.string()),
        ("file_name", pa.string()),
        ("file_category", pa.string()),
        ("size_bytes", pa.int64()),
        ("has_sdrf_annotation", pa.bool_()),
        ("material_link_precision", pa.string()),
    ),
}


class _Sink:
    def __init__(self, directory, name, *, batch_size=16384):
        self.schema = SCHEMAS[name]
        self.writer = pq.ParquetWriter(
            directory / f"{name}.parquet", self.schema, compression="zstd"
        )
        self.batch_size = batch_size
        self.rows = []
        self.n_rows = 0

    def add(self, row):
        self.rows.append(row)
        if len(self.rows) >= self.batch_size:
            self.flush()

    def flush(self):
        if self.rows:
            self.writer.write_table(pa.Table.from_pylist(self.rows, schema=self.schema))
            self.n_rows += len(self.rows)
            self.rows.clear()

    def close(self):
        self.flush()
        self.writer.close()


def _string(value):
    return None if value is None else str(value)


def _table_id(number):
    return f"pxd057588:t{number}"


def _peptide_row_id(number, row):
    return f"{_table_id(number)}:r{row}"


def _channel_id(number, column):
    return f"{_table_id(number)}:channel:{column}"


def _result_id(number, row, column):
    return f"{_peptide_row_id(number, row)}:channel:{column}"


def _source_context(number):
    pool, measured, partner, condition = _TABLE_CONTEXT[number]
    return {
        "measured_pool": pool,
        "measured_population": measured,
        "partner_population": partner,
        "condition_class": condition,
    }


def _excel_column(number):
    from openpyxl.utils import get_column_letter

    return get_column_letter(number)


def _excel_index(column):
    from openpyxl.utils import column_index_from_string

    return column_index_from_string(column)


def _source_cell_row(number, sheet_name, cell):
    value = cell.value
    if value is None:
        return None
    numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
    coordinate = cell.coordinate
    return {
        "cell_id": f"{_table_id(number)}:cell:{coordinate}",
        "asset_path": f"PXD057588/table{number}.xlsx",
        "sheet_name": sheet_name,
        "row_1based": cell.row,
        "column_1based": cell.column,
        "cell_ref": coordinate,
        "excel_type": cell.data_type,
        "value_text": str(value),
        "value_number": float(value) if numeric else None,
        "source_role": "header_or_method" if cell.row <= 5 else "source_value",
    }


def _write_panels(number, rows, spec, inventory, sinks):
    panel_count = 0
    primary_column = 1
    parent_index = _excel_index(spec["parent_column"]) - 1
    for panel in inventory["independently_keyed_sequence_blocks"]:
        header_cell = panel["sequence_header_cell"]
        column = "".join(char for char in header_cell if char.isalpha())
        column_index = _excel_index(column) - 1
        if rows[4][column_index] != "Sequence":
            raise ValueError(f"Unreviewed source panel header: table {number} {header_cell}")
        panel_id = f"{_table_id(number)}:panel:{column}"
        n_records = 0
        for row_number, row in enumerate(rows[5:], start=6):
            sequence = row[column_index]
            if sequence is None:
                continue
            if not isinstance(sequence, str) or not sequence:
                raise ValueError(f"Invalid panel sequence: table {number} row {row_number}")
            sinks["panel_records"].add(
                {
                    "panel_record_id": f"{panel_id}:r{row_number}",
                    "panel_id": panel_id,
                    "table_number": number,
                    "row_1based": row_number,
                    "sequence_cell": f"{column}{row_number}",
                    "sequence_raw": sequence,
                    "parent_label_raw": _string(row[parent_index])
                    if column_index == primary_column - 1
                    else None,
                }
            )
            n_records += 1
        if n_records != panel["populated_sequence_rows"]:
            raise ValueError(f"Source panel count changed: table {number} {header_cell}")
        sinks["panels"].add(
            {
                "panel_id": panel_id,
                "table_number": number,
                "sequence_header_cell": header_cell,
                "panel_role": "primary_quantity_block" if column_index == 0 else "additional_panel",
                "n_sequence_records": n_records,
            }
        )
        panel_count += n_records
    return panel_count


def _source_channels(number, rows, spec):
    columns = []
    for index in range(1, len(rows[3])):
        if rows[3][index] is None:
            break
        columns.append(index)
    samples = {}
    for index in columns:
        sample = rows[4][index]
        if not isinstance(sample, str) or not sample or sample in samples:
            raise ValueError(f"Invalid source channel sample label: table {number}")
        samples[sample] = index
    controls = {sample for contrast in spec["contrasts"] for sample in contrast["control_samples"]}
    if not controls <= samples.keys():
        raise ValueError(f"Source contrast lacks control channel: table {number}")
    totals = {}
    for index in columns:
        values = [row[index] for row in rows[5:]]
        if any(
            not isinstance(value, (int, float))
            or isinstance(value, bool)
            or not math.isfinite(value)
            or value <= 0
            for value in values
        ):
            raise ValueError(f"Source intensity is not a positive number: table {number}")
        totals[index] = math.fsum(values)
    return columns, samples, controls, totals


def _write_table(number, source_dir, spec, inventory, sinks, tolerance):
    try:
        import openpyxl
    except ImportError as error:
        raise ImportError("PXD057588 ingestion requires Hitlist's curation extra") from error

    path = source_dir / "assets" / "PXD057588" / f"table{number}.xlsx"
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        expected_sheet = f"Supplementary Table {number}"
        if workbook.sheetnames != [expected_sheet]:
            raise ValueError(f"Unreviewed workbook sheets: table {number}")
        sheet = workbook.active
        rows = list(sheet.values)
        if len(rows) != spec["n_rows"] + 5 or rows[4][0] != "Sequence":
            raise ValueError(f"Unreviewed workbook primary rows: table {number}")
        parent_index = _excel_index(spec["parent_column"]) - 1
        if rows[4][parent_index] != "Protein":
            raise ValueError(f"Unreviewed workbook parent header: table {number}")
        if spec["n_rows"] != inventory["primary_rows"]:
            raise ValueError(f"Pinned workbook inventory disagrees: table {number}")
        n_source_cells = 0
        for source_row in sheet.iter_rows():
            for cell in source_row:
                record = _source_cell_row(number, expected_sheet, cell)
                if record is not None:
                    sinks["source_cells"].add(record)
                    n_source_cells += 1
        context = _source_context(number)
        n_panel_records = _write_panels(number, rows, spec, inventory, sinks)
        columns, samples, controls, totals = _source_channels(number, rows, spec)
        if len(columns) != len(inventory["channels"]):
            raise ValueError(f"Source channel count differs from audit: table {number}")
        channel_map = {channel["column"]: channel for channel in inventory["channels"]}
        for index in columns:
            column = _excel_column(index + 1)
            audited = channel_map.get(column)
            if (
                audited is None
                or str(rows[3][index]) != str(audited["source_channel_label"])
                or rows[4][index] != audited["source_sample_label"]
            ):
                raise ValueError(f"Source channel map differs from audit: table {number}")
            partner_loss = (
                "reported_partner_mhc_i_loss"
                if number == 10 and rows[4][index] not in controls
                else "unreported"
                if context["partner_population"]
                else "not_applicable"
            )
            sinks["channels"].add(
                {
                    "channel_id": _channel_id(number, column),
                    "table_number": number,
                    "channel_column": column,
                    "channel_label_raw": str(rows[3][index]),
                    "sample_label_raw": rows[4][index],
                    "channel_label_cell": f"{column}4",
                    "sample_label_cell": f"{column}5",
                    "is_control_for_table": rows[4][index] in controls,
                    "measured_population": context["measured_population"],
                    "partner_population": context["partner_population"],
                    "partner_loss_status": partner_loss,
                    "material_link_precision": "assay_channel_unresolved_specimen",
                    "intensity_total_over_retained_rows": totals[index],
                    "normalization_policy": spec["policy"],
                }
            )
        contrast_by_sample = {}
        for contrast in spec["contrasts"]:
            sample = contrast["sample"]
            if sample not in samples or sample in contrast_by_sample:
                raise ValueError(f"Source contrast lacks unique treated channel: table {number}")
            if rows[4][_excel_index(contrast["fc_column"]) - 1] != contrast["expected_fc_header"]:
                raise ValueError(f"Source fold-change header differs: table {number}")
            contrast_by_sample[sample] = contrast
            treated_column = _excel_column(samples[sample] + 1)
            control_columns = [
                _excel_column(samples[control] + 1) for control in contrast["control_samples"]
            ]
            sinks["contrasts"].add(
                {
                    "contrast_id": f"{_table_id(number)}:contrast:{sample}",
                    "table_number": number,
                    "treated_channel_id": _channel_id(number, treated_column),
                    "control_channel_ids_json": json.dumps(
                        [_channel_id(number, column) for column in control_columns]
                    ),
                    "reported_fc_column": contrast["fc_column"],
                    "normalization_policy": spec["policy"],
                    "normalization_universe": "all_retained_complete_case_rows",
                    "operation_order": "channel_total_then_peptide_control_mean_then_log2_ratio"
                    if spec["policy"] == "total_signal"
                    else "peptide_control_mean_then_log2_ratio",
                    "method_source_locator": "Cui et al. 2025 / Materials and methods / MS data processing",
                }
            )
        n_fold_changes = 0
        max_error = 0.0
        for row_number, row in enumerate(rows[5:], start=6):
            sequence = row[0]
            parent = row[parent_index]
            if not isinstance(sequence, str) or not sequence:
                raise ValueError(f"Missing primary peptide: table {number} row {row_number}")
            if not isinstance(parent, str) or not parent:
                raise ValueError(f"Missing literal parent label: table {number} row {row_number}")
            peptide_row_id = _peptide_row_id(number, row_number)
            sinks["peptide_rows"].add(
                {
                    "peptide_row_id": peptide_row_id,
                    "table_number": number,
                    "row_1based": row_number,
                    "asset_path": f"PXD057588/table{number}.xlsx",
                    "sequence_cell": f"A{row_number}",
                    "sequence_raw": sequence,
                    "parent_label_cell": f"{spec['parent_column']}{row_number}",
                    "parent_label_raw": parent,
                    "parent_mapping_status": "literal_source_label_unverified",
                    "measured_pool": context["measured_pool"],
                    "measured_population": context["measured_population"],
                    "partner_population": context["partner_population"],
                    "material_link_precision": "assay_channel_unresolved_specimen",
                    "selection_scope": "source_retained_complete_case_only",
                }
            )
            for index in columns:
                column = _excel_column(index + 1)
                result_id = _result_id(number, row_number, column)
                sinks["quantity_results"].add(
                    {
                        "result_id": result_id,
                        "peptide_row_id": peptide_row_id,
                        "channel_id": _channel_id(number, column),
                        "table_number": number,
                        "source_intensity_cell": f"{column}{row_number}",
                        "measured_pool": context["measured_pool"],
                        "endpoint": "observed_source_intensity",
                        "material_link_precision": "assay_channel_unresolved_specimen",
                    }
                )
                sinks["quantity_measurements"].add(
                    {
                        "measurement_id": f"{result_id}:raw_tmt_intensity",
                        "result_id": result_id,
                        "contrast_id": None,
                        "representation": "raw_tmt_intensity",
                        "reported_value": float(row[index]),
                        "recomputed_value": None,
                        "absolute_replay_error": None,
                        "unit": "source_intensity_units",
                        "quantity_scale": "relative_tmt_signal",
                        "value_status": "source_reported",
                        "source_cell": f"{column}{row_number}",
                        "normalization_policy": "none_at_source_value",
                    }
                )
            for sample, contrast in contrast_by_sample.items():
                treated_index = samples[sample]
                control_indices = [samples[name] for name in contrast["control_samples"]]
                if spec["policy"] == "total_signal":
                    treated = row[treated_index] / totals[treated_index]
                    control_mean = math.fsum(
                        row[index] / totals[index] for index in control_indices
                    ) / len(control_indices)
                elif spec["policy"] == "unscaled":
                    treated = row[treated_index]
                    control_mean = math.fsum(row[index] for index in control_indices) / len(
                        control_indices
                    )
                else:
                    raise ValueError(f"Unreviewed source normalization policy: table {number}")
                recalculated = math.log2(treated / control_mean)
                fc_column = contrast["fc_column"]
                reported = row[_excel_index(fc_column) - 1]
                if not isinstance(reported, (int, float)) or not math.isfinite(reported):
                    raise ValueError(f"Invalid published fold change: table {number}")
                error = abs(recalculated - reported)
                if error > tolerance:
                    raise ValueError(
                        f"Published fold change fails arithmetic replay: table {number} "
                        f"{fc_column}{row_number}"
                    )
                max_error = max(max_error, error)
                treated_column = _excel_column(treated_index + 1)
                result_id = _result_id(number, row_number, treated_column)
                measurement_id = f"{result_id}:reported_log2_fold_change"
                contrast_id = f"{_table_id(number)}:contrast:{sample}"
                sinks["quantity_measurements"].add(
                    {
                        "measurement_id": measurement_id,
                        "result_id": result_id,
                        "contrast_id": contrast_id,
                        "representation": "reported_log2_fold_change",
                        "reported_value": float(reported),
                        "recomputed_value": recalculated,
                        "absolute_replay_error": error,
                        "unit": "log2_ratio",
                        "quantity_scale": "relative_to_contrast_controls",
                        "value_status": "source_reported_verified_derivation",
                        "source_cell": f"{fc_column}{row_number}",
                        "normalization_policy": spec["policy"],
                    }
                )
                sinks["derivation_inputs"].add(
                    {
                        "derived_measurement_id": measurement_id,
                        "input_measurement_id": f"{result_id}:raw_tmt_intensity",
                        "input_role": "treated",
                        "control_ordinal": None,
                    }
                )
                for ordinal, control_index in enumerate(control_indices, start=1):
                    control_column = _excel_column(control_index + 1)
                    sinks["derivation_inputs"].add(
                        {
                            "derived_measurement_id": measurement_id,
                            "input_measurement_id": f"{_result_id(number, row_number, control_column)}:raw_tmt_intensity",
                            "input_role": "control",
                            "control_ordinal": ordinal,
                        }
                    )
                n_fold_changes += 1
        expected_fc = spec["n_rows"] * len(spec["contrasts"])
        if n_fold_changes != expected_fc:
            raise ValueError(f"Source fold-change count changed: table {number}")
        return {
            "table_number": number,
            "n_primary_peptide_rows": spec["n_rows"],
            "n_panel_records": n_panel_records,
            "n_channels": len(columns),
            "n_source_cells": n_source_cells,
            "n_intensity_results": spec["n_rows"] * len(columns),
            "n_fold_change_measurements": n_fold_changes,
            "maximum_absolute_replay_error_log2": max_error,
        }
    finally:
        workbook.close()


def _assertion(sinks, subject, predicate, value, normalized, status, locator, scope):
    sinks["source_assertions"].add(
        {
            "assertion_id": f"{subject}:claim:{sinks['source_assertions'].n_rows + len(sinks['source_assertions'].rows)}",
            "subject_id": subject,
            "predicate": predicate,
            "value_raw": value,
            "normalized_value": normalized,
            "assertion_status": status,
            "source_locator": locator,
            "scope": scope,
        }
    )


def _write_assertions(source_dir, inventory_by_table, sinks):
    project = json.loads((source_dir / "assets" / "PXD057588" / "project.json").read_text())
    instruments = project.get("instruments") or []
    if len(instruments) != 1 or instruments[0].get("name") != "Orbitrap Exploris 480":
        raise ValueError("Pinned project instrument assertion changed")
    _assertion(
        sinks,
        "pxd057588:project",
        "instrument",
        instruments[0]["name"],
        instruments[0]["name"],
        "source_reported",
        "PXD057588/project.json::instruments/0/name",
        "project_not_individual_run",
    )
    _assertion(
        sinks,
        "pxd057588:project",
        "search_q_value_policy",
        "q <= 0.05",
        "q_le_0.05_study_method",
        "source_reported_method_not_row_q_value",
        "Cui et al. 2025 / Materials and methods / MS data processing",
        "study_method",
    )
    for number, inventory in sorted(inventory_by_table.items()):
        context = _source_context(number)
        subject = _table_id(number)
        header = {cell["cell"]: str(cell["value"]) for cell in inventory["header_cells"]}
        title_cells = [
            cell
            for cell in inventory["header_cells"]
            if cell["cell"].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ") == "1"
        ]
        if len(title_cells) != 1:
            raise ValueError(f"Unreviewed source title cells: table {number}")
        title_cell = title_cells[0]["cell"]
        for predicate, value in (
            ("measured_pool", context["measured_pool"]),
            ("measured_population", context["measured_population"]),
            ("condition_class", context["condition_class"]),
        ):
            _assertion(
                sinks,
                subject,
                predicate,
                header[title_cell],
                value,
                "curated_from_source_title_and_method",
                f"PXD057588/table{number}.xlsx!{title_cell}",
                "assay_table",
            )
        if context["partner_population"]:
            _assertion(
                sinks,
                subject,
                "partner_population",
                header.get("A2", ""),
                context["partner_population"],
                "curated_from_source_title_and_method",
                f"PXD057588/table{number}.xlsx!A2",
                "assay_table",
            )
        _assertion(
            sinks,
            subject,
            "selection_scope",
            header.get("A2", ""),
            "retained_complete_case_rows_only",
            "source_reported",
            f"PXD057588/table{number}.xlsx!A2",
            "assay_table",
        )
    table8 = inventory_by_table[8]
    header8 = {cell["cell"]: str(cell["value"]) for cell in table8["header_cells"]}
    if "anti-" not in header8["C1"].lower():
        raise ValueError("Pinned Table 8 title assertion changed")
    if "pro-inflammatory" not in header8["A2"]:
        raise ValueError("Pinned Table 8 conflicting assertion changed")
    for locator, value, normalized in (
        ("C1", header8["C1"], "anti_inflammatory"),
        ("A2", header8["A2"], "pro_inflammatory"),
    ):
        _assertion(
            sinks,
            _table_id(8),
            "polarization_class",
            value,
            normalized,
            "conflicting_source_assertion",
            f"PXD057588/table8.xlsx!{locator}",
            "assay_table",
        )
    table10 = inventory_by_table[10]
    header10 = {cell["cell"]: str(cell["value"]) for cell in table10["header_cells"]}
    if "n=3" not in header10["A3"]:
        raise ValueError("Pinned Table 10 replicate assertion changed")
    _assertion(
        sinks,
        _table_id(10),
        "replication_description",
        header10["A3"],
        "reported_n_3",
        "conflicting_source_assertion",
        "PXD057588/table10.xlsx!A3",
        "assay_table",
    )
    _assertion(
        sinks,
        _table_id(10),
        "replication_description",
        "G1-G4 and C1-C4 channel labels",
        "four_labeled_channels_per_partner_arm_not_donor_count",
        "conflicting_source_assertion",
        "PXD057588/table10.xlsx!B5:L5",
        "assay_table",
    )
    _assertion(
        sinks,
        _table_id(10),
        "mhc_i_loss_target_population",
        "tumor-cell MHC-I loss",
        "tumor_partner_not_measured_macrophage",
        "source_reported",
        "Cui et al. 2025 / Table 10 and Materials and methods",
        "partner_population",
    )
    _assertion(
        sinks,
        _table_id(11),
        "source_material_workflow",
        "post-IP lysate global proteomics",
        "same_workflow_no_exact_aliquot_id",
        "source_reported_method",
        "Cui et al. 2025 / Materials and methods / Global proteome data",
        "assay_table",
    )


def _write_sdrf_and_files(source_dir, sinks):
    path = source_dir / "assets" / "PXD057588" / "community.sdrf.tsv"
    with path.open(newline="") as stream:
        rows = list(csv.reader(stream, delimiter="\t"))
    if len(rows) != 73 or len(rows[0]) != 45:
        raise ValueError("Pinned community SDRF dimensions changed")
    header = rows[0]
    if header[20] != "comment[data file]":
        raise ValueError("Pinned community SDRF data-file column changed")
    annotated_files = set()
    for row_number, row in enumerate(rows[1:], start=2):
        if len(row) != 45:
            raise ValueError("Community SDRF positional fields changed")
        source_name, file_name = row[0], row[20]
        annotated_files.add(file_name)
        for column_number, (field, value) in enumerate(zip(header, row), start=1):
            sinks["sdrf_cells"].add(
                {
                    "sdrf_cell_id": f"pxd057588:sdrf:r{row_number}:c{column_number}",
                    "row_1based": row_number,
                    "column_1based": column_number,
                    "header_raw": field,
                    "value_raw": value,
                    "source_name": source_name,
                    "data_file": file_name,
                }
            )
    deposited = json.loads((source_dir / "assets" / "PXD057588" / "files.json").read_text())
    if not isinstance(deposited, list):
        raise ValueError("Pinned PRIDE file inventory changed")
    raw_files = set()
    categories = Counter()
    for source_file in deposited:
        name = source_file["fileName"]
        category = source_file["fileCategory"]["value"]
        categories[category] += 1
        if category == "RAW":
            raw_files.add(name)
        sinks["deposited_files"].add(
            {
                "source_file_id": source_file["accession"],
                "file_name": name,
                "file_category": category,
                "size_bytes": source_file["fileSizeBytes"],
                "has_sdrf_annotation": name in annotated_files,
                "material_link_precision": "sdrf_file_name_only"
                if name in annotated_files
                else "unannotated_in_community_sdrf",
            }
        )
    if not annotated_files <= raw_files or categories != Counter(
        {"RAW": 50, "PEAK": 7, "RESULT": 7, "EXPERIMENTAL DESIGN": 1}
    ):
        raise ValueError("Pinned PRIDE raw-file/SDRF scope changed")
    _assertion(
        sinks,
        "pxd057588:sdrf",
        "treatment",
        "not available",
        "unreported",
        "source_reported_missingness",
        "PXD057588/community.sdrf.tsv::column/8",
        "six_annotated_tumor_raw_files_only",
    )
    return {
        "n_sdrf_rows": len(rows) - 1,
        "n_sdrf_cells": (len(rows) - 1) * len(header),
        "n_deposited_files": len(deposited),
        "n_deposited_raw_files": len(raw_files),
        "n_sdrf_annotated_raw_files": len(annotated_files),
    }


def _write_source_assets(source_dir, sinks):
    receipt = json.loads((source_dir / "contracts" / "source_assets.json").read_text())
    retrieval = json.loads((source_dir / "contracts" / "retrieval_receipts.json").read_text())
    by_name = {entry["name"]: entry for entry in retrieval["files"]}
    for relative, detail in sorted(receipt["assets"].items()):
        source_license = by_name.get(relative, {}).get("license")
        license_raw = f"{source_license['name']} 4.0" if source_license is not None else None
        sinks["source_assets"].add(
            {
                "asset_id": f"pxd057588:asset:{relative}",
                "relative_path": relative,
                "sha256": detail["sha256"],
                "size_bytes": detail["size"],
                "original_url": detail["url"],
                "license_raw": license_raw,
                "license_status": "publisher_reported"
                if source_license is not None
                else "unverified_for_snapshot",
            }
        )
    return len(receipt["assets"])


class PXD057588Adapter:
    """Materialize one pinned external source using the shared package contract."""

    source_id = SOURCE_ID
    schema_version = 1
    asset_receipt_path = ASSET_RECEIPT
    support_files = SUPPORT_FILES

    def write(self, source_dir):
        source_dir = Path(source_dir)
        typed = source_dir / "typed"
        typed.mkdir()
        sinks = {name: _Sink(typed, name) for name in SCHEMAS}
        try:
            contract = json.loads((source_dir / "contracts" / "contrasts.json").read_text())
            inventory = json.loads(
                (source_dir / "contracts" / "workbook_inventory.json").read_text()
            )
            inventory_by_table = {entry["table"]: entry for entry in inventory}
            if contract["input_assets_sha256"] != (
                hashlib.sha256(
                    (source_dir / "contracts" / "source_assets.json").read_bytes()
                ).hexdigest()
            ):
                raise ValueError("Quantitative contract source receipt hash changed")
            if set(inventory_by_table) != {7, 8, 9, 10, 11, 12}:
                raise ValueError("Unreviewed external workbook inventory")
            n_source_assets = _write_source_assets(source_dir, sinks)
            table_counts = []
            for spec in contract["tables"]:
                number = spec["table"]
                table_counts.append(
                    _write_table(
                        number,
                        source_dir,
                        spec,
                        inventory_by_table[number],
                        sinks,
                        contract["absolute_tolerance_log2"],
                    )
                )
            _write_assertions(source_dir, inventory_by_table, sinks)
            sdrf = _write_sdrf_and_files(source_dir, sinks)
        finally:
            for sink in sinks.values():
                sink.close()
        reconciliation = {
            "table_counts": table_counts,
            "maximum_absolute_replay_error_log2": max(
                row["maximum_absolute_replay_error_log2"] for row in table_counts
            ),
            "n_source_assets": n_source_assets,
            "n_primary_peptide_rows": sinks["peptide_rows"].n_rows,
            "n_panels": sinks["panels"].n_rows,
            "n_panel_records": sinks["panel_records"].n_rows,
            "n_channels": sinks["channels"].n_rows,
            "n_source_cells": sinks["source_cells"].n_rows,
            "n_intensity_results": sinks["quantity_results"].n_rows,
            "n_quantity_measurements": sinks["quantity_measurements"].n_rows,
            "n_fold_change_measurements": sum(
                row["n_fold_change_measurements"] for row in table_counts
            ),
            "n_contrasts": sinks["contrasts"].n_rows,
            "n_derivation_inputs": sinks["derivation_inputs"].n_rows,
            "n_source_assertions": sinks["source_assertions"].n_rows,
            **sdrf,
        }
        self.verify(source_dir, reconciliation)
        return reconciliation

    def verify(self, source_dir, reconciliation):
        source_dir = Path(source_dir)
        for name, schema in SCHEMAS.items():
            file = source_dir / "typed" / f"{name}.parquet"
            parquet = pq.ParquetFile(file)
            key = {
                "source_assets": "n_source_assets",
                "source_cells": "n_source_cells",
                "panels": "n_panels",
                "panel_records": "n_panel_records",
                "peptide_rows": "n_primary_peptide_rows",
                "channels": "n_channels",
                "quantity_results": "n_intensity_results",
                "quantity_measurements": "n_quantity_measurements",
                "contrasts": "n_contrasts",
                "derivation_inputs": "n_derivation_inputs",
                "source_assertions": "n_source_assertions",
                "sdrf_cells": "n_sdrf_cells",
                "deposited_files": "n_deposited_files",
            }[name]
            if parquet.schema_arrow != schema or parquet.metadata.num_rows != reconciliation[key]:
                raise ValueError(f"External typed relation mismatch: {name}")
        for key, expected in EXPECTED.items():
            if reconciliation.get(key) != expected:
                raise ValueError(f"Pinned external source count changed: {key}")
        if (
            reconciliation["n_quantity_measurements"]
            != reconciliation["n_intensity_results"] + reconciliation["n_fold_change_measurements"]
            or reconciliation["n_derivation_inputs"]
            != 4 * reconciliation["n_fold_change_measurements"]
            or reconciliation["n_deposited_raw_files"] != 50
            or reconciliation["maximum_absolute_replay_error_log2"] > 1e-8
        ):
            raise ValueError("External quantitative conservation mismatch")

    def registry(self, reconciliation):
        return {
            "source_id": self.source_id,
            "dataset_accession": "PXD057588",
            "organism_scope": "mouse_tables_7_to_12_only",
            "measured_pools": ["mhc_bound", "bulk_digest_peptide"],
            "measurement_grain": "peptide_row_by_assay_scoped_channel",
            "rights": {
                "publisher_workbooks": "CC BY 4.0",
                "PRIDE_deposit": "CC0",
                "API_and_SDRF_snapshots": "unverified",
            },
            "training_admission": "deferred_pending_material_mapping_and_selection_review",
            "n_primary_peptide_rows": reconciliation["n_primary_peptide_rows"],
            "n_intensity_results": reconciliation["n_intensity_results"],
            "n_fold_change_measurements": reconciliation["n_fold_change_measurements"],
            "candidate_sources": {
                "PXD009092": "discovery_snapshots_only_no_verified_measurement_table",
                "PXD017407": "contract_only_absolute_calibration_scope_unreviewed",
                "PXD013455": "contract_only_fraction_independence_unresolved",
                "PXD056426": "contract_only_control_loss_quantities_unreviewed",
            },
            "source_cells_artifact": "typed/source_cells.parquet",
        }
