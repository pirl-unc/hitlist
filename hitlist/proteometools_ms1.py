"""Original nontryptic synthesis-pool MS1 signals, not peptide concentration labels."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import pyarrow.parquet as pq

from .evidence_snapshot import snapshot_file
from .mapp_evidence import _schema, _Sink
from .provenance import file_digest

PACKAGE_SCHEMA_VERSION = 23
CATALOG_SCHEMA_VERSION = 28
PACKAGE_FAMILY = "proteometools_ms1"
REVISION = "pxd021013_ms1_v1"
ROOT = Path("proteometools_ms1")
CORE = Path(__file__).parent / "data/evidence_revisions/proteometools_ms1_core.json"
BASE_URL = "https://ftp.pride.ebi.ac.uk/pride/data/archive/2021/04/PXD021013/"
MAPPING_SHEETS = {"HLA Class I": 168689, "HLA Class II": 73465, "AspN": 31745, "LysN": 31436}
CONTEXT = {
    "context_id": "pxd021013:synthetic_pool_protocol",
    "source_url": "https://doi.org/10.1038/s41467-021-23713-9",
    "material_type": "chemically_synthesized_peptide_pool",
    "sequence_species": "human_library_origin_not_measured_cell_species",
    "instrument": "Orbitrap_Fusion_Lumos_primary_protocol",
    "laboratory_digest": "not_applicable_to_chemical_synthesis",
    "nominal_load_alternatives_fmol_json": "[200,500]",
    "per_run_load_status": "unresolved",
    "individual_purity_status": "unreported",
    "individual_molar_amount_status": "unresolved",
    "sdrf_conflicts_json": '["material=cell","instrument=Fusion","all_HCD","all_run1"]',
    "interpretation": "synthesis_handling_LC_MS_response_not_pure_detector_or_APC_concentration",
}
SCHEMAS = {
    "assets": _schema(
        ("asset_id", "source_file", "source_url", "source_asset_role", "sha256"), ("size_bytes",)
    ),
    "source_records": _schema(
        ("source_record_id", "source_file", "source_locator", "source_row_json", "source_role"),
        ("source_row_index",),
    ),
    "planned_products": _schema(
        (
            "planned_id",
            "source_file",
            "source_locator",
            "pool_name",
            "peptide_sequence",
            "library_origin",
        ),
        ("source_row_index",),
    ),
    "contexts": _schema(tuple(CONTEXT)),
    "runs": _schema(
        (
            "run_id",
            "source_record_id",
            "context_id",
            "source_file",
            "source_name",
            "raw_file",
            "pool_name",
            "library_origin",
            "search_index_id",
            "raw_index_id",
            "experiment_raw",
            "fixed_modifications_raw",
            "variable_modifications_raw",
            "search_enzyme_raw",
            "search_enzyme_mode_raw",
            "software_version_raw",
            "acquisition_label_raw",
        )
    ),
}
_QUANTITY_TEXT = (
    "measurement_id",
    "source_record_id",
    "run_id",
    "peptide_sequence",
    "source_id_raw",
    "quantitative_raw",
    "value_status",
    "measurement_unit",
    "quantity_scale",
    "pep_raw",
    "pep_status",
    "reverse_raw",
    "contaminant_raw",
    "parent_label_raw",
)
SCHEMAS["peptides"] = _schema(_QUANTITY_TEXT, numbers=("quantitative_value", "pep"))
SCHEMAS["forms"] = _schema(
    (*_QUANTITY_TEXT, "peptide_measurement_id", "peptide_modifications_raw", "charges_raw"),
    numbers=("quantitative_value", "pep"),
)


def _json(value):
    return json.dumps(value, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _id(kind, *values):
    return f"pxd021013:{kind}:" + hashlib.sha256(_json(values).encode()).hexdigest()


def _number(raw):
    if raw == "":
        return None, "missing"
    try:
        value = float(raw)
    except ValueError:
        return None, "invalid_literal"
    if not math.isfinite(value):
        return None, "nonfinite"
    return value, "zero" if value == 0 else "positive" if value > 0 else "negative"


def _rows(stream, file, member, sinks, *, positional=False):
    reader = csv.reader(stream, delimiter="\t")
    headers = next(reader)
    if not positional and len(set(headers)) != len(headers):
        raise ValueError(f"Duplicate source headers: {file}:{member}")
    for index, cells in enumerate(reader, 2):
        if len(cells) != len(headers):
            raise ValueError(f"Malformed source row: {file}:{member}:{index}")
        record = _id("record", file, member, index)
        sinks["source_records"].add(
            {
                "source_record_id": record,
                "source_file": file,
                "source_locator": f"{member}:row:{index}",
                "source_row_index": index,
                "source_row_json": _json({"headers": headers, "cells": cells}),
                "source_role": member,
            }
        )
        yield record, dict(zip(headers, cells, strict=True)) if not positional else (headers, cells)


def _read_zip(archive, member, filename, sinks):
    with archive.open(member) as stream:
        return list(
            _rows(
                io.TextIOWrapper(stream, encoding="utf-8-sig", newline=""), filename, member, sinks
            )
        )


def _planned(assets, sinks):
    name = "article-2021-planned-peptides.xlsx"
    book = openpyxl.load_workbook(assets / name, read_only=True, data_only=False)
    pools = {}
    try:
        for sheet_name, last_row in MAPPING_SHEETS.items():
            sheet = book[sheet_name]
            rows = sheet.iter_rows(
                max_row=min(sheet.max_row, last_row), max_col=2, values_only=True
            )
            if next(rows) != ("Pool name", "Sequence"):
                raise ValueError("Intended-product header changed")
            for index, (pool, sequence) in enumerate(rows, 2):
                if (
                    not isinstance(pool, str)
                    or not isinstance(sequence, str)
                    or not pool
                    or not sequence
                    or pool.startswith("=")
                    or sequence.startswith("=")
                ):
                    raise ValueError("Invalid intended-product literal")
                if pool in pools and pools[pool] != sheet_name:
                    raise ValueError("Pool has ambiguous library origin")
                pools[pool] = sheet_name
                sinks["planned_products"].add(
                    {
                        "planned_id": _id("planned", sheet_name, index),
                        "source_file": name,
                        "source_locator": f"{sheet_name}!A{index}:B{index}",
                        "source_row_index": index,
                        "pool_name": pool,
                        "peptide_sequence": sequence,
                        "library_origin": sheet_name,
                    }
                )
    finally:
        book.close()
    return pools


def _quantity(row, record, run_id, kind):
    value, status = _number(row["Intensity"])
    pep, pep_status = _number(row["PEP"])
    channel_columns = [key for key in row if key.startswith("Intensity ")]
    if len(channel_columns) != 1 or row[channel_columns[0]] != row["Intensity"]:
        raise ValueError("Single-run channel and aggregate intensity disagree")
    return {
        "measurement_id": _id(kind, run_id, row["id"]),
        "source_record_id": record,
        "run_id": run_id,
        "peptide_sequence": row["Sequence"],
        "source_id_raw": row["id"],
        "quantitative_raw": row["Intensity"],
        "quantitative_value": value,
        "value_status": status,
        "measurement_unit": "arbitrary_MS1_signal",
        "quantity_scale": "linear_reported_MaxQuant_intensity",
        "pep_raw": row["PEP"],
        "pep": pep,
        "pep_status": pep_status,
        "reverse_raw": row["Reverse"],
        "contaminant_raw": row["Potential contaminant"],
        "parent_label_raw": row["Proteins"],
    }


def _compile(assets, typed, receipt):
    typed.mkdir(parents=True)
    sinks = {name: _Sink(typed, name, schema=schema) for name, schema in SCHEMAS.items()}
    try:
        for spec in receipt["assets"]:
            sinks["assets"].add(
                {
                    "asset_id": _id("asset", spec["filename"]),
                    "source_file": spec["filename"],
                    "source_url": spec["source_url"],
                    "source_asset_role": spec["role"],
                    "sha256": spec["sha256"],
                    "size_bytes": spec["size_bytes"],
                }
            )
        sinks["contexts"].add(CONTEXT)
        pools = _planned(assets, sinks)
        with (assets / "PXD021013-README.txt").open(encoding="utf-8-sig", newline="") as stream:
            index_rows = [row for _, row in _rows(stream, "PXD021013-README.txt", "README", sinks)]
        index = {row["ID"]: row for row in index_rows}
        if len(index) != len(index_rows):
            raise ValueError("Duplicate deposited file ID")
        searches = {row["NAME"]: row for row in index_rows if row["TYPE"] == "SEARCH"}
        if len(searches) != sum(row["TYPE"] == "SEARCH" for row in index_rows):
            raise ValueError("Duplicate deposited SEARCH filename")
        with (assets / "PXD021013-sdrf.tsv").open(encoding="utf-8-sig", newline="") as stream:
            list(_rows(stream, "PXD021013-sdrf.tsv", "SDRF", sinks, positional=True))
        for spec in receipt["assets"]:
            if spec["role"] != "search_archive":
                continue
            source_name, filename = spec["source_name"], spec["filename"]
            if source_name not in searches:
                raise ValueError("Archive is not a deposited SEARCH filename")
            search = searches[source_name]
            raw = index.get(search["MAPPINGS"])
            if raw is None or raw["TYPE"] != "RAW":
                raise ValueError("SEARCH does not map to exactly one RAW file")
            matched = [pool for pool in pools if source_name.startswith(pool + "_")]
            if len(matched) != 1:
                raise ValueError("SEARCH pool is unresolved or ambiguous")
            pool = matched[0]
            run_id = _id("run", source_name)
            with zipfile.ZipFile(assets / filename) as archive:
                if len(archive.namelist()) != len(set(archive.namelist())):
                    raise ValueError("Duplicate ZIP member")
                summaries = _read_zip(archive, "summary.txt", filename, sinks)
                runs = [
                    (record, row)
                    for record, row in summaries
                    if row["Raw file"] + ".raw" == raw["NAME"]
                ]
                if len(runs) != 1:
                    raise ValueError("Summary RAW identity disagrees with deposited mapping")
                record, run = runs[0]
                aggregate_labels = set()
                for other_record, row in summaries:
                    if other_record == record:
                        continue
                    if (
                        row["Raw file"] not in (run["Experiment"], "Total")
                        or row["Experiment"]
                        or row["Raw file"] in aggregate_labels
                    ):
                        raise ValueError("Unsupported extra run or ambiguous summary aggregate")
                    aggregate_labels.add(row["Raw file"])
                parameters = _read_zip(archive, "parameters.txt", filename, sinks)
                params = {}
                for key in ("Version", "Fixed modifications"):
                    values = {row["Value"] for _, row in parameters if row["Parameter"] == key}
                    if len(values) > 1:
                        raise ValueError(f"Conflicting required search parameter: {key}")
                    params[key] = next(iter(values)) if values else None
                sinks["runs"].add(
                    {
                        "run_id": run_id,
                        "source_record_id": record,
                        "context_id": CONTEXT["context_id"],
                        "source_file": filename,
                        "source_name": source_name,
                        "raw_file": raw["NAME"],
                        "pool_name": pool,
                        "library_origin": pools[pool],
                        "search_index_id": search["ID"],
                        "raw_index_id": raw["ID"],
                        "experiment_raw": run["Experiment"],
                        "fixed_modifications_raw": params.get("Fixed modifications"),
                        "variable_modifications_raw": run["Variable modifications"],
                        "search_enzyme_raw": run["Enzyme"],
                        "search_enzyme_mode_raw": run["Enzyme mode"],
                        "software_version_raw": params.get("Version"),
                        "acquisition_label_raw": source_name[len(pool) + 1 :].removesuffix(".zip"),
                    }
                )
                peptides = {}
                for record, row in _read_zip(archive, "peptides.txt", filename, sinks):
                    value = _quantity(row, record, run_id, "peptide")
                    if row["id"] in peptides:
                        raise ValueError("Duplicate peptide ID")
                    peptides[row["id"]] = value
                    sinks["peptides"].add(value)
                forms = set()
                for record, row in _read_zip(
                    archive, "modificationSpecificPeptides.txt", filename, sinks
                ):
                    value = _quantity(row, record, run_id, "form")
                    if row["id"] in forms:
                        raise ValueError("Duplicate molecular-form ID")
                    forms.add(row["id"])
                    peptide = peptides.get(row["Peptide ID"])
                    if peptide is None or peptide["peptide_sequence"] != row["Sequence"]:
                        raise ValueError("Molecular-form/peptide identity disagrees")
                    if row["Raw file"] != run["Raw file"]:
                        raise ValueError("Molecular-form RAW identity disagrees")
                    value.update(
                        peptide_measurement_id=peptide["measurement_id"],
                        peptide_modifications_raw=row["Modifications"],
                        charges_raw=row["Charges"],
                    )
                    sinks["forms"].add(value)
        return {"n_" + name + "_rows": sink.n_rows for name, sink in sinks.items()}
    finally:
        for sink in sinks.values():
            sink.close()


def _inventory(directory):
    return {
        str(p.relative_to(directory)): file_digest(p)
        for p in sorted(directory.rglob("*"))
        if p.is_file() and p != directory / "manifest.json"
    }


def _receipt(archives):
    core = json.loads(CORE.read_text())
    required = {"filename", "source_name", "sha256", "size_bytes"}
    assets = list(core["assets"])
    for spec in archives:
        if set(spec) != required:
            raise ValueError(
                "Archive receipt requires filename, source_name, sha256 and size_bytes"
            )
        for field in ("filename", "source_name"):
            if (
                not isinstance(spec[field], str)
                or Path(spec[field]).name != spec[field]
                or not spec[field].endswith(".zip")
            ):
                raise ValueError("Unsafe archive filename")
        if (
            not re.fullmatch("[0-9a-f]{64}", spec["sha256"])
            or type(spec["size_bytes"]) is not int
            or spec["size_bytes"] <= 0
        ):
            raise ValueError("Invalid archive digest")
        assets.append(
            {**spec, "role": "search_archive", "source_url": BASE_URL + spec["source_name"]}
        )
    if (
        not archives
        or len({a["filename"] for a in assets}) != len(assets)
        or len({a["source_name"] for a in archives}) != len(archives)
    ):
        raise ValueError("Empty or duplicate archive receipts")
    return {"revision": REVISION, "assets": assets}


def _source_registry():
    return {
        "pxd021013": {
            "source_url": CONTEXT["source_url"],
            "source_role": "non_MHC_synthesis_pool_MS1",
            "training_status": "registered_relative_response_recipe",
            "quantity_scope": CONTEXT["interpretation"],
            "archive_hash_authority": "caller_pinned_acquisition_receipt_not_remote_signature",
        }
    }


def verify_proteometools_ms1_package(directory):
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    if (
        manifest.get("schema_version"),
        manifest.get("catalog_schema_version"),
        manifest.get("package_family"),
        manifest.get("revision"),
    ) != (PACKAGE_SCHEMA_VERSION, CATALOG_SCHEMA_VERSION, PACKAGE_FAMILY, REVISION):
        raise ValueError("Unsupported ProteomeTools package")
    if _inventory(directory) != manifest.get("files"):
        raise ValueError("ProteomeTools inventory or hashes differ")
    receipt = json.loads((directory / ROOT / "assets.json").read_text())
    archives = [
        {k: spec[k] for k in ("filename", "source_name", "sha256", "size_bytes")}
        for spec in receipt["assets"]
        if spec["role"] == "search_archive"
    ]
    if receipt != _receipt(archives) or manifest.get("source_registry") != _source_registry():
        raise ValueError("ProteomeTools source contract differs")
    for spec in receipt["assets"]:
        if file_digest(directory / ROOT / "source_assets" / spec["filename"]) != {
            k: spec[k] for k in ("sha256", "size_bytes")
        }:
            raise ValueError("ProteomeTools original source mismatch")
    with tempfile.TemporaryDirectory(prefix="hitlist-ms1-verify-") as temporary:
        replay = Path(temporary) / "typed"
        counts = _compile(directory / ROOT / "source_assets", replay, receipt)
        if counts != manifest.get("reconciliation"):
            raise ValueError("ProteomeTools source reconciliation differs")
        for name, schema in SCHEMAS.items():
            expected = pq.ParquetFile(replay / f"{name}.parquet")
            actual = pq.ParquetFile(directory / ROOT / "typed" / f"{name}.parquet")
            if not actual.schema_arrow.equals(schema):
                raise ValueError(f"ProteomeTools schema differs: {name}")
            for left, right in zip(
                expected.iter_batches(batch_size=8192),
                actual.iter_batches(batch_size=8192),
                strict=True,
            ):
                if not left.equals(right):
                    raise ValueError(f"ProteomeTools typed rows differ from original: {name}")
    return manifest


def write_proteometools_ms1_package(source_directory, destination, *, archives):
    """Snapshot a caller-pinned deposited subset and reviewed common source assets."""
    source_directory, destination = Path(source_directory), Path(destination)
    if destination.exists():
        raise FileExistsError(destination)
    receipt = _receipt(archives)
    destination.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{destination.name}-", dir=destination.parent))
    try:
        for spec in receipt["assets"]:
            source = source_directory / spec["filename"]
            if file_digest(source) != {k: spec[k] for k in ("sha256", "size_bytes")}:
                raise ValueError(f"ProteomeTools source mismatch: {spec['filename']}")
            snapshot_file(source, staging / ROOT / "source_assets" / spec["filename"])
        (staging / ROOT / "assets.json").write_text(_json(receipt) + "\n")
        counts = _compile(staging / ROOT / "source_assets", staging / ROOT / "typed", receipt)
        from .version import __version__

        manifest = {
            "schema_version": PACKAGE_SCHEMA_VERSION,
            "catalog_schema_version": CATALOG_SCHEMA_VERSION,
            "package_family": PACKAGE_FAMILY,
            "revision": REVISION,
            "hitlist_version": __version__,
            "source_registry": _source_registry(),
            "reconciliation": counts,
            "files": _inventory(staging),
        }
        (staging / "manifest.json").write_text(
            json.dumps(manifest, indent=2, sort_keys=True) + "\n"
        )
        verify_proteometools_ms1_package(staging)
        staging.rename(destination)
        return destination / "manifest.json"
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
