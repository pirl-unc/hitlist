"""Pinned MAPP fragments: preserve raw observations, filtering and context.

No source row is a concentration label. Filtered zeros, missing evidence and
imputed publisher values retain distinct statuses. Parent occurrences are a
separate relation and never multiply quantity observations.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import statistics
from collections import Counter, defaultdict
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq

SOURCE_ID = "pxd009092"
DATA = Path(__file__).parent / "data" / "evidence_revisions"
ASSET_RECEIPT = DATA / "mapp_assets.json"
SUPPORT_FILES = {"contract.json": DATA / "mapp_contract.json"}


def _schema(text=(), integers=(), numbers=(), flags=()):
    return pa.schema(
        [(key, pa.string()) for key in text]
        + [(key, pa.int64()) for key in integers]
        + [(key, pa.float64()) for key in numbers]
        + [(key, pa.bool_()) for key in flags]
    )


SCHEMAS = {
    "source_records": _schema(("source_record_id", "asset_path", "fields_json"), ("row_1based",)),
    "peptide_rows": _schema(
        (
            "peptide_row_id",
            "source_record_id",
            "peptide_id_raw",
            "sequence_raw",
            "pep_raw",
            "pep_status",
            "parent_labels_raw",
            "leading_parent_raw",
            "group_ids_raw",
            "n_flank_raw",
            "c_flank_raw",
            "measured_pool",
            "material_link_precision",
        ),
        ("start_1based_reported", "end_1based_reported"),
        ("pep",),
        ("reverse", "potential_contaminant"),
    ),
    "groups": _schema(
        (
            "group_record_id",
            "source_record_id",
            "group_id_raw",
            "protein_ids_raw",
            "peptide_ids_raw",
            "status",
        ),
        flags=("reverse", "potential_contaminant"),
    ),
    "channels": _schema(
        (
            "channel_id",
            "source_record_id",
            "experiment_raw",
            "raw_file",
            "condition",
            "condition_role",
            "measured_species",
            "measured_cell",
            "measured_pool",
            "material_link_precision",
            "search_enzyme_mode_raw",
        ),
        ("replicate_label",),
    ),
    "evidence": _schema(
        (
            "evidence_id",
            "source_record_id",
            "peptide_row_id",
            "channel_id",
            "sequence_raw",
            "modified_sequence_raw",
            "modifications_raw",
            "type_raw",
            "group_ids_raw",
            "pep_raw",
            "pep_status",
            "intensity_raw",
            "intensity_status",
        ),
        ("msms_count",),
        ("intensity", "pep"),
        ("reverse", "potential_contaminant"),
    ),
    "quantity_results": _schema(
        (
            "result_id",
            "peptide_row_id",
            "channel_id",
            "measured_pool",
            "endpoint",
            "material_link_precision",
        ),
    ),
    "quantity_measurements": _schema(
        (
            "measurement_id",
            "result_id",
            "peptide_row_id",
            "representation",
            "value_status",
            "unit",
            "quantity_scale",
            "source_locator",
            "derivation_rule",
        ),
        numbers=("value", "recomputed_value", "absolute_replay_error"),
    ),
    "derivation_inputs": _schema(
        ("derived_measurement_id", "input_measurement_id", "input_role"),
    ),
    "source_assertions": _schema(
        (
            "assertion_id",
            "subject_id",
            "predicate",
            "value_raw",
            "normalized_value",
            "assertion_status",
            "source_locator",
            "scope",
        ),
    ),
    "publisher_cells": _schema(
        ("cell_id", "asset_path", "sheet_name", "cell_ref", "excel_type", "value_text"),
        ("row_1based", "column_1based"),
        ("value_number",),
    ),
    "publisher_quantities": _schema(
        (
            "measurement_id",
            "peptide_row_id",
            "source_cell",
            "representation",
            "condition_header_raw",
            "condition",
            "value_status",
            "unit",
            "derivation_rule",
        ),
        numbers=("value", "recomputed_value", "absolute_replay_error"),
    ),
}


class _Sink:
    def __init__(self, directory, name, *, schema=None):
        self.schema = SCHEMAS[name] if schema is None else schema
        self.writer = pq.ParquetWriter(
            directory / f"{name}.parquet", self.schema, compression="zstd"
        )
        self.rows = []
        self.n_rows = 0

    def add(self, row):
        if set(row) - set(self.schema.names):
            raise ValueError("Unknown typed MAPP field")
        self.rows.append(row)
        self.n_rows += 1
        if len(self.rows) >= 8192:
            self.flush()

    def flush(self):
        if self.rows:
            self.writer.write_table(pa.Table.from_pylist(self.rows, schema=self.schema))
            self.rows.clear()

    def close(self):
        self.flush()
        self.writer.close()


def _rows(directory, filename, sinks):
    with (directory / "assets" / filename).open(newline="") as stream:
        reader = csv.DictReader(stream, delimiter="\t")
        if len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise ValueError("Duplicate source column")
        for number, row in enumerate(reader, 2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError("Malformed source row")
            key = f"{SOURCE_ID}:{filename}:row:{number}"
            sinks["source_records"].add(
                {
                    "source_record_id": key,
                    "asset_path": filename,
                    "row_1based": number,
                    "fields_json": json.dumps(row, sort_keys=True),
                }
            )
            yield key, row


def _number(value, *, nonnegative=True):
    if value == "":
        return None
    value = float(value)
    if not math.isfinite(value) or (nonnegative and value < 0):
        raise ValueError("Invalid numeric source value")
    return value


def _integer(value):
    return int(value) if value.isdigit() else None


def _peptide_id(value):
    return f"{SOURCE_ID}:peptide:{value}"


def _channel_id(value):
    return f"{SOURCE_ID}:run:{value}"


def _result_id(pid, experiment):
    return f"{_peptide_id(pid)}:run:{experiment}"


def _measurement(
    sinks, result, representation, value, status, *, source=None, replay=None, rule=None
):
    key = f"{result}:{representation}"
    sinks["quantity_measurements"].add(
        {
            "measurement_id": key,
            "result_id": result if ":run:" in result else None,
            "peptide_row_id": result.split(":run:", 1)[0],
            "representation": representation,
            "value": value,
            "value_status": status,
            "unit": "arbitrary_intensity",
            "quantity_scale": "relative_label_free_signal",
            "source_locator": source,
            "recomputed_value": replay,
            "absolute_replay_error": abs(value - replay)
            if value is not None and replay is not None
            else None,
            "derivation_rule": rule,
        }
    )
    return key


def _link(sinks, target, source, role):
    sinks["derivation_inputs"].add(
        {"derived_measurement_id": target, "input_measurement_id": source, "input_role": role}
    )


def _write_tables(directory, sinks):
    runs, conditions = {}, defaultdict(list)
    for locator, row in _rows(directory, "txt/summary.txt", sinks):
        experiment = row["Experiment"]
        if not experiment:
            continue
        parts = experiment.rsplit("_", 2)
        if (
            len(parts) != 3
            or parts[1] not in {"UT", "TNF+IFN", "CHX", "NO-DSP", "NO-INH", "BEADS"}
            or not parts[2].isdigit()
        ):
            raise ValueError("Unknown MAPP source run condition")
        if experiment in runs or row["Raw file"] in {r["Raw file"] for r in runs.values()}:
            raise ValueError("Duplicate source run")
        condition = parts[1]
        runs[experiment] = row
        conditions[condition].append(experiment)
        sinks["channels"].add(
            {
                "channel_id": _channel_id(experiment),
                "source_record_id": locator,
                "experiment_raw": experiment,
                "raw_file": row["Raw file"],
                "condition": condition,
                "condition_role": "capture_background"
                if condition == "BEADS"
                else (
                    "procedural_control"
                    if condition in {"NO-DSP", "NO-INH"}
                    else "biological_condition"
                ),
                "replicate_label": int(parts[2]),
                "measured_species": "Homo sapiens",
                "measured_cell": "HEK293",
                "measured_pool": "proteasome_associated_fragment",
                "material_link_precision": "acquisition_run_unresolved_specimen",
                "search_enzyme_mode_raw": row["Enzyme mode"],
            }
        )
    if set(conditions) != {"UT", "TNF+IFN", "CHX", "NO-DSP", "NO-INH", "BEADS"} or any(
        len(v) != 3 for v in conditions.values()
    ):
        raise ValueError("Expected three source runs in each of six conditions")
    peptides, locators, sequences = {}, {}, set()
    for locator, row in _rows(directory, "txt/peptides.txt", sinks):
        pid = row["id"]
        if not pid.isdigit() or pid in peptides or row["Sequence"] in sequences:
            raise ValueError("Duplicate or invalid source peptide identity")
        peptides[pid], locators[pid] = row, locator
        sequences.add(row["Sequence"])
        sinks["peptide_rows"].add(
            {
                "peptide_row_id": _peptide_id(pid),
                "source_record_id": locator,
                "peptide_id_raw": pid,
                "sequence_raw": row["Sequence"],
                "parent_labels_raw": row["Proteins"],
                "leading_parent_raw": row["Leading razor protein"],
                "group_ids_raw": row["Protein group IDs"],
                "n_flank_raw": row["Amino acid before"],
                "c_flank_raw": row["Amino acid after"],
                "start_1based_reported": _integer(row["Start position"]),
                "end_1based_reported": _integer(row["End position"]),
                "pep": None if row["PEP"] in ("", "NaN") else _number(row["PEP"]),
                "pep_raw": row["PEP"],
                "pep_status": "unreported_source_nan"
                if row["PEP"] == "NaN"
                else ("unreported" if row["PEP"] == "" else "reported"),
                "reverse": row["Reverse"] == "+",
                "potential_contaminant": row["Potential contaminant"] == "+",
                "measured_pool": "proteasome_associated_fragment",
                "material_link_precision": "acquisition_run_unresolved_specimen",
            }
        )
    groups, quarantined = {}, 0
    for locator, row in _rows(directory, "txt/proteinGroups.txt", sinks):
        gid = row["id"]
        valid = gid.isdigit()
        if valid and gid in groups:
            raise ValueError("Duplicate source group ID")
        if valid:
            groups[gid] = row
        else:
            quarantined += 1
        sinks["groups"].add(
            {
                "group_record_id": locator,
                "source_record_id": locator,
                "group_id_raw": gid,
                "protein_ids_raw": row["Protein IDs"],
                "peptide_ids_raw": row["Peptide IDs"],
                "status": "source_group" if valid else "quarantined_missing_or_nonnumeric_id",
                "reverse": row["Reverse"] == "+",
                "potential_contaminant": row["Potential contaminant"] == "+",
            }
        )
    totals, counts, numeric_counts, seen = defaultdict(float), Counter(), Counter(), set()
    reciprocal = defaultdict(set)
    for locator, row in _rows(directory, "txt/evidence.txt", sinks):
        pid, experiment, eid = row["Peptide ID"], row["Experiment"], row["id"]
        if eid in seen or pid not in peptides or experiment not in runs:
            raise ValueError("Invalid evidence relationship")
        if (
            row["Sequence"] != peptides[pid]["Sequence"]
            or row["Raw file"] != runs[experiment]["Raw file"]
        ):
            raise ValueError("Evidence sequence/run relationship mismatch")
        if any(gid not in groups for gid in row["Protein group IDs"].split(";") if gid):
            raise ValueError("Evidence references missing source group")
        seen.add(eid)
        reciprocal[pid].add(eid)
        value = _number(row["Intensity"])
        totals[pid, experiment] += value or 0
        counts[pid, experiment] += 1
        numeric_counts[pid, experiment] += value is not None
        evidence_id = f"{SOURCE_ID}:evidence:{eid}"
        sinks["evidence"].add(
            {
                "evidence_id": evidence_id,
                "source_record_id": locator,
                "peptide_row_id": _peptide_id(pid),
                "channel_id": _channel_id(experiment),
                "sequence_raw": row["Sequence"],
                "modified_sequence_raw": row["Modified sequence"],
                "modifications_raw": row["Modifications"],
                "type_raw": row["Type"],
                "group_ids_raw": row["Protein group IDs"],
                "intensity_raw": row["Intensity"],
                "intensity": value,
                "intensity_status": "unreported" if value is None else "reported",
                "msms_count": _integer(row["MS/MS Count"]),
                "pep": None if row["PEP"] in ("", "NaN") else _number(row["PEP"]),
                "pep_raw": row["PEP"],
                "pep_status": "unreported_source_nan"
                if row["PEP"] == "NaN"
                else ("unreported" if row["PEP"] == "" else "reported"),
                "reverse": row["Reverse"] == "+",
                "potential_contaminant": row["Potential contaminant"] == "+",
            }
        )
        _link(
            sinks,
            _result_id(pid, experiment) + ":evidence_sum",
            evidence_id,
            "source_evidence_intensity",
        )
    columns = {
        c: c.removeprefix("Intensity ").replace("-MAPPbeadFilt", "")
        for c in next(iter(peptides.values()))
        if c.startswith("Intensity ")
    }
    if set(columns.values()) != set(runs) - set(conditions["BEADS"]) or len(columns) != 15:
        raise ValueError("Filtered quantity run coverage mismatch")
    filtered_zeros, maximum_error = 0, 0.0
    for pid, row in peptides.items():
        if set(filter(None, row["Evidence IDs"].split(";"))) != reciprocal[pid]:
            raise ValueError("Reciprocal peptide/evidence identity mismatch")
        for gid in filter(None, row["Protein group IDs"].split(";")):
            if gid not in groups or pid not in groups[gid]["Peptide IDs"].split(";"):
                raise ValueError("Reciprocal peptide/group identity mismatch")
        for experiment in runs:
            result = _result_id(pid, experiment)
            sinks["quantity_results"].add(
                {
                    "result_id": result,
                    "peptide_row_id": _peptide_id(pid),
                    "channel_id": _channel_id(experiment),
                    "measured_pool": "proteasome_associated_fragment",
                    "endpoint": "proteasome_associated_fragment_intensity",
                    "material_link_precision": "acquisition_run_unresolved_specimen",
                }
            )
            status = (
                "no_reported_evidence"
                if counts[pid, experiment] == 0
                else (
                    "all_evidence_intensities_unreported"
                    if numeric_counts[pid, experiment] == 0
                    else (
                        "partial_evidence_intensity_sum"
                        if numeric_counts[pid, experiment] < counts[pid, experiment]
                        else "reported_evidence_sum"
                    )
                )
            )
            _measurement(
                sinks,
                result,
                "evidence_sum",
                totals[pid, experiment],
                status,
                rule="sum_reported_intensities_missing_contributes_zero_only_for_filter_replay",
            )
        median = statistics.median(totals[pid, e] for e in conditions["BEADS"])
        background = _measurement(
            sinks,
            _peptide_id(pid),
            "beads_median",
            median,
            "derived_background",
            rule="median_three_source_evidence_sums",
        )
        for experiment in conditions["BEADS"]:
            _link(
                sinks, background, _result_id(pid, experiment) + ":evidence_sum", "background_run"
            )
        for column, experiment in columns.items():
            measured, raw = _number(row[column]), totals[pid, experiment]
            replay = raw if raw >= 2 * median else 0.0
            if measured is None or not math.isclose(measured, replay, rel_tol=1e-4, abs_tol=1):
                raise ValueError("Published bead filter does not replay")
            maximum_error = max(maximum_error, abs(measured - replay))
            filtered_zeros += measured == 0 and raw > 0
            status = (
                "filtered_zero"
                if measured == 0 and raw > 0
                else (
                    "reported_zero_no_positive_evidence" if measured == 0 else "reported_positive"
                )
            )
            target = _measurement(
                sinks,
                _result_id(pid, experiment),
                "bead_filtered",
                measured,
                status,
                source=locators[pid] + ":" + column,
                replay=replay,
                rule="retain_run_sum_when_greater_equal_two_times_beads_median",
            )
            _link(
                sinks, target, _result_id(pid, experiment) + ":evidence_sum", "unfiltered_run_sum"
            )
            _link(sinks, target, background, "background_threshold")
    for locator, row in _rows(directory, "txt/parameters.txt", sinks):
        sinks["source_assertions"].add(
            {
                "assertion_id": locator,
                "subject_id": SOURCE_ID,
                "predicate": "search_parameter:" + row["Parameter"],
                "value_raw": row["Value"],
                "normalized_value": None,
                "assertion_status": "source_reported",
                "source_locator": locator,
                "scope": "MaxQuant_search",
            }
        )
    return (
        peptides,
        conditions,
        {
            "n_quarantined_groups": quarantined,
            "n_filtered_zeros_with_positive_evidence": filtered_zeros,
            "maximum_absolute_filter_replay_error": maximum_error,
        },
    )


def _write_publisher(directory, sinks, peptides, conditions):
    import openpyxl

    by_sequence = {r["Sequence"]: pid for pid, r in peptides.items()}
    numerical = Counter()
    for asset in ("publisher/suppdata2.xlsx", "publisher/suppdata3.xlsx"):
        workbook = openpyxl.load_workbook(
            directory / "assets" / asset, read_only=True, data_only=False
        )
        try:
            for sheet in workbook:
                for row in sheet:
                    for cell in row:
                        if cell.value is None:
                            continue
                        value = cell.value
                        numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
                        sinks["publisher_cells"].add(
                            {
                                "cell_id": f"{asset}:{sheet.title}:{cell.coordinate}",
                                "asset_path": asset,
                                "sheet_name": sheet.title,
                                "cell_ref": cell.coordinate,
                                "row_1based": cell.row,
                                "column_1based": cell.column,
                                "excel_type": cell.data_type,
                                "value_text": str(value),
                                "value_number": float(value) if numeric else None,
                            }
                        )
                if asset.endswith("suppdata2.xlsx"):
                    if sheet.title != "Sheet1":
                        raise ValueError("Unknown publisher sheet")
                    headers = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
                    for index, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
                        if row[0] is None:
                            continue
                        if row[0] not in by_sequence:
                            raise ValueError(
                                "Publisher sequence absent from original peptide table"
                            )
                        pid = by_sequence[row[0]]
                        for column, condition in ((2, "UT"), (3, "TNF+IFN"), (4, None), (5, None)):
                            value = row[column]
                            if value is not None and (
                                not isinstance(value, (int, float)) or not math.isfinite(value)
                            ):
                                raise ValueError("Invalid publisher intensity")
                            source = f"{asset}:Sheet1:{chr(65 + column)}{index}"
                            replay = None
                            if condition is not None and value is not None:
                                vals = [
                                    float(
                                        peptides[pid][
                                            "Intensity " + e.replace("T+I_", "T+I-MAPPbeadFilt_")
                                        ]
                                    )
                                    for e in conditions[condition]
                                ]
                                median = statistics.median(vals)
                                if median <= 0 or not math.isclose(
                                    10**value, median, rel_tol=1e-4, abs_tol=1
                                ):
                                    raise ValueError("Publisher log-median does not replay")
                                replay = math.log10(median)
                                numerical[condition] += 1
                            sinks["publisher_quantities"].add(
                                {
                                    "measurement_id": source,
                                    "peptide_row_id": _peptide_id(pid),
                                    "source_cell": source,
                                    "representation": "log10_median_filtered"
                                    if condition
                                    else "imputed_header_conflict",
                                    "condition_header_raw": str(headers[column]),
                                    "condition": condition,
                                    "value_status": "unreported"
                                    if value is None
                                    else (
                                        "reported"
                                        if condition
                                        else "imputed_conflicting_condition_header"
                                    ),
                                    "unit": "log10_arbitrary_intensity"
                                    if condition
                                    else "arbitrary_intensity",
                                    "value": value,
                                    "recomputed_value": replay,
                                    "absolute_replay_error": abs(value - replay)
                                    if replay is not None
                                    else None,
                                    "derivation_rule": "log10_median_three_filtered_runs_including_zeros"
                                    if condition
                                    else None,
                                }
                            )
                            if condition is not None:
                                for experiment in conditions[condition]:
                                    _link(
                                        sinks,
                                        source,
                                        _result_id(pid, experiment) + ":bead_filtered",
                                        "median_run",
                                    )
        finally:
            workbook.close()
    return {
        "n_publisher_UT_numeric": numerical["UT"],
        "n_publisher_TNF_IFN_numeric": numerical["TNF+IFN"],
    }


def write_mapp_tables(directory):
    directory = Path(directory)
    typed = directory / "typed"
    typed.mkdir()
    sinks = {name: _Sink(typed, name) for name in SCHEMAS}
    try:
        peptides, conditions, counts = _write_tables(directory, sinks)
        counts.update(_write_publisher(directory, sinks, peptides, conditions))
        contract = json.loads((directory / "contracts/contract.json").read_text())
        for index, row in enumerate(contract["assertions"]):
            sinks["source_assertions"].add({"assertion_id": f"{SOURCE_ID}:curated:{index}", **row})
    finally:
        for sink in sinks.values():
            sink.close()
    return {**counts, **{f"n_{name}_rows": sink.n_rows for name, sink in sinks.items()}}


class MAPPAdapter:
    source_id = SOURCE_ID
    schema_version = 2
    asset_receipt_path = ASSET_RECEIPT
    support_files = SUPPORT_FILES

    def write(self, source_dir):
        source_dir = Path(source_dir)
        contract = json.loads((source_dir / "contracts/contract.json").read_text())
        receipt = source_dir / "contracts/source_assets.json"
        if hashlib.sha256(receipt.read_bytes()).hexdigest() != contract["source_receipt_sha256"]:
            raise ValueError("MAPP contract receipt hash changed")
        from .sequence_reference_evidence import write_reference_relations

        reconciliation = write_mapp_tables(source_dir)
        peptides = pq.read_table(source_dir / "typed/peptide_rows.parquet").to_pylist()
        reconciliation.update(
            write_reference_relations(source_dir, contract["references"], peptides)
        )
        return reconciliation

    def verify(self, source_dir, reconciliation):
        from .sequence_reference_evidence import verify_reference_relations

        verify_reference_relations(source_dir, reconciliation)
        if reconciliation["reference_relation_rows"]["peptide_reference_occurrences"] != 37721:
            raise ValueError("Pinned MAPP reference occurrence count changed")
        if reconciliation["reference_relation_rows"]["source_parent_claims"] != 64538:
            raise ValueError("Pinned MAPP source parent claim count changed")
        for name, schema in SCHEMAS.items():
            parquet = pq.ParquetFile(Path(source_dir) / "typed" / f"{name}.parquet")
            if (
                parquet.schema_arrow != schema
                or parquet.metadata.num_rows != reconciliation[f"n_{name}_rows"]
            ):
                raise ValueError(f"MAPP typed relation mismatch: {name}")
        expected = {
            "n_peptide_rows_rows": 10673,
            "n_evidence_rows": 149959,
            "n_channels_rows": 18,
            "n_quantity_results_rows": 192114,
            "n_quarantined_groups": 36,
            "n_filtered_zeros_with_positive_evidence": 9276,
            "n_publisher_UT_numeric": 4983,
            "n_publisher_TNF_IFN_numeric": 5311,
        }
        for key, value in expected.items():
            if reconciliation.get(key) != value:
                raise ValueError(f"Pinned MAPP source count changed: {key}")

    def registry(self, reconciliation):
        return {
            "source_id": SOURCE_ID,
            "dataset_accession": "PXD009092",
            "organism_scope": "human_HEK293_measured_cells",
            "measured_pools": ["proteasome_associated_fragment"],
            "measurement_grain": "peptide_by_source_acquisition_run",
            "training_admission": "deferred_pending_endpoint_recipe_and_parent_alternative_policy",
            "quantity_artifact": "typed/quantity_measurements.parquet",
            "reference_occurrence_artifact": "typed/peptide_reference_occurrences.parquet",
            "reference_ids": [r["reference_id"] for r in reconciliation["reference_counts"]],
            "rights": {
                "PRIDE_deposit": "EBI terms of use",
                "publisher_workbooks": "unverified_not_bundled_in_distribution",
            },
            "n_primary_peptide_rows": reconciliation["n_peptide_rows_rows"],
            "limitations": [
                "captured_pool_not_free_concentration",
                "unresolved_specimens",
                "source_FDR_conflict",
                "search_scope_6_to_40",
                "imputed_header_conflict",
                "parallel_tryptic_specimen_and_condition_map_unresolved",
            ],
        }
