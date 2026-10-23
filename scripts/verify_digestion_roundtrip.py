"""Independently compare a Soh 2024 package with original CSV and RData exports.

The decoded-root inputs are produced separately from the original RData by the
pinned Presto source audit. This script intentionally imports no Hitlist code.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import io
import json
import math
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
import pyarrow.parquet as pq

SOURCE_ID = "soh2024_digestion"
MS1_UNIT = "source_MS1_peak_area_arbitrary_units"
TABLES = {
    "QUANTITIES_raw.RData": "raw_quantities",
    "QUANTITIES_filtered.RData": "filtered_vectors",
    "MS1_HPR.csv": "ms1",
    "ASSIGNMENTS.csv": "identifications",
    "finalKinetics.csv": "final_vectors",
}


def rows(path, columns):
    parquet = pq.ParquetFile(path)
    for batch in parquet.iter_batches(batch_size=4096, columns=columns):
        yield from batch.to_pylist()


def source_number(value):
    if value in ("", "NA", "#N/A", "NaN", "nan"):
        return None, "source_missing"
    number = float(value)
    if not math.isfinite(number):
        return None, "source_nonfinite"
    if number == 0:
        return number, "source_zero"
    return number, "source_positive" if number > 0 else "source_negative"


def assert_original_rows(assets, decoded_root, typed):
    readers = {
        name: rows(typed / f"{name}.parquet", ["source_locator", "fields_json"])
        for name in TABLES.values()
    }
    member_rows = rows(
        typed / "members.parquet",
        [
            "source_locator",
            "fields_json",
            "member_size_bytes",
            "member_crc32",
            "materialization_status",
        ],
    )
    counts = Counter()
    with zipfile.ZipFile(assets / "aSPIRE_outputs_zip.zip") as outer:
        for member in outer.infolist():
            if member.is_dir():
                continue
            condition = Path(member.filename).stem
            with zipfile.ZipFile(io.BytesIO(outer.read(member))) as inner:
                if inner.testzip() is not None:
                    raise AssertionError((condition, "nested archive CRC failure"))
                for entry in inner.infolist():
                    locator = f"aSPIRE_outputs_zip.zip!{member.filename}!{entry.filename}"
                    archived = next(member_rows)
                    if (
                        archived["source_locator"] != locator
                        or archived["member_size_bytes"] != entry.file_size
                        or archived["member_crc32"] != entry.CRC
                    ):
                        raise AssertionError((locator, "archive member metadata differs"))
                    original_info = json.loads(archived["fields_json"])
                    if original_info != {
                        "filename": entry.filename,
                        "file_size": entry.file_size,
                        "compress_size": entry.compress_size,
                        "CRC": entry.CRC,
                        "date_time": list(entry.date_time),
                        "external_attr": entry.external_attr,
                    }:
                        raise AssertionError((locator, "archive member fields differ"))
                    base = Path(entry.filename).name
                    table = next(
                        (value for suffix, value in TABLES.items() if base.endswith(suffix)), None
                    )
                    status = (
                        "archive_directory"
                        if entry.is_dir()
                        else f"typed_{table}"
                        if table
                        else "original_asset_only"
                    )
                    if archived["materialization_status"] != status:
                        raise AssertionError((locator, "member materialization status differs"))
                    counts["members"] += 1
                    if entry.is_dir():
                        continue
                    if table is None:
                        continue
                    with (
                        gzip.open(
                            decoded_root / condition / base.replace(".RData", ".csv.gz"),
                            "rt",
                            newline="",
                        )
                        if base.endswith(".RData")
                        else io.TextIOWrapper(inner.open(entry), newline="")
                    ) as handle:
                        for number, original in enumerate(csv.DictReader(handle), 2):
                            actual = next(readers[table])
                            if actual["source_locator"] != f"{locator}:row:{number}":
                                raise AssertionError(
                                    (table, counts[table], "source locator differs")
                                )
                            if actual["fields_json"] != json.dumps(original, sort_keys=True):
                                raise AssertionError(
                                    (table, counts[table], "original fields differ")
                                )
                            counts[table] += 1
    for table, reader in readers.items():
        if next(reader, None) is not None:
            raise AssertionError((table, "extra packaged row"))
    if next(member_rows, None) is not None:
        raise AssertionError("Extra packaged archive member")
    return dict(counts)


def assert_workbook(assets, typed):
    actual = rows(typed / "metadata.parquet", ["source_locator", "fields_json"])
    counts = Counter()
    book = openpyxl.load_workbook(assets / "Identifier.xlsx", read_only=True, data_only=False)
    try:
        for sheet in book:
            source = sheet.iter_rows(values_only=True)
            header = list(next(source))
            for number, values in enumerate(source, 2):
                row = next(actual)
                if row["source_locator"] != f"Identifier.xlsx:{sheet.title}:row:{number}":
                    raise AssertionError((sheet.title, number, "locator differs"))
                if json.loads(row["fields_json"]) != {
                    "columns": header,
                    "values": list(values),
                }:
                    raise AssertionError((sheet.title, number, "workbook cells differ"))
                counts[sheet.title] += 1
    finally:
        book.close()
    if next(actual, None) is not None:
        raise AssertionError("Extra packaged workbook row")
    return dict(counts)


def assert_run_links(assets, typed):
    book = openpyxl.load_workbook(assets / "Identifier.xlsx", read_only=True, data_only=False)
    try:
        sheet = book["Main_searches"]
        source = sheet.iter_rows(values_only=True)
        header = next(source)
        runs = {
            row["raw_file"].removesuffix(".raw"): row
            for row in (dict(zip(header, values)) for values in source)
        }
    finally:
        book.close()
    seen = set()
    for row in rows(
        typed / "runs.parquet", ["run_id", "source_name", "time_hours", "instrument_raw"]
    ):
        original = runs[row["source_name"]]
        if (
            row["run_id"] != f"{SOURCE_ID}:run:{original['raw_file']}"
            or row["time_hours"] != float(original["digestTime"])
            or row["instrument_raw"] != original["Mass spectrometry"]
        ):
            raise AssertionError((row["source_name"], "main-search run differs"))
        seen.add(row["source_name"])
    if seen != set(runs):
        raise AssertionError("Main-search acquisition coverage differs")
    conflicts = 0
    for row in rows(
        typed / "raw_quantities.parquet",
        [
            "run_id",
            "source_name",
            "quantity_time_hours",
            "metadata_time_hours",
            "time_status",
            "measurement_unit",
            "quantity_scale",
            "fields_json",
            "peptide_sequence",
            "intensity",
            "value_status",
        ],
    ):
        source = json.loads(row["fields_json"])
        original = runs[row["source_name"]]
        differing = row["quantity_time_hours"] != float(original["digestTime"])
        intensity, value_status = source_number(source["intensity"])
        if (
            row["run_id"] != f"{SOURCE_ID}:run:{original['raw_file']}"
            or row["source_name"] != source["source"]
            or row["quantity_time_hours"] != float(source["digestTime"])
            or row["peptide_sequence"] != source["pepSeq"]
            or row["intensity"] != intensity
            or row["value_status"] != value_status
            or row["metadata_time_hours"] != float(original["digestTime"])
            or row["time_status"]
            != ("source_metadata_conflict" if differing else "source_agreement")
            or row["measurement_unit"] != MS1_UNIT
            or row["quantity_scale"] != "aSPIRE_raw_peptide_run_quantity"
        ):
            raise AssertionError((row["source_name"], "raw quantity run/time differs"))
        conflicts += differing
    statuses = Counter()
    for row in rows(
        typed / "ms1.parquet",
        [
            "run_id",
            "source_name",
            "run_link_status",
            "measurement_unit",
            "quantity_scale",
            "fields_json",
            "peptide_sequence",
            "modified_sequence",
            "ms1_area",
            "ms1_background",
            "value_status",
        ],
    ):
        name = row["source_name"]
        original = runs.get(name)
        source = json.loads(row["fields_json"])
        area, value_status = source_number(source["TotalAreaMs1"])
        background, _ = source_number(source["TotalBackgroundMs1"])
        status = (
            "matched_main_search"
            if original
            else "blank_original_source_name"
            if not name
            else "outside_main_search_roster"
        )
        if (
            row["run_link_status"] != status
            or name != source["FileName"].removesuffix(".raw")
            or row["peptide_sequence"] != source["PeptideSequence"]
            or row["modified_sequence"] != source["PeptideModifiedSequence"]
            or row["ms1_area"] != area
            or row["ms1_background"] != background
            or row["value_status"] != value_status
            or row["run_id"] != (f"{SOURCE_ID}:run:{original['raw_file']}" if original else None)
            or row["measurement_unit"] != MS1_UNIT
            or row["quantity_scale"] != "Skyline_MS1_peak_area"
        ):
            raise AssertionError((name, "Skyline run link differs"))
        statuses[status] += 1
    return {
        "n_main_search_runs": len(runs),
        "n_quantity_time_conflicts": conflicts,
        "ms1_run_link_status": dict(statuses),
    }


def assert_constructs(assets, typed):
    actual = {
        row["source_locator"]: row
        for row in rows(
            typed / "constructs.parquet",
            ["source_locator", "header_raw", "sequence", "sequence_sha256"],
        )
    }
    n_constructs = 0
    for path in (assets / "substrate-fastas").glob("*.fasta"):
        lines = path.read_text().splitlines()
        sequence = "".join(line.strip() for line in lines[1:])
        row = actual.pop(f"substrate-fastas/{path.name}")
        if (row["header_raw"], row["sequence"], row["sequence_sha256"]) != (
            lines[0],
            sequence,
            hashlib.sha256(sequence.encode()).hexdigest(),
        ):
            raise AssertionError((path.name, "construct differs"))
        n_constructs += 1
    if actual:
        raise AssertionError(("extra packaged constructs", sorted(actual)))
    return n_constructs


def assert_final_slots(typed):
    values = rows(
        typed / "final_values.parquet",
        [
            "source_record_id",
            "vector_id",
            "slot_0based",
            "time_hours",
            "intensity",
            "value_raw",
            "value_status",
            "time_status",
            "measurement_unit",
            "quantity_scale",
        ],
    )
    counts = Counter()
    for vector in rows(
        typed / "final_vectors.parquet",
        ["source_record_id", "fields_json", "row_status"],
    ):
        source = json.loads(vector["fields_json"])
        if source["biological_replicate"] == "NA":
            if vector["row_status"] != "annotation_only":
                raise AssertionError("Annotation-only vector misclassified")
            counts["n_annotation_only_rows"] += 1
            continue
        if vector["row_status"] != "reported_vector":
            raise AssertionError("Quantified vector misclassified")
        times = source["digestTimes"].split(";")
        intensities = source["intensities"].split(";")
        if len(times) != len(intensities):
            raise AssertionError("Original vector lengths differ")
        repeated = len(times) != len(set(times))
        for slot, (time_raw, intensity_raw) in enumerate(zip(times, intensities)):
            row = next(values)
            if (row["vector_id"], row["source_record_id"], row["slot_0based"]) != (
                vector["source_record_id"],
                f"{vector['source_record_id']}:slot:{slot}",
                slot,
            ):
                raise AssertionError((vector["source_record_id"], slot, "slot identity differs"))
            if row["time_hours"] != float(time_raw) or row["value_raw"] != intensity_raw:
                raise AssertionError((vector["source_record_id"], slot, "original slot differs"))
            if row["time_status"] != (
                "repeated_source_time_slots" if repeated else "source_reported"
            ):
                raise AssertionError((vector["source_record_id"], slot, "time status differs"))
            if (
                row["measurement_unit"] != MS1_UNIT
                or row["quantity_scale"] != "aSPIRE_background_adjusted_biological_mean"
            ):
                raise AssertionError((vector["source_record_id"], slot, "quantity scale differs"))
            expected = float(intensity_raw)
            if math.isnan(expected):
                if row["intensity"] is not None or row["value_status"] != (
                    "postprocessing_missing_unresolved_contributors"
                ):
                    raise AssertionError((vector["source_record_id"], slot, "missing differs"))
            elif row["intensity"] != expected or row["value_status"] != (
                "postprocessing_forced_zero"
                if expected == 0 and row["time_hours"] == 0
                else "source_zero"
                if expected == 0
                else "source_positive"
            ):
                raise AssertionError((vector["source_record_id"], slot, "numeric/status differs"))
            counts[row["value_status"]] += 1
        counts["n_reported_vectors"] += 1
    if next(values, None) is not None:
        raise AssertionError("Extra packaged final value")
    return dict(counts)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("package", type=Path)
    parser.add_argument("original_assets", type=Path)
    parser.add_argument("decoded_rdata", type=Path)
    args = parser.parse_args()
    typed = args.package / "external" / SOURCE_ID / "typed"
    result = {
        "source_id": SOURCE_ID,
        "original_row_roundtrip": assert_original_rows(
            args.original_assets, args.decoded_rdata, typed
        ),
        "workbook_rows": assert_workbook(args.original_assets, typed),
        "run_links": assert_run_links(args.original_assets, typed),
        "n_constructs": assert_constructs(args.original_assets, typed),
        "final_slot_states": assert_final_slots(typed),
    }
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
