"""Original-shaped synthetic workbooks exercise offline canine source lineage."""

import copy
import hashlib
import json
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
import yaml

from hitlist import canine_ligands
from hitlist.curation import detect_monoallelic, load_pmid_overrides
from hitlist.export import generate_observations_table
from hitlist.lineage import attach_lineage
from hitlist.provenance import ContributorCollector, file_digest
from hitlist.supplement import scan_supplementary


@pytest.fixture
def sources(tmp_path, monkeypatch):
    profile = copy.deepcopy(yaml.safe_load(canine_ligands._PROFILE_PATH.read_text()))
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    for asset in profile["assets"]:
        book = openpyxl.Workbook()
        book.remove(book.active)
        for arm in asset["sheets"]:
            arm["n_observations"] = 1
            sheet = book.create_sheet(arm["sheet"])
            sheet.append(["Synthetic original-shaped observation fixture"])
            sheet.append(asset["header"])
            sheet.append(["ACDEFGHIK", 9, "protein-a:protein-b", "DB Search"])
        other = book.create_sheet("Human TAA comparison")
        other.append(["AAAAAAAAA", "This is deliberately not an observation worksheet"])
        path = inputs / asset["file"]
        book.save(path)
        book.close()
        asset.update(file_digest(path), md5=hashlib.md5(path.read_bytes()).hexdigest())
        asset["url"] = "fixture:" + asset["file"]
    profile_path = tmp_path / "profile.yaml"
    profile_path.write_text(yaml.safe_dump(profile))
    monkeypatch.setattr(canine_ligands, "_PROFILE_PATH", profile_path)

    def forbidden(*args, **kwargs):
        pytest.fail("Offline canine curation attempted a network call")

    monkeypatch.setattr("requests.get", forbidden)
    monkeypatch.setattr("hitlist.downloads.fetch_data_asset", forbidden)
    return inputs, profile_path


def _scan(inputs, output, collector=None):
    manifest = canine_ligands.curate_canine_ligands(inputs, output)
    return scan_supplementary(
        entries=json.loads(manifest.read_text()),
        directory=manifest.parent,
        allow_download=False,
        provenance=collector,
    )


def test_original_sheets_keep_species_restrictions_and_contributors(sources, tmp_path):
    inputs, _ = sources
    output = tmp_path / "curated"
    with ContributorCollector(scratch_dir=tmp_path) as collector:
        frame = _scan(inputs, output, collector)
        collector.write([frame], tmp_path / "contributors.parquet")
    assert len(frame) == 8
    assert set(frame.peptide) == {"ACDEFGHIK"}
    human = frame[frame.host.eq("Homo sapiens")]
    dogs = frame[frame.host.eq("Canis lupus familiaris")]
    assert len(human) == 3 and len(dogs) == 5
    assert human.source_organism.eq("Homo sapiens").all()
    assert human.species.eq("Homo sapiens").all()
    assert frame.mhc_species.str.startswith("Canis").all()
    assert human.is_monoallelic.all()
    assert human.restriction_evidence.eq("monoallelic").all()
    assert not dogs.is_monoallelic.any()
    assert dogs.restriction_evidence.eq("unknown").all()
    assert dogs.mhc_allele_set.eq("").all()
    contributors = pd.read_parquet(tmp_path / "contributors.parquet")
    assert len(contributors) == 8
    assert set(contributors.attributed_sample_label) == set(frame.attributed_sample_label)
    rows = [json.loads(value) for value in contributors.original_fields]
    for row in rows:
        assert row["source_row"] == "3"
        assert row["source_protein_mappings"] == "protein-a:protein-b"
        assert row["source_sha256"] == row["manifest"]["source_asset"]["sha256"]
        assert row["source_url"].startswith("fixture:")
        assert row["il_ambiguity"] == "unresolved"
        assert row["peptide_q_value"] == row["spectrum_id"] == ""
    assert sum("AA21 P > L" in row["construct_description"] for row in rows) == 1


def test_exact_sample_join_preserves_four_dogs_and_shared_lola_specimen(
    sources, tmp_path, monkeypatch
):
    frame = _scan(sources[0], tmp_path / "curated")
    frame["source"] = "supplement"
    monkeypatch.setattr("hitlist.observations.load_observations", lambda **kwargs: frame.copy())
    result = generate_observations_table(exclude_non_peptide_ligand=False)
    assert result.sample_attribution.eq("curated_sample_label").all()
    assert result.condition_id.nunique() == 8
    result["evidence_kind"] = "ms"
    result = attach_lineage(result)
    dogs = result[result.donor_status.eq("resolved")]
    assert len(dogs) == 5 and dogs.donor_ids.nunique() == 4
    lola = dogs[dogs.sample_label.str.startswith("Lola")]
    assert len(lola) == 2
    assert lola.donor_ids.nunique() == lola.specimen_ids.nunique() == 1
    assert lola.experimental_origin_ids.nunique() == 2
    assert result.acquisition_status.eq("unknown").all()
    human = result[result.sample_label.str.startswith("HCT116")]
    assert human.donor_status.eq("unknown").all()
    assert human.presenting_species.eq("Homo sapiens").all()
    assert human.introduced_mhc_species.eq("Canis lupus familiaris").all()
    assert human.species_context_kind.eq("mhc_transfectant").all()
    assert dogs.species_context_kind.eq("native").all()
    assert dogs.restriction_evidence.eq("unknown").all()


def test_curation_is_deterministic_and_never_replaces_an_export(sources, tmp_path):
    inputs, _ = sources
    first, second = tmp_path / "first", tmp_path / "second"
    canine_ligands.curate_canine_ligands(inputs, first)
    canine_ligands.curate_canine_ligands(inputs, second)
    assert {p.name: p.read_bytes() for p in first.iterdir()} == {
        p.name: p.read_bytes() for p in second.iterdir()
    }
    with pytest.raises(FileExistsError):
        canine_ligands.curate_canine_ligands(inputs, first)


@pytest.mark.parametrize("failure", ["missing", "checksum", "header", "row", "count"])
def test_invalid_sources_publish_nothing(sources, tmp_path, failure):
    inputs, profile_path = sources
    profile = yaml.safe_load(profile_path.read_text())
    asset = profile["assets"][0]
    path = inputs / asset["file"]
    if failure == "missing":
        path.unlink()
    elif failure == "checksum":
        path.write_bytes(path.read_bytes() + b"changed")
    elif failure == "count":
        asset["sheets"][0]["n_observations"] += 1
    else:
        book = openpyxl.load_workbook(path)
        sheet = book[asset["sheets"][0]["sheet"]]
        sheet["A2" if failure == "header" else "A3"] = "invalid"
        book.save(path)
        book.close()
        asset.update(file_digest(path), md5=hashlib.md5(path.read_bytes()).hexdigest())
    profile_path.write_text(yaml.safe_dump(profile))
    with pytest.raises((ValueError, FileNotFoundError)):
        canine_ligands.curate_canine_ligands(inputs, tmp_path / "output")
    assert not (tmp_path / "output").exists()
    assert not list(tmp_path.glob(".output-*"))


def test_only_engineered_hct116_variant_is_a_null_host():
    assert not detect_monoallelic("HCT116", "DLA-88*003:02")[0]
    assert detect_monoallelic("HCT116 HLA-I KO", "DLA-88*003:02")[0]
    study = load_pmid_overrides()[42199926]
    assert study["species"] == ""  # Mixed source proteomes: no paper-wide fill.
    assert not study.get("hla_alleles")  # No transductant allele pool on tumor rows.


def test_profile_pins_every_original_source_and_observation_count():
    profile = yaml.safe_load(
        (Path(canine_ligands.__file__).parent / "data" / "canine_ligands.yaml").read_text()
    )
    arms = [arm for asset in profile["assets"] for arm in asset["sheets"]]
    assert len(profile["assets"]) == 6 and len(arms) == 8
    assert sum(arm["n_observations"] for arm in arms) == 16359
    assert sum(arm["n_observations"] for arm in arms if arm["donor_label"]) == 12779
    assert len({arm["donor_label"] for arm in arms} - {""}) == 4
    assert all(asset["md5"] and asset["sha256"] for asset in profile["assets"])
