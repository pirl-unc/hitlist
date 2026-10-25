"""Source identity, literal workbook semantics, unresolved context and replay."""

import csv
import io
import json
import zipfile

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from hitlist import goldberg_mapp as source
from hitlist.evidence_catalog import FIELDS, GOLDBERG_RELATIONS
from hitlist.evidence_package import package_family
from hitlist.evidence_store import EvidenceStore
from hitlist.external_evidence import _copy_source
from hitlist.goldberg_workbooks import workbook_rows
from hitlist.material_context import links_with_minimum_precision
from hitlist.provenance import file_digest


def tsv(rows):
    output = io.StringIO()
    writer = csv.DictWriter(output, list(rows[0]), delimiter="\t", lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue()


def workbook(sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.fixture
def original(tmp_path, monkeypatch):
    contract = json.loads(source.CONTRACT.read_text())
    directory = tmp_path / "original"
    directory.mkdir()
    run1, run2 = "sourceA_Ctrl_1", "sourceB_Ctrl4_1"
    (directory / contract["deposition"]).write_bytes(
        workbook(
            {
                "Sheet1": [
                    ["Raw File", "Experiment", "Sample Name", "Cell Type", "Comments"],
                    ["rawA", "EC infection Degradome MAPP Peptides", "Ctrl_1", "A549", None],
                    ["rawB", "EC infection Degradome MAPP Peptides", "Ctrl4_1", "A549", None],
                ]
            }
        )
    )
    (directory / contract["sdrf"]).write_text(
        tsv(
            [
                {"comment[data file]": "rawA.raw", "comment[cleavage agent details]": "trypsin"},
                {
                    "comment[data file]": "unmatched.raw",
                    "comment[cleavage agent details]": "trypsin",
                },
            ]
        )
    )
    (directory / contract["project_metadata"]).write_text('{"species": "Homo sapiens"}')
    (directory / contract["file_metadata"]).write_text('[{"fileName":"rawA.raw"}]')
    peptides = []
    for native, parents, first, second in [
        ("0", "P1", "10", "0"),
        ("1", "P2", "", "4"),
        ("2", "P1;P2", "NaN", "invalid"),
    ]:
        peptides.append(
            {
                "id": native,
                "Sequence": "ACDEFGH",
                "Proteins": parents,
                "Protein group IDs": "0",
                "Reverse": "",
                "Potential contaminant": "",
                "PEP": "0.001",
                "Evidence IDs": "0" if native == "0" else "",
                "Mod. peptide IDs": "0" if native == "0" else "",
                "Start position": "1",
                "End position": "7",
                "Amino acid before": "-",
                "Amino acid after": "R",
                "Intensity": first,
                "Intensity " + run1: first,
                "Intensity " + run2: second,
                "Identification type " + run1: "By MS/MS",
                "Identification type " + run2: "By matching",
            }
        )
    form = {
        "id": "0",
        "Peptide ID": "0",
        "Sequence": "ACDEFGH",
        "Proteins": "P1",
        "Protein group IDs": "0",
        "Reverse": "",
        "Potential contaminant": "",
        "PEP": "0.001",
        "Modifications": "Unmodified",
        "Evidence IDs": "0",
        "Intensity": "8",
        "Intensity " + run1: "8",
        "Intensity " + run2: "0",
    }
    evidence = {
        "id": "0",
        "Peptide ID": "0",
        "Mod. peptide ID": "0",
        "Sequence": "ACDEFGH",
        "Proteins": "P1",
        "Protein group IDs": "0",
        "Reverse": "",
        "Potential contaminant": "",
        "PEP": "0.001",
        "Raw file": "rawA",
        "Experiment": run1,
        "Type": "MULTI-MSMS",
        "Modifications": "Unmodified",
        "Modified sequence": "_ACDEFGH_",
        "Intensity": "7",
    }
    with zipfile.ZipFile(directory / contract["archive"], "w") as archive:
        archive.writestr(
            "parameters.txt",
            tsv(
                [
                    {"Parameter": "Fasta file", "Value": "human_plus_other.fasta"},
                    {"Parameter": "Fasta file", "Value": "second.fasta"},
                ]
            ),
        )
        archive.writestr(
            "summary.txt",
            tsv(
                [
                    {"Raw file": "rawA", "Experiment": run1, "Enzyme mode": "Unspecific"},
                    {"Raw file": "rawB", "Experiment": run2, "Enzyme mode": "Unspecific"},
                    {"Raw file": run1, "Experiment": "", "Enzyme mode": ""},
                ]
            ),
        )
        archive.writestr(
            "proteinGroups.txt",
            tsv(
                [
                    {
                        "id": "0",
                        "Protein IDs": "P1;P2",
                        "Fasta headers": "sp|P1|HOST Protein OS=Homo sapiens OX=9606 GN=A;sp|P2|HOST Other OS=Homo sapiens OX=960",
                        "Peptide IDs": "0;1;2",
                        "Evidence IDs": "0",
                        "Reverse": "",
                        "Potential contaminant": "",
                    }
                ]
            ),
        )
        archive.writestr("peptides.txt", tsv(peptides))
        archive.writestr("modificationSpecificPeptides.txt", tsv([form]))
        archive.writestr("evidence.txt", tsv([evidence]))
    with zipfile.ZipFile(directory / contract["publisher"], "w") as archive:
        archive.writestr(
            contract["publisher_workbooks"][0]["member"],
            workbook(
                {
                    "MAPP database": [
                        [
                            "Sequence",
                            "Leading razor protein",
                            "Start position",
                            "End position",
                            "Gene names",
                            "Treatment",
                        ],
                        ["ACDEFGH", "P1", 1, 7, 41000, "repeat:3;repeat:0;unknown:bad"],
                    ]
                }
            ),
        )
        archive.writestr(
            contract["publisher_workbooks"][1]["member"],
            workbook(
                {
                    "MAPP_peptidomics": [
                        [
                            "Sequence",
                            "Proteins",
                            "id",
                            "Intensity " + run1,
                            "Intensity renamed_run",
                            "Intensity " + run2,
                        ],
                        ["ACDEFGH", "P1", 0, 10, 5, 0],
                    ],
                    "Shotgun_proteomics": [
                        ["Protein IDs", "LFQ intensity unrelated_run"],
                        ["P1", 1000],
                    ],
                }
            ),
        )
    for asset in contract["assets"]:
        asset.update(file_digest(directory / asset["filename"]))
    path = tmp_path / "contract.json"
    path.write_text(json.dumps(contract))
    monkeypatch.setattr(source, "CONTRACT", path)
    return directory


def tables(package):
    return {
        name: pq.read_table(package / source.ROOT / "typed" / f"{name}.parquet").to_pylist()
        for name in source.SCHEMAS
    }


def test_external_adapter_receipt_matches_the_standalone_contract():
    contract = json.loads(source.CONTRACT.read_text())
    receipt = json.loads(source.ASSET_RECEIPT.read_text())
    assert receipt == {
        "source_id": "goldberg2025",
        "assets": {
            asset["filename"]: {"sha256": asset["sha256"], "size": asset["size_bytes"]}
            for asset in contract["assets"]
        },
    }


def test_external_adapter_replays_the_same_source_faithful_relations(
    original, tmp_path, monkeypatch
):
    contract = json.loads(source.CONTRACT.read_text())
    receipt = tmp_path / "receipt.json"
    receipt.write_text(
        json.dumps(
            {
                "source_id": "goldberg2025",
                "assets": {
                    asset["filename"]: {
                        "sha256": asset["sha256"],
                        "size": asset["size_bytes"],
                    }
                    for asset in contract["assets"]
                },
            }
        )
    )
    monkeypatch.setattr(source.GoldbergMAPPAdapter, "asset_receipt_path", receipt)
    monkeypatch.setattr(
        source.GoldbergMAPPAdapter, "support_files", {"contract.json": source.CONTRACT}
    )
    adapter = source.GoldbergMAPPAdapter()
    external = tmp_path / "external"
    _copy_source(adapter, original, external)
    reconciliation = adapter.write(external)
    adapter.verify(external, reconciliation)
    assert reconciliation["n_peptides_rows"] == 3
    assert reconciliation["n_quantities_rows"] == 13
    assert adapter.registry(reconciliation)["dataset_accession"] == "PXD052314"

    path = external / "typed/quantities.parquet"
    rows = pq.read_table(path).to_pylist()
    rows[0]["quantitative_value"] = 999
    pq.write_table(pa.Table.from_pylist(rows, schema=source.SCHEMAS["quantities"]), path)
    with pytest.raises(ValueError, match="typed rows differ"):
        adapter.verify(external, reconciliation)


def test_original_identity_grains_and_ambiguous_material(original, tmp_path):
    package = tmp_path / "package"
    source.write_goldberg_mapp_package(original, package)
    store = EvidenceStore.open(package)
    assert package_family(store.manifest) == "goldberg_mapp"
    assert set(store.bindings) == {"goldberg_" + name for name in source.SCHEMAS}
    assert store.catalog()["source_registry"]["goldberg2025"]["training_status"].startswith(
        "source_evidence_only"
    )
    data = tables(package)
    assert len(data["peptides"]) == 3
    assert len({r["peptide_id"] for r in data["peptides"]}) == 3
    assert {r["peptide_sequence"] for r in data["peptides"]} == {"ACDEFGH"}
    assert [r["source_organism_status"] for r in data["peptides"]] == [
        "source_human_only",
        "mixed_or_unresolved",
        "mixed_or_unresolved",
    ]
    assert data["references"][1]["organism_taxon_id"] is None
    assert all(r["reference_sequence"] is None for r in data["references"])
    assert len(data["quantities"]) == 13
    assert {r["value_status"] for r in data["quantities"]} == {
        "positive",
        "zero",
        "missing",
        "nonfinite",
        "invalid_literal",
    }
    assert {r["quantity_kind"] for r in data["quantities"]} == {
        "peptide_intensity",
        "form_intensity",
        "identification_intensity",
    }
    assert len(data["acquisitions"]) == 2
    assert not links_with_minimum_precision(data["material_links"], "exact_specimen")
    assert all(r["specimen_id"] is None and r["aliquot_id"] is None for r in data["acquisitions"])
    assert all(
        r["component_state_status"] == "unreported_not_assumed_wildtype"
        for r in data["acquisitions"]
    )
    assert any(r["assertion_status"] == "unmatched_secondary_raw_name" for r in data["assertions"])
    assert sum(r["assertion_predicate"] == "Fasta file" for r in data["assertions"]) == 2
    assert len(data["pooled_slots"]) == 3
    assert [r["sample_label_raw"] for r in data["pooled_slots"]] == ["repeat", "repeat", "unknown"]
    assert [r["quantitative_value"] for r in data["pooled_slots"]] == [3, 0, None]
    assert all(r["acquisition_id"] is None for r in data["pooled_slots"])
    assert {r["link_status"] for r in data["publisher_links"]} == {
        "exact_native_id_and_sequence",
        "exact_source_value_agreement",
        "unresolved_exact_column_name",
    }
    assert {r["measured_pool"] for r in data["publisher_quantities"]} == {
        "proteasome_associated_peptides",
        "whole_cell_digested_protein",
    }
    queried = store.query(
        "goldberg_quantities", select=["quantitative_value", "value_status"], limit=None
    ).table
    assert queried.num_rows == 13
    known = {f.name for f in FIELDS}
    assert all(set(fields) <= known for _, fields in GOLDBERG_RELATIONS.values())
    with pytest.raises(FileExistsError):
        source.write_goldberg_mapp_package(original, package)


def test_replay_rejects_tampering_even_with_rehashed_manifest(original, tmp_path):
    package = tmp_path / "package"
    source.write_goldberg_mapp_package(original, package)
    path = package / source.ROOT / "typed/quantities.parquet"
    data = pq.read_table(path).to_pylist()
    data[0]["quantitative_value"] = 999
    pq.write_table(pa.Table.from_pylist(data, schema=source.SCHEMAS["quantities"]), path)
    manifest_path = package / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["files"] = source._inventory(package)
    manifest_path.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="typed rows differ"):
        source.verify_goldberg_mapp_package(package)


def test_wrong_original_refused_before_compilation(original, tmp_path):
    contract = json.loads(source.CONTRACT.read_text())
    (original / contract["project_metadata"]).write_text("{}")
    with pytest.raises(ValueError, match="source mismatch"):
        source.write_goldberg_mapp_package(original, tmp_path / "package")
    assert not (tmp_path / "package").exists()


def test_candidate_recipe_separates_review_axes_and_rejects_rehashed_tamper(original, tmp_path):
    from hitlist.evidence_recipe import RecipeError, materialize_recipe, verify_recipe_package
    from hitlist.goldberg_candidates import RECIPE_ID

    package, output = tmp_path / "package", tmp_path / "candidates"
    source.write_goldberg_mapp_package(original, package)
    store = EvidenceStore.open(package)
    recipe = {"schema_version": 1, "name": "fixture", "registered_recipe": RECIPE_ID}
    materialize_recipe(store, recipe, output)
    manifest = verify_recipe_package(output, source_package=package)
    assert manifest["coverage"]["n_source_peptide_cells"] == 9
    assert manifest["coverage"]["n_baseline_cells"] == 6
    assert manifest["coverage"]["n_candidate_cells"] == 1
    assert manifest["coverage"]["n_excluded_baseline_cells"] == 5
    assert manifest["coverage"]["n_outside_scope_cells"] == 3
    candidates = pq.read_table(output / "typed/candidates.parquet").to_pylist()
    assert candidates[0]["quantitative_value"] == 10
    assert candidates[0]["log10_intensity"] == 1
    assert candidates[0]["confidence_policy_status"] == "retained_not_yet_selected"
    assert (
        candidates[0]["training_eligibility_status"] == "not_fully_qualified_processing_supervision"
    )
    assert candidates[0]["full_reference_status"] == "full_original_search_FASTA_absent"
    assert candidates[0]["precursor_pairing_status"] == "exact_physical_pairing_unestablished"
    assert candidates[0]["publisher_selection_status"] == "published_exact_identity"
    assert candidates[0]["source_evidence_ids_raw"] == "0"
    path = output / "typed/candidates.parquet"
    schema = pq.read_schema(path)
    candidates[0]["log10_intensity"] = 999
    pq.write_table(pa.Table.from_pylist(candidates, schema=schema), path)
    manifest["files"]["typed/candidates.parquet"] = file_digest(path)
    (output / "manifest.json").write_text(json.dumps(manifest))
    with pytest.raises(RecipeError, match="rows differ"):
        verify_recipe_package(output, source_package=package)


def test_excel_date_style_formula_and_sparse_coordinates_stay_literal():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["A1"] = "Gene names"
    sheet["A2"] = 41000
    sheet["A2"].number_format = "mm-dd-yy"
    sheet["C2"] = "=1+2"
    sheet["F3"] = "repeat:1;repeat:0"
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    rows = list(workbook_rows(buffer))
    assert [(r, [c["source_column_index"] for c in cells]) for _, r, cells in rows] == [
        (1, [1]),
        (2, [1, 3]),
        (3, [6]),
    ]
    date, formula = rows[1][2]
    assert date["cell_value_raw"] == date["cell_value_resolved"] == "41000"
    assert date["cell_style_raw"] is not None
    assert formula["cell_formula_raw"] == "1+2"
    assert formula["cell_value_raw"] == ""  # Original <v/> is present but empty.


@pytest.mark.parametrize(
    "literal,status",
    [
        ("2", "outside_probability_range"),
        ("-1", "outside_probability_range"),
        ("0", "reported_probability"),
        ("NaN", "nonfinite"),
    ],
)
def test_confidence_is_not_an_implicit_threshold(literal, status):
    assert source._confidence(literal)[1] == status
