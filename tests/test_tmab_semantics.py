"""Real-notation regressions and a two-paper source package for #791."""

import hashlib
import json

import openpyxl
import pyarrow.parquet as pq
import pytest

from hitlist import EvidenceStore
from hitlist import tmab_peptidomics as tmab


def test_tentative_acetylation_does_not_insert_an_alanine():
    assert tmab._parse_sequence("(Ac-AGKQAVSASGK)") == ("AGKQAVSASGK", "N_terminal_acetylation")


def test_oxidation_positions_survive_normalization():
    assert tmab._parse_sequence("Ac-MoxAGMox") == (
        "MAGM",
        "N_terminal_acetylation;oxidation@1;oxidation@4",
    )


@pytest.mark.parametrize(
    "raw",
    [
        "APIKVGDAIPAVE -18 (dehydro? - maps to Cterm E)",
        "pyroE-AEWDDYVPKLYEQLSGK",
        "ABC",
        "PEP(TIDE)",
    ],
)
def test_unresolved_annotation_does_not_become_a_confident_sequence(raw):
    assert tmab._parse_sequence(raw)[0] is None


def test_multiword_compounds_remain_distinct():
    assert tmab._treatment("Bestatin Methyl Ester 75min 100uM")[1] == "bestatin methyl ester"
    assert tmab._treatment("Bestatin 75min 10uM")[1] == "bestatin"


@pytest.fixture
def source_package(tmp_path, monkeypatch):
    sources = tmp_path / "inputs"
    sources.mkdir()
    specs = []
    for filename, config in {**tmab._PRIMARY, **tmab._HEATMAP}.items():
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = config["sheet"]
        source = "dasgupta_2014" if filename.startswith("dasgupta") else "gelman_2013"
        if filename in tmab._PRIMARY:
            sequences = [
                "LGDGPKLV",
                None,
                "(Ac-AGKQAVSASGK)",
                "Ac-MoxAGMox",
                "APIKVGDAIPAVE -18 (dehydro? - maps to Cterm E)",
                "PEPTIDE",
            ]
            for offset, sequence in enumerate(sequences):
                row = config["data_row"] + offset
                for col, value in {
                    1: "HEK293T",
                    2: "Bortezomib 1h, 500nM",
                    3: 2,
                    4: 1,
                    5: 797.5163335836667 + offset,
                    7: "parent",
                }.items():
                    sheet.cell(row, col, value)
                sheet.cell(row, config["sequence"], sequence)
                for (col, _, _), value in zip(
                    config["ratio_columns"], [1.2, 0.8, None, 1.1, 0.9, 1.0]
                ):
                    if value is not None:
                        sheet.cell(row, col, value)
                average = 22 if source == "dasgupta_2014" else 21
                sheet.cell(row, average, "=Avg Bort/Cont (1h, 500nM)")
                if offset == 5:
                    sheet.cell(row, 2, "MG132 35min 2.5 uM")
                    sheet.cell(row, average, "=Avg MG132/Cont (20min+15'wash, 2.5 uM)")
                # Real supplements store '=Avg...' as literal strings, not formulas.
                sheet.cell(row, average).data_type = "s"
        else:
            row = config["data_row"]
            sheet.cell(row, 1, "GENE")
            sheet.cell(row, 5, "A" if source == "gelman_2013" else "PEPTIDE")
            sheet.cell(row, 6, "PEPTIDE" if source == "gelman_2013" else "G")
        path = sources / filename
        book.save(path)
        specs.append(
            {
                "filename": filename,
                "asset_id": filename,
                "source_id": source,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "url": "https://example.invalid/" + filename,
                "role": "test_primary_or_heatmap",
            }
        )
    monkeypatch.setattr(tmab, "_asset_specs", lambda: specs)
    destination = tmp_path / "package"
    tmab.write_tmab_peptidomics_package(sources, destination)
    return destination, EvidenceStore.open(destination)


def test_cross_paper_candidates_keep_unique_ids_and_unresolved_independence(source_package):
    root, _ = source_package
    measurements = pq.read_table(root / tmab.MEASUREMENTS_TABLE).to_pylist()
    lineage = pq.read_table(root / tmab.LINEAGE_TABLE).to_pylist()
    paired = [m for m in measurements if m["peptide_sequence_raw"] == "LGDGPKLV"]
    assert len({m["measurement_id"] for m in paired}) == 12
    assert len({m["lineage_group_id"] for m in paired}) == 6
    assert {m["source_reuse_status"] for m in paired} == {"possible_cross_paper_reuse"}
    assert all(x["independent_specimen_status"] == "unknown" for x in lineage)
    assert {m["shared_control_group_status"] for m in measurements} == {
        "source_label_proxy_unresolved_denominator"
    }
    assert not any(x["lineage_status"] == "source_unique" for x in lineage)


def test_raw_numeric_exactness_does_not_override_identity_or_role(source_package):
    root, _ = source_package
    rows = pq.read_table(root / tmab.MEASUREMENTS_TABLE).to_pylist()
    for row in rows:
        if row["peptide_sequence_raw"] is None or "(" in row["peptide_sequence_raw"]:
            assert row["eligibility_status"].startswith("deferred_")
        if row["ratio_role"] != "treated_vs_control":
            assert row["eligibility_status"].startswith("deferred_")
    tentative = next(row for row in rows if row["peptide_sequence_raw"] == "(Ac-AGKQAVSASGK)")
    assert tentative["peptide_sequence"] == "AGKQAVSASGK"
    assert tentative["sequence_identity_status"] == "tentative"
    assert tentative["quantitative_value"] == 1.2


def test_heatmap_projection_uses_source_specific_sequence_column(source_package):
    root, _ = source_package
    rows = pq.read_table(root / tmab.ROWS_TABLE).to_pylist()
    heatmap = [r for r in rows if r["source_row_disposition"] == "deferred_heatmap_republication"]
    assert len(heatmap) == 2
    assert {r["source_sequence_raw"] for r in heatmap} == {"PEPTIDE"}


def test_protocol_scope_and_timeline_are_queryable(source_package):
    _, store = source_package
    contexts = store.query(
        "tmab_protocol_context",
        select=[
            "source_id",
            "treatment_name",
            "instrument",
            "publishing_protocol_instrument",
            "instrument_status",
            "duration_hours",
            "incubation_duration_hours",
            "harvest_duration_hours",
            "harvest_compound_status",
            "average_annotation_raw",
        ],
        limit=None,
    ).table.to_pylist()
    new = next(
        c for c in contexts if c["source_id"] == "dasgupta_2014" and c["treatment_name"] == "mg132"
    )
    assert new["instrument"] == new["publishing_protocol_instrument"] == "Waters Synapt G2"
    assert new["duration_hours"] == pytest.approx(35 / 60)
    assert new["incubation_duration_hours"] == pytest.approx(20 / 60)
    assert new["harvest_duration_hours"] == pytest.approx(15 / 60)
    assert new["harvest_compound_status"] == "present_per_study_protocol"
    reused = next(
        c
        for c in contexts
        if c["source_id"] == "dasgupta_2014" and c["treatment_name"] == "bortezomib"
    )
    assert reused["instrument"] is None
    assert reused["instrument_status"] == "unresolved_reused_acquisition"
    assert reused["publishing_protocol_instrument"] == "Waters Synapt G2"
    asset = store.query(
        "tmab_source_asset", select=["source_asset_role"], limit=None
    ).table.to_pylist()
    assert {a["source_asset_role"] for a in asset} == {"test_primary_or_heatmap"}


def test_public_package_rejects_mutated_raw_sources(source_package):
    root, _ = source_package
    source = next((root / tmab.SOURCE_ROOT).glob("*.xlsx"))
    source.write_bytes(source.read_bytes() + b"tampered")
    with pytest.raises(ValueError, match="Changed TMAB artifact"):
        EvidenceStore.open(root)


def test_manifest_keeps_correction_revision_explicit(source_package):
    root, _ = source_package
    manifest = json.loads((root / "manifest.json").read_text())
    assert manifest["revision"] == "tmab_native_peptidomics_v3"


def test_numeric_zero_stays_exact_but_is_deferred_from_log_ratio_candidates():
    identity = tmab._sequence_identity("PEPTIDE")
    assert tmab._parse_ratio(0)[:4] == (0.0, None, "exact", "reported_exact")
    assert (
        tmab._eligibility(identity, "treated_vs_control", 0, "exact")
        == "deferred_nonpositive_ratio"
    )
    assert (
        tmab._eligibility(identity, "treated_vs_control", 0.2, "source_declared_bound")
        == "candidate_censored"
    )


def test_unknown_original_acquisition_is_not_filled_from_republishing_protocol():
    context = tmab._context_for("gelman_2013", "asset", "file", "HEK293T", "epoximicin 1hr 0.2 uM")
    assert context["instrument"] is None
    assert context["publishing_protocol_instrument"] == "Waters Synapt G1"
    assert context["harvest_compound_status"] == "unresolved"


def test_handling_variants_do_not_overwrite_one_source_context():
    first = tmab._context_for(
        "dasgupta_2014",
        "asset",
        "file",
        "HEK293T",
        "MG132 35min 2.5uM",
        "20min+15'wash",
    )
    second = tmab._context_for(
        "dasgupta_2014",
        "asset",
        "file",
        "HEK293T",
        "MG132 35min 2.5uM",
        "20min+10'wash",
    )
    assert first["context_id"] != second["context_id"]
    assert first["harvest_duration_hours"] == 0.25
    assert second["harvest_duration_hours"] == pytest.approx(1 / 6)


def test_unidentified_peaks_can_have_candidate_reuse_without_sequence_labels():
    key = tmab._candidate_fingerprint("HEK293T", None, 2, 1234.5, [1.1, 0.9, None, None, None])
    assert key is not None
    assert (
        tmab._candidate_fingerprint("HEK293T", None, 2, 1234.5, [True, False, None, None, None])
        is None
    )
