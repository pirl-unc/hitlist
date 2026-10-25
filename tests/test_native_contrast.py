"""Physical contrast invariants, source decisions and public recipe replay."""

import csv
import hashlib
import json
import math
import shutil
from collections import defaultdict

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from hitlist import native_contrast as recipe_module
from hitlist import native_peptidomics as source
from hitlist.evidence_recipe import RecipeError, verify_recipe_package
from hitlist.evidence_store import EvidenceStore
from hitlist.native_contrast_compiler import compile_contrasts
from hitlist.native_contrast_schema import RECIPE_ID, SAMPLES, TABLES
from hitlist.provenance import file_digest
from tests import test_native_peptidomics as source_fixture

original = source_fixture.original
SIGNALS = {
    "0": [10, 20, 40, 100, 200, 400],
    "1": [5, 4, 20, 2, 80, 100],
}


def refresh_contract(assets, contract):
    for asset in contract["assets"]:
        asset.update(file_digest(assets / asset["filename"]))
    source.CONTRACT.write_text(json.dumps(contract))


@pytest.fixture
def qualified(original):
    assets, package, contract = original
    (assets / contract["reference"]).write_text(">sp|R1|PARENT_HUMAN OX=9606\nMACDEFGHIKLMNPQRST\n")
    for name in ("peptides.txt", "modificationSpecificPeptides.txt", "evidence.txt"):
        with (assets / name).open(newline="") as stream:
            rows = list(csv.DictReader(stream, delimiter="\t"))
        for row in rows:
            row["Proteins"] = "R1"
            if name == "peptides.txt" and row["id"] in SIGNALS:
                for sample, value in zip(SAMPLES, SIGNALS[row["id"]], strict=True):
                    row["Intensity " + sample] = str(value)
                    row["LFQ intensity " + sample] = str(100000 / value)
        source_fixture.write_tsv(assets / name, rows)
    workbook = openpyxl.load_workbook(assets / contract["publisher_workbook"])
    sheet = workbook.active
    columns = {cell.value: cell.column for cell in sheet[1]}
    for row in sheet.iter_rows(min_row=2):
        sheet.cell(row[0].row, columns["Proteins"], "R1")
    workbook.save(assets / contract["publisher_workbook"])
    refresh_contract(assets, contract)
    source.write_native_peptidomics_package(assets, package)
    return EvidenceStore.open(package)


def recipe():
    return {"schema_version": 1, "name": "software-fixture", "registered_recipe": RECIPE_ID}


def table(directory, name):
    return pq.read_table(directory / "typed" / (name + ".parquet")).to_pylist()


def materialize(store, destination):
    store.materialize(recipe(), destination)
    return verify_recipe_package(destination, source_package=store.directory)


def test_public_recipe_exact_unpaired_target_lineage_and_support(qualified, tmp_path):
    output = tmp_path / "contrast"
    manifest = materialize(qualified, output)
    examples = table(output, "examples")
    assert len(examples) == 1
    row = examples[0]
    peptides = qualified.query("native_peptides", limit=None).table.to_pylist()
    native = {p["peptide_id"]: p["native_id"] for p in peptides}
    a, b = SIGNALS[native[row["peptide_id"]]], SIGNALS[native[row["reference_peptide_id"]]]
    # Independent product-of-ratios arithmetic, without pairing specimen labels.
    expected = (
        math.log10((math.prod(a[3:]) / math.prod(b[3:])) / (math.prod(a[:3]) / math.prod(b[:3])))
        / 3
    )
    assert abs(expected) > 0.1
    assert row["target_log10_ratio_change"] == pytest.approx(expected, abs=1e-14)
    assert len(json.loads(row["acquisition_log_ratios_json"])) == 6
    contributors = table(output, "contributors")
    assert len(contributors) == len(json.loads(row["measurement_ids_json"])) == 12
    assert len({r["measurement_id"] for r in contributors}) == 12
    assert {r["sample_label_raw"] for r in contributors} == set(SAMPLES)
    replay = math.fsum(
        r["target_coefficient"] * math.log10(r["quantitative_value"]) for r in contributors
    )
    assert replay == pytest.approx(expected, abs=1e-14)
    assert row["weight"] == 1
    assert manifest["output_kind"] == "registered_native_contrast"
    assert manifest["coverage"]["qualified_for_screen"] is False
    assert manifest["coverage"]["numeric_target_summaries_disclosed"] is False
    assert any(
        c["contract"]["recipe_id"] == RECIPE_ID for c in qualified.catalog()["registered_recipes"]
    )
    covariates = table(output, "covariates")
    assert "target_log10_ratio_change" not in covariates[0]
    assert "acquisition_log_ratios_json" not in covariates[0]


def test_original_dependency_members_and_all_exclusion_reasons_survive(qualified, tmp_path):
    output = tmp_path / "contrast"
    materialize(qualified, output)
    decisions = table(output, "decisions")
    assert len(decisions) == 3
    excluded = next(r for r in decisions if r["native_id"] == "2")
    assert excluded["status"] == "excluded"
    assert set(json.loads(excluded["reasons_json"])) >= {
        "source_contaminant_flag",
        "no_deposited_reference_origin",
        "nonpositive_or_missing_source_signal",
    }
    components = table(output, "components")
    assert len(components) == 3 and len({r["dependency_group_id"] for r in components}) == 1
    assert len({r["split"] for r in components}) == 1
    reference = next(r for r in decisions if r["status"] == "reference")
    eligible = [r for r in decisions if r["status"] != "excluded"]
    expected_reference = min(
        eligible,
        key=lambda r: hashlib.sha256(
            (
                "pxd018722-native-reference-v1|"
                + r["parent_sequence_sha256"]
                + "|"
                + r["peptide_sequence"]
            ).encode()
        ).digest(),
    )
    assert reference["peptide_id"] == expected_reference["peptide_id"]
    assert any(json.loads(r["modified_form_support_json"])["Cont1"] for r in eligible)
    assert any(
        "MULTI-MATCH" in json.loads(r["feature_types_by_acquisition_json"])["Cont1"]
        for r in eligible
    )


def copy_tables(store, destination):
    destination.mkdir()
    for name in TABLES:
        shutil.copyfile(
            store.directory / source.ROOT / "typed" / (name + ".parquet"),
            destination / (name + ".parquet"),
        )


def change_table(directory, name, transform):
    path = directory / (name + ".parquet")
    original = pq.read_table(path)
    rows = original.to_pylist()
    transform(rows)
    pq.write_table(pa.Table.from_pylist(rows, schema=original.schema), path)


def test_precursor_run_loading_and_fixed_peptide_response_cancel(qualified, tmp_path):
    original_tables, changed_tables = tmp_path / "original", tmp_path / "changed"
    copy_tables(qualified, original_tables)
    copy_tables(qualified, changed_tables)
    acqs = pq.read_table(changed_tables / "acquisitions.parquet").to_pylist()
    factors = {r["acquisition_id"]: f for r, f in zip(acqs, [7, 0.3, 2, 11, 0.2, 3], strict=True)}

    def rescale(rows):
        peptide_factors = dict(
            zip(sorted({r["peptide_id"] for r in rows}), [1e-4, 1e4, 1.0], strict=True)
        )
        for row in rows:
            if row["acquisition_id"] and row["quantitative_value"] is not None:
                row["quantitative_value"] *= (
                    factors[row["acquisition_id"]] * peptide_factors[row["peptide_id"]]
                )
                row["quantitative_raw"] = str(row["quantitative_value"])

    change_table(changed_tables, "quantities", rescale)
    compile_contrasts(original_tables, tmp_path / "a" / "typed")
    compile_contrasts(changed_tables, tmp_path / "b" / "typed")
    a, b = table(tmp_path / "a", "examples"), table(tmp_path / "b", "examples")
    assert len(a) == len(b) == 1
    for field in ("example_id", "reference_peptide_id", "split", "weight"):
        assert a[0][field] == b[0][field]
    assert a[0]["target_log10_ratio_change"] == pytest.approx(
        b[0]["target_log10_ratio_change"], abs=1e-14
    )


@pytest.mark.parametrize(
    "defect", ["zero", "missing", "unresolved_parent", "engineered", "il_origin"]
)
def test_source_uncertainty_excludes_without_imputation(qualified, tmp_path, defect):
    snapshot = tmp_path / "source"
    copy_tables(qualified, snapshot)
    peptides = pq.read_table(snapshot / "peptides.parquet").to_pylist()
    pid = next(p["peptide_id"] for p in peptides if p["native_id"] == "0")
    if defect in {"zero", "missing"}:

        def mutate(rows):
            row = next(
                r
                for r in rows
                if r["peptide_id"] == pid
                and r["observation_kind"] == "peptide_sequence"
                and r["quantity_kind"] == "source_intensity"
                and r["experiment_raw"] == "KD1"
            )
            row.update(
                quantitative_value=0 if defect == "zero" else None,
                quantitative_raw="0" if defect == "zero" else "",
                value_status=defect,
            )

        change_table(snapshot, "quantities", mutate)
    elif defect == "unresolved_parent":

        def mutate(rows):
            row = next(r for r in rows if r["peptide_id"] == pid)
            row["reference_record_id"] = None

        change_table(snapshot, "declared_parents", mutate)
    elif defect == "engineered":
        change_table(
            snapshot,
            "references",
            lambda rows: rows[0].update(reference_role="engineered_construct"),
        )
    else:

        def mutate(rows):
            next(r for r in rows if r["peptide_id"] == pid)["origin_match_type"] = (
                "isoleucine_leucine_equivalent"
            )

        change_table(snapshot, "origins", mutate)
    summary = compile_contrasts(snapshot, tmp_path / "result" / "typed")
    assert summary["n_contrasts"] == 0 and not summary["qualified_for_screen"]
    decisions = table(tmp_path / "result", "decisions")
    assert all(r["status"] == "excluded" for r in decisions)
    assert len(table(tmp_path / "result", "components")) == 3


def test_tampered_target_rejected_even_after_refreshing_artifact_hashes(qualified, tmp_path):
    output = tmp_path / "contrast"
    materialize(qualified, output)
    change_table(
        output / "typed", "examples", lambda rows: rows[0].update(target_log10_ratio_change=99)
    )
    manifest = json.loads((output / "manifest.json").read_text())
    manifest["files"] = recipe_module._inventory(output)
    (output / "manifest.json").write_text(json.dumps(manifest))
    with pytest.raises(RecipeError, match="source replay differs"):
        verify_recipe_package(output, source_package=qualified.directory)


def test_recipe_policy_is_fixed_and_existing_output_is_preserved(qualified, tmp_path):
    output = tmp_path / "contrast"
    materialize(qualified, output)
    with pytest.raises(FileExistsError):
        qualified.materialize(recipe(), output)
    changed = {**recipe(), "allow_zero": True}
    with pytest.raises(RecipeError, match="Invalid native same-parent"):
        qualified.materialize(changed, tmp_path / "invalid")


def test_weights_balance_components_then_parents_then_contrasts(tmp_path):
    # Exercise weighting on source-generated contrasts with unequal parent sizes.
    # This fixture uses canonical table schemas but has no real biological labels.
    from hitlist.native_peptidomics_schema import SCHEMAS as INPUT_SCHEMAS

    source_tables = {name: [] for name in TABLES}
    for label in SAMPLES:
        source_tables["acquisitions"].append({"acquisition_id": label, "sample_label_raw": label})
    parent_peptides = [
        ["ACDEFGHIK", "FGHIKLMNP"],
        ["QWERTYIPK", "YIPKASDFG", "ASDFGHVCK"],
        ["CVNMLEASQ", "KPFHDAWYL", "RQPNKDAVE", "VTSMFGEHI"],
        ["NYVCRDMLT", "KLFQADSMH"],
    ]
    for parent_index, sequences in enumerate(parent_peptides):
        parent = "parent" + str(parent_index)
        sequence = "".join(sequences)
        digest = hashlib.sha256(sequence.encode()).hexdigest()
        source_tables["references"].append(
            {
                "reference_record_id": parent,
                "reference_sequence_sha256": digest,
                "reference_sequence": sequence,
                "reference_role": "deposited_reference_endogenous_status_not_guaranteed_by_taxonomy",
                "organism_taxon_id": "9606",
                "accession_raw": parent,
                "source_record_id": parent,
            }
        )
        for i, seq in enumerate(sequences):
            pid = parent + ":" + str(i)
            source_tables["peptides"].append(
                {
                    "peptide_id": pid,
                    "native_id": pid,
                    "source_record_id": pid,
                    "peptide_sequence": seq,
                    "posterior_error_probability": 0.001,
                    "reverse_raw": "",
                    "contaminant_raw": "",
                }
            )
            source_tables["origins"].append(
                {"peptide_id": pid, "reference_record_id": parent, "origin_match_type": "literal"}
            )
            source_tables["dependency_groups"].append(
                {
                    "peptide_id": pid,
                    "dependency_group_id": "shared" if parent_index < 3 else "independent",
                    "dependency_policy": "all_source",
                }
            )
            for label in SAMPLES:
                source_tables["quantities"].append(
                    {
                        "measurement_id": pid + ":" + label,
                        "peptide_id": pid,
                        "source_record_id": pid,
                        "acquisition_id": label,
                        "observation_kind": "peptide_sequence",
                        "quantity_kind": "source_intensity",
                        "source_field": "Intensity " + label,
                        "quantitative_value": float(i + 1),
                        "quantitative_raw": str(i + 1),
                        "value_status": "positive",
                    }
                )
    # An excluded source claim connects the first three parents; excluding it
    # must not split their source-wide component. The fourth stays independent.
    source_tables["peptides"].append(
        {
            "peptide_id": "bridge",
            "native_id": "bridge",
            "source_record_id": "bridge",
            "peptide_sequence": "AAAAAAAA",
            "posterior_error_probability": 1.0,
            "reverse_raw": "",
            "contaminant_raw": "+",
        }
    )
    source_tables["dependency_groups"].append(
        {
            "peptide_id": "bridge",
            "dependency_group_id": "shared",
            "dependency_policy": "all_source",
        }
    )
    for i in range(3):
        source_tables["declared_parents"].append(
            {
                "peptide_id": "bridge",
                "reference_record_id": "parent" + str(i),
            }
        )
    directory = tmp_path / "source"
    directory.mkdir()
    for name, rows in source_tables.items():
        pq.write_table(
            pa.Table.from_pylist(rows, schema=INPUT_SCHEMAS[name]), directory / (name + ".parquet")
        )
    summary = compile_contrasts(directory, tmp_path / "result" / "typed")
    assert summary["n_contrasts"] == 7
    assert summary["split_support"]["train"]["n_components"] == 2
    parent_weights = defaultdict(float)
    component_weights = defaultdict(float)
    for row in table(tmp_path / "result", "examples"):
        parent_weights[row["parent_sequence_sha256"]] += row["weight"]
        component_weights[row["dependency_group_id"]] += row["weight"]
    assert sorted(parent_weights.values()) == pytest.approx([1 / 6] * 3 + [1 / 2])
    assert dict(component_weights) == pytest.approx({"shared": 0.5, "independent": 0.5})
