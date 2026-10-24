"""Small source fixtures exercise identity, missingness and anti-tampering contracts."""

import json

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from hitlist import nanaware_lymph as source
from hitlist.evidence_package import package_family
from hitlist.evidence_store import EvidenceStore
from hitlist.provenance import file_digest


@pytest.fixture
def original(tmp_path, monkeypatch):
    contract = json.loads(source.CONTRACT.read_text())
    directory = tmp_path / "source"
    directory.mkdir()
    books = {}
    for block in contract["blocks"]:
        book = books.setdefault(block["filename"], openpyxl.Workbook())
        if block["sheet"] not in book:
            book.create_sheet(block["sheet"])
        sheet = book[block["sheet"]]
        for column, value in block["description"].items():
            sheet[column + "1"] = value
        for column, value in block["headers"].items():
            sheet[column + "2"] = value
    human = books["NIHMS2005243-supplement-9.xlsx"]["Suppl Table S6 (a)"]
    for index, sequence, value in [
        (3, "K.ACD(+15.99)EF.R", 10),
        (4, "K.ACD(+15.99)EF.R", None),
        (5, "AXX", 0),
        (6, "bad", "=1+1"),
    ]:
        for column, cell in {
            "A": "P1|HOST",
            "B": sequence,
            "C": True,
            "D": 12,
            "F": 5,
            "I": index - 1,
            "K": value,
            "M": index,
            "P": 1,
            "Q": 5,
        }.items():
            human[f"{column}{index}"] = cell
    dc = books["NIHMS2005243-supplement-7.xlsx"]["Suppl Table S4(e)"]
    for column, value in {
        "A": "P2|HOST",
        "B": "FGHIK",
        "C": False,
        "D": 19,
        "F": 5,
        "I": 2,
        "L": 33,
        "M": 7,
        "N": 0,
        "AF": 1,
        "AG": 5,
        "AK": "P3|HOST",
        "AL": "LMNPQ",
        "AM": True,
        "AN": 2,
        "AP": 5,
        "AS": 3,
        "AV": 44,
        "AW": 9,
        "BP": 1,
        "BQ": 5,
    }.items():
        dc[column + "3"] = value
    for column, value in {
        "AK": "P2|HOST",
        "AL": "FGHIK",
        "AM": True,
        "AN": 3,
        "AP": 5,
        "AS": 2,
        "AV": 33,
        "AW": 2,
        "BP": 1,
        "BQ": 5,
    }.items():
        dc[column + "4"] = value
    for filename, book in books.items():
        book.remove(book["Sheet"])
        book.save(directory / filename)
    xml = """<MzIdentML xmlns="http://psidev.info/psi/pi/mzIdentML/1.2">
      <SequenceCollection>
        <DBSequence id="ref1" accession="P1|HOST" searchDatabase_ref="db1"><Seq>ACDEF</Seq></DBSequence>
        <Peptide id="pep1"><PeptideSequence>ACDEF</PeptideSequence>
          <Modification location="3" monoisotopicMassDelta="15.99"/></Peptide>
        <PeptideEvidence id="ev1" peptide_ref="pep1" dBSequence_ref="ref1" start="1" end="5" isDecoy="false"/>
      </SequenceCollection></MzIdentML>"""
    (directory / "test.mzid").write_text(xml)
    contract["assets"] = [
        {
            "filename": path.name,
            "role": "original_identifications_and_reference_sequences"
            if path.suffix == ".mzid"
            else "original_quantitative_workbook",
            "source_url": "https://example.org/" + path.name,
            **file_digest(path),
        }
        for path in sorted(directory.iterdir())
    ]
    contract_path = tmp_path / "contract.json"
    contract_path.write_text(json.dumps(contract))
    monkeypatch.setattr(source, "CONTRACT", contract_path)
    return directory


def tables(package):
    return {
        name: pq.read_table(package / source.ROOT / "typed" / f"{name}.parquet").to_pylist()
        for name in source.SCHEMAS
    }


def rehash(package):
    path = package / "manifest.json"
    manifest = json.loads(path.read_text())
    manifest["files"] = source._inventory(package)
    path.write_text(json.dumps(manifest))


def test_original_blocks_forms_missingness_and_shared_samples(original, tmp_path):
    package = tmp_path / "package"
    source.write_nanaware_lymph_package(original, package)
    store = EvidenceStore.open(package)
    assert package_family(store.manifest) == "nanaware_lymph"
    assert set(store.bindings) == {"lymph_" + name for name in source.SCHEMAS}
    assert (
        store.catalog()["source_registry"]["nanaware2024"]["training_status"]
        == "source_evidence_only_no_registered_recipe"
    )
    data = tables(package)
    assert len(data["observations"]) == 7
    assert len(data["quantities"]) == 40
    assert len(data["samples"]) == 7
    assert len({r["sample_id"] for r in data["quantities"]}) == 7
    parsed = [r for r in data["observations"] if r["peptide_sequence"] == "ACDEF"]
    assert len(parsed) == 2 and {r["charge_raw"] for r in parsed} == {"2", "3"}
    assert {r["flank_n_raw"] for r in parsed} == {"K"}
    assert {r["peptide_raw"] for r in parsed} == {"K.ACD(+15.99)EF.R"}
    assert data["references"][0]["reference_sequence"] == "ACDEF"
    assert (
        data["identifications"][0]["reference_record_id"]
        == data["references"][0]["reference_record_id"]
    )
    rows = store.query(
        "lymph_quantities",
        select=["peptide_sequence", "value_status", "quantitative_value"],
        limit=None,
    ).table.to_pylist()
    assert {r["value_status"] for r in rows} == {
        "positive",
        "zero",
        "missing",
        "formula_not_evaluated",
    }
    assert any(r["peptide_sequence"] is None and r["value_status"] == "zero" for r in rows)
    dc = [r for r in data["observations"] if r["block_id"].startswith("dc_")]
    assert [(r["block_id"], r["peptide_sequence"]) for r in dc] == [
        ("dc_dia_database", "FGHIK"),
        ("dc_dia_library", "LMNPQ"),
        ("dc_dia_library", "FGHIK"),
    ]
    assert all(r["specimen_id"] is None for r in data["samples"])
    assert all(
        r["component_state_status"] == "unreported_not_assumed_wildtype" for r in data["contexts"]
    )
    with pytest.raises(FileExistsError):
        source.write_nanaware_lymph_package(original, package)


def test_typed_tamper_detected_even_with_rehashed_inventory(original, tmp_path):
    package = tmp_path / "package"
    source.write_nanaware_lymph_package(original, package)
    path = package / source.ROOT / "typed/quantities.parquet"
    table = pq.read_table(path)
    rows = table.to_pylist()
    rows[0]["quantitative_value"] += 1
    pq.write_table(pa.Table.from_pylist(rows, schema=table.schema), path)
    rehash(package)
    with pytest.raises(ValueError, match="typed rows differ"):
        source.verify_nanaware_lymph_package(package)


def test_source_tamper_stops_atomic_write(original, tmp_path):
    path = original / "NIHMS2005243-supplement-9.xlsx"
    path.write_bytes(path.read_bytes() + b"changed")
    destination = tmp_path / "package"
    with pytest.raises(ValueError, match="source mismatch"):
        source.write_nanaware_lymph_package(original, destination)
    assert not destination.exists()
    assert not list(tmp_path.glob(".package-*"))


def test_header_drift_is_rejected_even_with_updated_fixture_digest(original, tmp_path):
    path = original / "NIHMS2005243-supplement-9.xlsx"
    book = openpyxl.load_workbook(path)
    book["Suppl Table S6 (a)"]["K2"] = "Different kind of quantity"
    book.save(path)
    contract = json.loads(source.CONTRACT.read_text())
    for asset in contract["assets"]:
        if asset["filename"] == path.name:
            asset.update(file_digest(path))
    source.CONTRACT.write_text(json.dumps(contract))
    with pytest.raises(ValueError, match="header drift"):
        source.write_nanaware_lymph_package(original, tmp_path / "package")


@pytest.mark.parametrize(
    "raw,status",
    [
        (False, "invalid_literal"),
        ("NaN", "nonfinite"),
        (-3, "negative"),
        ("=2+2", "formula_not_evaluated"),
    ],
)
def test_quantity_state_is_not_silently_coerced(raw, status):
    assert source._number(raw)[1] == status
